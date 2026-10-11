import os
import sys
import tempfile
import textwrap
import unittest

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from processors import loan_tracker_processor


class LoanTrackerProcessorTests(unittest.TestCase):
    @staticmethod
    def _tracking_headers():
        return [
            "First Name",
            "Other Names",
            "Account No",
            "Field Officer",
            "Disbursement Date",
            "Disbursed Amount",
            "Initial / Current OLB",
            "Expected Repayment Days",
            "Current Week",
            "Expected Balance",
            "Missed Days",
            "Required Catch-up Amount",
        ]

    @staticmethod
    def _worksheet_row_values(worksheet, row_number, column_count):
        return [
            worksheet.cell(row=row_number, column=column_index).value
            for column_index in range(1, column_count + 1)
        ]

    @staticmethod
    def _load_tracking_sheet(output_file):
        workbook = load_workbook(output_file)
        return workbook, workbook["Loan Tracking"]

    def test_run_generates_week_tracking_and_catchup_columns(self):
        csv_text = textwrap.dedent(
            """
            FullNames,Amount Disbursed,Borrow Date,Product,OLB,Clear Date
            "Jane Doe","Ksh 3,000.00",18/03/2026,"5 WEEKS CREDIT EASE","Ksh 3,600.00",22/04/2026
            "John Doe","Ksh 9,999.00",18/03/2026,"5 WEEKS CREDIT EASE","Ksh 11,000.00",22/04/2026
            """
        ).strip()

        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = os.path.join(temp_dir, "report.csv")
            output_dir = os.path.join(temp_dir, "out")

            with open(report_path, "w", encoding="utf-8", newline="") as handle:
                handle.write(csv_text)

            result = loan_tracker_processor.run(
                report_path,
                output_dir,
                as_of_date="2026-03-24",
            )

            self.assertTrue(result["success"], result)
            self.assertTrue(os.path.isfile(result["output_file"]))
            self.assertEqual(result["stats"]["loans_processed"], 1)
            self.assertEqual(result["stats"]["rows_with_errors"], 1)
            self.assertEqual(result["stats"]["loans_in_arrears"], 1)

            workbook, worksheet = self._load_tracking_sheet(result["output_file"])
            headers = self._tracking_headers()

            self.assertEqual(worksheet["A1"].value, "LOAN TRACKER REPORT")
            self.assertEqual(worksheet["A2"].value, "As Of Date: 24/03/2026")
            self.assertEqual(worksheet["A3"].value, "Disbursement Date: 18/03/2026")
            self.assertEqual(
                self._worksheet_row_values(worksheet, 4, len(headers)),
                headers,
            )

            self.assertEqual(worksheet["A5"].value, "Jane")
            self.assertEqual(worksheet["B5"].value, "Doe")
            self.assertIsNone(worksheet["C5"].value)
            self.assertIsNone(worksheet["D5"].value)
            self.assertEqual(worksheet["E5"].value.strftime("%d/%m/%Y"), "18/03/2026")
            self.assertEqual(int(worksheet["F5"].value), 3000)
            self.assertEqual(worksheet["G5"].value, "4,000 / 3,600")
            self.assertEqual(int(worksheet["H5"].value), 5)
            self.assertEqual(worksheet["I5"].value, "Week 1")
            self.assertAlmostEqual(float(worksheet["J5"].value), 3333.35, places=2)
            self.assertAlmostEqual(float(worksheet["K5"].value), 2.0, places=2)
            self.assertAlmostEqual(float(worksheet["L5"].value), 266.65, places=2)

            self.assertNotIn("Product", headers)
            self.assertNotIn("Term Weeks", headers)
            self.assertNotIn("Review Note", headers)
            workbook.close()

            errors_df = pd.read_excel(result["output_file"], sheet_name="Errors")
            self.assertIn("Review Note", errors_df.columns)
            self.assertIn("Unsupported loan size", str(errors_df.loc[0, "Review Note"]))

    def test_run_can_infer_term_from_clear_date(self):
        csv_text = textwrap.dedent(
            """
            FullNames,Amount Disbursed,Borrow Date,OLB,Clear Date
            "Mary Doe","Ksh 5,000.00",18/03/2026,"Ksh 500.00",22/04/2026
            """
        ).strip()

        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = os.path.join(temp_dir, "report.csv")
            output_dir = os.path.join(temp_dir, "out")

            with open(report_path, "w", encoding="utf-8", newline="") as handle:
                handle.write(csv_text)

            result = loan_tracker_processor.run(
                report_path,
                output_dir,
                as_of_date="2026-04-30",
            )

            self.assertTrue(result["success"], result)

            workbook, worksheet = self._load_tracking_sheet(result["output_file"])
            self.assertEqual(worksheet["I5"].value, "Week 5")
            self.assertEqual(int(worksheet["H5"].value), 30)
            self.assertAlmostEqual(float(worksheet["J5"].value), 0.0, places=2)
            self.assertAlmostEqual(float(worksheet["L5"].value), 500.0, places=2)
            workbook.close()

    def test_run_rejects_unsupported_term_logic(self):
        csv_text = textwrap.dedent(
            """
            FullNames,Amount Disbursed,Borrow Date,Product,OLB
            "Alex Doe","Ksh 5,000.00",18/03/2026,"8 WEEKS CREDIT EASE","Ksh 6,500.00"
            """
        ).strip()

        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = os.path.join(temp_dir, "report.csv")
            output_dir = os.path.join(temp_dir, "out")

            with open(report_path, "w", encoding="utf-8", newline="") as handle:
                handle.write(csv_text)

            result = loan_tracker_processor.run(
                report_path,
                output_dir,
                as_of_date="2026-03-24",
            )

            self.assertTrue(result["success"], result)
            errors_df = pd.read_excel(result["output_file"], sheet_name="Errors")
            self.assertEqual(len(errors_df), 1)
            self.assertIn("Unsupported term 8 weeks", str(errors_df.loc[0, "Review Note"]))

    def test_run_groups_multiple_dates_sorts_officers_and_colors_catchup(self):
        csv_text = textwrap.dedent(
            """
            FullNames,Account No,Field Officer,Amount Disbursed,Borrow Date,Product,OLB
            "Zed Alpha",111,"B Officer","Ksh 3,000.00",19/03/2026,"5 WEEKS CREDIT EASE","Ksh 4,000.00"
            "Amy Beta",222,"A Officer","Ksh 3,000.00",19/03/2026,"5 WEEKS CREDIT EASE","Ksh 3,400.00"
            "Chris Gamma",333,"C Officer","Ksh 3,000.00",18/03/2026,"5 WEEKS CREDIT EASE","Ksh 3,600.00"
            """
        ).strip()

        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = os.path.join(temp_dir, "report.csv")
            output_dir = os.path.join(temp_dir, "out")

            with open(report_path, "w", encoding="utf-8", newline="") as handle:
                handle.write(csv_text)

            result = loan_tracker_processor.run(
                report_path,
                output_dir,
                as_of_date="2026-03-24",
            )

            self.assertTrue(result["success"], result)

            workbook, worksheet = self._load_tracking_sheet(result["output_file"])
            self.assertEqual(worksheet["A3"].value, "Disbursement Date: 19/03/2026")
            self.assertEqual(worksheet["A8"].value, "Disbursement Date: 18/03/2026")

            self.assertEqual(worksheet["A5"].value, "Amy")
            self.assertEqual(worksheet["D5"].value, "A Officer")
            self.assertEqual(worksheet["A6"].value, "Zed")
            self.assertEqual(worksheet["D6"].value, "B Officer")
            self.assertEqual(worksheet["A10"].value, "Chris")

            self.assertEqual(worksheet["L5"].fill.fgColor.rgb, "FFC6E0B4")
            self.assertEqual(worksheet["L6"].fill.fgColor.rgb, "FFC00000")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
