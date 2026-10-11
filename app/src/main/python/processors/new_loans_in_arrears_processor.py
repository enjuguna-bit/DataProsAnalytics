"""
New Loans in Arrears Processor
MVP enhancement of the New Loans in Arrears report.
Adds OLB-based arrears rate (alongside the existing count-based rate),
a centralized portfolio classifier (RC / SME / UNKNOWN), and
portfolio + officer breakdowns with a validation summary.

Report layout (v2):
  1. Executive Summary  - headline figures
  2. Portfolio Summary  - portfolio totals, per-officer breakdown,
                          and the full new-loans-in-arrears list
                          grouped by officer (A-Z) with risk tiers
                          (Early/Elevated/High/Critical) ascending
                          within each officer block
  3. Validation         - data-quality checks with OK/REVIEW status
  4. Data               - full audit-trail dataset as an Excel table
"""

import os
import re
from datetime import datetime
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from processors.tabular_io import read_tabular_file

RC_PATTERN = re.compile(r"\d+\s*WEEKS?\s*CREDIT\s*EASE", re.IGNORECASE)


def classify_portfolio(product_name: str) -> str:
    """Centralized portfolio classifier. Add new rules here only."""
    if not isinstance(product_name, str) or not product_name.strip():
        return "UNKNOWN"
    name = product_name.upper()
    if "SME" in name:
        return "SME"
    if RC_PATTERN.search(name):
        return "RC"
    return "UNKNOWN"


def parse_money(value) -> float:
    if pd.isna(value):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace("Ksh", "").replace(",", "").strip() or 0)


def _normalize_cols(df: pd.DataFrame, mapping: dict) -> pd.DataFrame:
    """Finds matching columns by checking aliases (case-insensitive, ignoring non-alphanumeric) and renames them."""
    rename_dict = {}
    normalized_cols = {re.sub(r"[^a-z0-9]", "", c.lower()): c for c in df.columns}
    for target, aliases in mapping.items():
        for alias in aliases:
            norm_alias = re.sub(r"[^a-z0-9]", "", alias.lower())
            if norm_alias in normalized_cols:
                rename_dict[normalized_cols[norm_alias]] = target
                break
    return df.rename(columns=rename_dict)


def _normalize_name(name):
    if not isinstance(name, str):
        return str(name) if name else ""
    parts = re.split(r"\s+", name.strip())
    return " ".join(p.capitalize() for p in parts if p)


def load_disbursement(path: str) -> pd.DataFrame:
    df = read_tabular_file(path, dtype=str)
    # Normalize standard column names
    df = _normalize_cols(df, {
        "LoanId": ["LoanId", "Loan Id", "LoanID", "Loan_ID", "LoanNo", "Loan No"],
        "OLB": ["OLB", "AmountDisbursed", "Amount Disbursed", "LoanAmount", "Loan Amount"],
        "Product": ["Product", "ProductName", "Product Name", "ProductType"],
        "Disbursement Date": ["Disbursement Date", "DisbursementDate", "Disb Date", "DateDisbursed", "Date", "BorrowDate", "Borrow Date"],
        "FieldOfficer": ["FieldOfficer", "Field Officer", "SalesRep", "Sales Rep", "Officer", "Agent"],
        "FullNames": ["FullNames", "Full Name", "Customer Name", "Name", "BorrowerName"],
        "Branch": ["Branch", "BranchName", "Branch Name"],
    })

    # Fill in missing columns with default/empty values if not found, to prevent crash
    for col in ["LoanId", "OLB", "Product", "Disbursement Date", "FieldOfficer", "FullNames", "Branch"]:
        if col not in df.columns:
            df[col] = ""

    df["LoanId"] = df["LoanId"].fillna("").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    df["OLB"] = df["OLB"].apply(parse_money)
    df["Portfolio"] = df["Product"].apply(classify_portfolio)
    df["Disbursement Date"] = pd.to_datetime(df["Disbursement Date"], errors="coerce")
    return df


def load_arrears(path: str) -> pd.DataFrame:
    df = read_tabular_file(path, dtype=str)
    # Normalize standard column names
    df = _normalize_cols(df, {
        "LoanId": ["LoanId", "Loan Id", "LoanID", "Loan_ID", "LoanNo", "Loan No"],
        "LoanBalance": ["LoanBalance", "Loan Balance", "Balance", "Outstanding"],
        "ProductName": ["ProductName", "Product", "Product Name", "ProductType"],
        "DaysInArrears": ["DaysInArrears", "Days In Arrears", "Days Overdue", "DaysOverdue", "Overdue Days"],
        "SalesRep": ["SalesRep", "Sales Rep", "Officer", "FieldOfficer", "Field Officer", "Agent"],
        "FullNames": ["FullNames", "Full Name", "Customer Name", "Name", "BorrowerName"],
        "Branch": ["Branch", "BranchName", "Branch Name"],
        "ArrearsAmount": ["ArrearsAmount", "Arrears Amount", "Amount In Arrears", "AmountInArrears",
                          "OverdueAmount", "Overdue Amount", "ArrearsDue", "Arrears Due",
                          "InstallmentArrears", "Installment Arrears", "ArrearsInstallment"],
        "Phone": ["Phone", "PhoneNumber", "Phone Number", "MobileNumber", "Mobile Number",
                  "MobileNo", "Mobile No", "Contact", "ContactNumber", "Contact Number",
                  "PhoneNo", "Phone No", "Telephone", "Tel"],
    })

    for col in ["LoanId", "LoanBalance", "ProductName", "DaysInArrears", "SalesRep", "FullNames",
                "Branch", "ArrearsAmount", "Phone"]:
        if col not in df.columns:
            df[col] = ""

    df["LoanId"] = df["LoanId"].fillna("").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    df["LoanBalance"] = df["LoanBalance"].apply(parse_money)
    df["Portfolio"] = df["ProductName"].apply(classify_portfolio)
    df["DaysInArrears"] = pd.to_numeric(df["DaysInArrears"], errors="coerce").fillna(0)
    df["ArrearsAmount"] = df["ArrearsAmount"].apply(parse_money)
    df["Phone"] = df["Phone"].fillna("").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    return df


def build_joined_dataset(disb: pd.DataFrame, arr: pd.DataFrame, window_days: Optional[int] = None,
                          as_of: Optional[datetime] = None) -> pd.DataFrame:
    """
    Left-joins new disbursements (optionally windowed to the last
    `window_days`) against the FULL arrears book on LoanId.
    One row per new loan. InArrears=1 if the LoanId also appears
    in the arrears export.
    """
    d = disb.copy()
    if window_days is not None:
        as_of = as_of or d["Disbursement Date"].max()
        cutoff = as_of - pd.Timedelta(days=window_days)
        d = d[d["Disbursement Date"] >= cutoff]

    a_cols = ["LoanId", "LoanBalance", "DaysInArrears", "ProductName", "SalesRep", "ArrearsAmount", "Phone"]
    if "FullNames" in arr.columns:
        a_cols.append("FullNames")
    if "Branch" in arr.columns:
        a_cols.append("Branch")

    a = arr[a_cols].copy()
    a = a.sort_values("DaysInArrears", ascending=False).drop_duplicates(subset=["LoanId"])

    rename_map = {"ProductName": "ArrearsProductName", "SalesRep": "ArrearsOfficer"}
    if "FullNames" in a.columns:
        rename_map["FullNames"] = "ArrearsFullNames"
    if "Branch" in a.columns:
        rename_map["Branch"] = "ArrearsBranch"
    a = a.rename(columns=rename_map)

    joined = d.merge(a, on="LoanId", how="left", indicator=True)
    joined["InArrears"] = (joined["_merge"] == "both").astype(int)
    joined["ArrearsOLB"] = joined["LoanBalance"].where(joined["InArrears"] == 1, 0.0)
    joined["ArrearsAmount"] = joined["ArrearsAmount"].where(joined["InArrears"] == 1, 0.0)
    joined["Phone"] = joined["Phone"].fillna("").astype(str)
    joined["DaysInArrears"] = joined["DaysInArrears"].where(joined["InArrears"] == 1)
    joined = joined.drop(columns=["_merge", "LoanBalance"])

    # Fill in FullNames/Branch from arrears if missing on disbursement side
    if "ArrearsFullNames" in joined.columns:
        joined["FullNames"] = joined["FullNames"].fillna("").astype(str)
        mask = (joined["FullNames"] == "") | (joined["FullNames"].isna())
        joined.loc[mask, "FullNames"] = joined.loc[mask, "ArrearsFullNames"]
        joined = joined.drop(columns=["ArrearsFullNames"])

    if "ArrearsBranch" in joined.columns:
        joined["Branch"] = joined["Branch"].fillna("").astype(str)
        mask = (joined["Branch"] == "") | (joined["Branch"].isna())
        joined.loc[mask, "Branch"] = joined.loc[mask, "ArrearsBranch"]
        joined = joined.drop(columns=["ArrearsBranch"])

    # Portfolio mismatch between disbursement-side and arrears-side product name
    # (e.g. a loan later refinanced onto a different product). Flag, don't overwrite.
    joined["PortfolioMismatch"] = (
        (joined["InArrears"] == 1)
        & (joined["ArrearsProductName"].apply(classify_portfolio) != joined["Portfolio"])
    )
    return joined


def _native(value):
    """Cast numpy/pandas scalar types to plain Python types.
    Chaquopy marshals return values across the JNI boundary to Kotlin.
    """
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def compute_validation(disb: pd.DataFrame, arr: pd.DataFrame, joined: pd.DataFrame) -> dict:
    return {
        "Total New Loans": _native(len(disb)),
        "Matched to Arrears Book": _native(joined["InArrears"].sum()),
        "Unmatched (not in arrears)": _native((joined["InArrears"] == 0).sum()),
        "Duplicate Loan IDs (disbursement)": _native(disb["LoanId"].duplicated().sum()),
        "Duplicate Loan IDs (arrears)": _native(arr["LoanId"].duplicated().sum()),
        "Unknown Products (disbursement)": _native((disb["Portfolio"] == "UNKNOWN").sum()),
        "Unknown Products (arrears)": _native((arr["Portfolio"] == "UNKNOWN").sum()),
        "Missing Officer (disbursement)": _native(disb["FieldOfficer"].isna().sum() if "FieldOfficer" in disb.columns else 0),
        "Missing OLB (disbursement)": _native(disb["OLB"].isna().sum() if "OLB" in disb.columns else 0),
        "Portfolio Mismatches (disb vs arrears product)": _native(joined["PortfolioMismatch"].sum()),
    }


# Checks where a nonzero value signals something worth reviewing (vs. a plain count like "Total New Loans")
VALIDATION_FLAG_KEYS = {
    "Duplicate Loan IDs (disbursement)", "Duplicate Loan IDs (arrears)",
    "Unknown Products (disbursement)", "Unknown Products (arrears)",
    "Missing Officer (disbursement)", "Missing OLB (disbursement)",
    "Portfolio Mismatches (disb vs arrears product)",
}


def _rate_table(joined: pd.DataFrame, group_col) -> pd.DataFrame:
    g = joined.groupby(group_col, dropna=False)
    out = g.agg(
        New_Loans=("LoanId", "count"),
        New_OLB=("OLB", "sum"),
        Loans_In_Arrears=("InArrears", "sum"),
        Arrears_OLB=("ArrearsOLB", "sum"),
    ).reset_index()
    out["Count_Rate"] = out["Loans_In_Arrears"] / out["New_Loans"]
    out["OLB_Rate"] = out["Arrears_OLB"] / out["New_OLB"].replace(0, pd.NA)
    total_arrears_olb = out["Arrears_OLB"].sum()
    out["Contribution_Pct"] = out["Arrears_OLB"] / total_arrears_olb if total_arrears_olb else 0
    return out.sort_values("OLB_Rate", ascending=False, na_position="last")


def compute_portfolio_summary(joined: pd.DataFrame) -> pd.DataFrame:
    return _rate_table(joined, "Portfolio")


def compute_officer_summary(joined: pd.DataFrame) -> pd.DataFrame:
    return _rate_table(joined, "FieldOfficer")


def compute_officer_portfolio_summary(joined: pd.DataFrame) -> pd.DataFrame:
    return _rate_table(joined, ["FieldOfficer", "Portfolio"])


def compute_executive_summary(joined: pd.DataFrame, portfolio_summary: pd.DataFrame,
                               officer_summary: pd.DataFrame) -> dict:
    total_new_olb = joined["OLB"].sum()
    total_arrears_olb = joined["ArrearsOLB"].sum()
    ranked_pf = portfolio_summary.dropna(subset=["OLB_Rate"])
    ranked_of = officer_summary.dropna(subset=["OLB_Rate"])
    return {
        "Total New Loans": _native(len(joined)),
        "Total New OLB": _native(total_new_olb),
        "Total Loans in Arrears": _native(joined["InArrears"].sum()),
        "Total OLB in Arrears": _native(total_arrears_olb),
        "Count-Based Arrears Rate": _native(joined["InArrears"].sum() / len(joined) if len(joined) else 0),
        "OLB-Based Arrears Rate": _native(total_arrears_olb / total_new_olb if total_new_olb else 0),
        "Highest Risk Portfolio": ranked_pf.iloc[0]["Portfolio"] if not ranked_pf.empty else "N/A",
        "Lowest Risk Portfolio": ranked_pf.iloc[-1]["Portfolio"] if not ranked_pf.empty else "N/A",
        "Highest Risk Officer": ranked_of.iloc[0]["FieldOfficer"] if not ranked_of.empty else "N/A",
        "Lowest Risk Officer": ranked_of.iloc[-1]["FieldOfficer"] if not ranked_of.empty else "N/A",
    }


def risk_tier(days) -> tuple:
    """Returns (label, fill_hex) for a DaysInArrears value."""
    try:
        d = float(days)
    except (TypeError, ValueError):
        d = 0.0
    if d <= 3:
        return "Early", "FFF2CC"
    elif d <= 7:
        return "Elevated", "FCE4D6"
    elif d <= 14:
        return "High", "F8CBAD"
    else:
        return "Critical", "F4A6A6"


def loans_in_arrears_list(joined: pd.DataFrame) -> pd.DataFrame:
    """Full list of new loans matched to the arrears book (no cap),
    sorted by Field Officer (A-Z) then Days In Arrears (ascending / risk
    tier ascending) so each officer's own loans read mildest-to-worst.
    """
    cols = ["FullNames", "Phone", "FieldOfficer", "Portfolio", "ArrearsOLB", "ArrearsAmount", "DaysInArrears"]
    for c in cols:
        if c not in joined.columns:
            joined[c] = ""
    return (
        joined[joined["InArrears"] == 1][cols]
        .sort_values(["FieldOfficer", "DaysInArrears"], ascending=[True, True])
    )


# ---------------------------------------------------------------------------
# Report writer
# ---------------------------------------------------------------------------

NAVY = "1F4E78"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", bold=True, size=14, color=NAVY)
SUBTITLE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
SECTION_FONT = Font(name="Arial", bold=True, size=11, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor=NAVY)
LABEL_FONT = Font(name="Arial", bold=True)
BODY_FONT = Font(name="Arial")
TOTAL_FONT = Font(name="Arial", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
OFFICER_HDR_FONT = Font(name="Arial", bold=True, color=NAVY)
OFFICER_HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT_FMT = "0.0%"
MONEY_FMT = '"KSh" #,##0;("KSh" #,##0);"-"'
DATE_FMT = "dd-mmm-yyyy"


def _write_table(ws, df: pd.DataFrame, start_row: int, pct_cols=(), money_cols=()) -> int:
    headers = [str(c).replace("_", " ") for c in df.columns]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            col_name = df.columns[j - 1]
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            col_name_cleaned = col_name.replace("_", " ")
            if col_name in pct_cols or col_name_cleaned in pct_cols:
                cell.number_format = PCT_FMT
            elif col_name in money_cols or col_name_cleaned in money_cols:
                cell.number_format = MONEY_FMT
    for j, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(14, len(h) + 2)
    return start_row + len(df) + 1


def _heat_map(ws, col_letter: str, first_row: int, last_row: int):
    if last_row < first_row:
        return
    rule = ColorScaleRule(
        start_type="min", start_color="63BE7B",
        mid_type="percentile", mid_value=50, mid_color="FFEB84",
        end_type="max", end_color="F8696B",
    )
    ws.conditional_formatting.add(f"{col_letter}{first_row}:{col_letter}{last_row}", rule)


def _write_section_title(ws, row, text, span=8):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    for c in range(1, span + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
    return row + 1


def _write_table_header(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    return row + 1


def _write_portfolio_summary_block(ws, row: int, portfolio_summary: pd.DataFrame) -> int:
    row = _write_section_title(ws, row, "PORTFOLIO SUMMARY")
    headers = ["Portfolio", "New Loans", "New OLB", "Arrears OLB", "OLB Rate",
               "Loans In Arrears", "Count Rate", "Contribution %"]
    row = _write_table_header(ws, row, headers)
    start = row
    for _, r in portfolio_summary.iterrows():
        ws.cell(row=row, column=1, value=r["Portfolio"]).font = BODY_FONT
        ws.cell(row=row, column=2, value=int(r["New_Loans"]))
        ws.cell(row=row, column=3, value=r["New_OLB"])
        ws.cell(row=row, column=4, value=r["Arrears_OLB"])
        ws.cell(row=row, column=5, value=r["OLB_Rate"] if pd.notna(r["OLB_Rate"]) else 0)
        ws.cell(row=row, column=6, value=int(r["Loans_In_Arrears"]))
        ws.cell(row=row, column=7, value=r["Count_Rate"])
        ws.cell(row=row, column=8, value=r["Contribution_Pct"])
        for c in range(1, 9):
            cell = ws.cell(row=row, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
        ws.cell(row=row, column=3).number_format = MONEY_FMT
        ws.cell(row=row, column=4).number_format = MONEY_FMT
        ws.cell(row=row, column=5).number_format = PCT_FMT
        ws.cell(row=row, column=7).number_format = PCT_FMT
        ws.cell(row=row, column=8).number_format = PCT_FMT
        row += 1
    end = row - 1

    # Total row (formulas so it stays correct if rows are edited later)
    ws.cell(row=row, column=1, value="Total")
    ws.cell(row=row, column=2, value=f"=SUM(B{start}:B{end})")
    ws.cell(row=row, column=3, value=f"=SUM(C{start}:C{end})")
    ws.cell(row=row, column=4, value=f"=SUM(D{start}:D{end})")
    ws.cell(row=row, column=5, value=f"=IFERROR(D{row}/C{row},0)")
    ws.cell(row=row, column=6, value=f"=SUM(F{start}:F{end})")
    ws.cell(row=row, column=7, value=f"=IFERROR(F{row}/B{row},0)")
    ws.cell(row=row, column=8, value=1)
    for c in range(1, 9):
        cell = ws.cell(row=row, column=c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
    ws.cell(row=row, column=3).number_format = MONEY_FMT
    ws.cell(row=row, column=4).number_format = MONEY_FMT
    ws.cell(row=row, column=5).number_format = PCT_FMT
    ws.cell(row=row, column=7).number_format = PCT_FMT
    ws.cell(row=row, column=8).number_format = PCT_FMT
    return row + 2


def _write_officer_breakdown_block(ws, row: int, officer_portfolio_summary: pd.DataFrame,
                                    rc_flag_threshold: float = 0.20) -> int:
    row = _write_section_title(ws, row, "PER OFFICER BREAKDOWN")
    headers = ["Field Officer", "Portfolio", "New Loans", "New OLB", "Arrears OLB",
               "OLB Rate", "Loans In Arrears", "Count Rate", "Contribution %"]
    row = _write_table_header(ws, row, headers)

    ordered = officer_portfolio_summary.sort_values(
        ["Portfolio", "Contribution_Pct"], ascending=[True, False]
    )
    for _, r in ordered.iterrows():
        olb_rate = r["OLB_Rate"] if pd.notna(r["OLB_Rate"]) else 0
        ws.cell(row=row, column=1, value=r["FieldOfficer"])
        ws.cell(row=row, column=2, value=r["Portfolio"])
        ws.cell(row=row, column=3, value=int(r["New_Loans"]))
        ws.cell(row=row, column=4, value=r["New_OLB"])
        ws.cell(row=row, column=5, value=r["Arrears_OLB"])
        ws.cell(row=row, column=6, value=olb_rate)
        ws.cell(row=row, column=7, value=int(r["Loans_In_Arrears"]))
        ws.cell(row=row, column=8, value=r["Count_Rate"])
        ws.cell(row=row, column=9, value=r["Contribution_Pct"])
        for c in range(1, 10):
            cell = ws.cell(row=row, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
        ws.cell(row=row, column=4).number_format = MONEY_FMT
        ws.cell(row=row, column=5).number_format = MONEY_FMT
        ws.cell(row=row, column=6).number_format = PCT_FMT
        ws.cell(row=row, column=8).number_format = PCT_FMT
        ws.cell(row=row, column=9).number_format = PCT_FMT
        # Flag RC officers whose OLB Rate exceeds the threshold
        if r["Portfolio"] == "RC" and olb_rate > rc_flag_threshold:
            for c in range(1, 10):
                ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="FDEDEC")
        row += 1
    return row + 1


def _write_arrears_list_block(ws, row: int, arrears_df: pd.DataFrame) -> int:
    row = _write_section_title(
        ws, row, f"NEW LOANS IN ARREARS  ({len(arrears_df)} loans, by officer then risk tier)"
    )
    headers = ["Borrower", "Phone", "Field Officer", "Portfolio", "Arrears OLB",
               "Current Arrears Amount", "Days In Arrears", "Risk Tier"]
    row = _write_table_header(ws, row, headers)

    officers_in_order = sorted(arrears_df["FieldOfficer"].unique())
    for oi, officer in enumerate(officers_in_order):
        officer_rows = arrears_df[arrears_df["FieldOfficer"] == officer]
        officer_total = officer_rows["ArrearsOLB"].sum()

        # Officer sub-header: name + loan count + total exposure
        ws.cell(row=row, column=1, value=f"{officer}  —  {len(officer_rows)} loans")
        ws.cell(row=row, column=5, value=officer_total)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        for c in range(1, 9):
            cell = ws.cell(row=row, column=c)
            cell.font = OFFICER_HDR_FONT
            cell.fill = OFFICER_HDR_FILL
            cell.border = BORDER
        ws.cell(row=row, column=5).number_format = MONEY_FMT
        ws.cell(row=row, column=5).alignment = Alignment(horizontal="center")
        row += 1

        for _, r in officer_rows.iterrows():
            tier, color = risk_tier(r["DaysInArrears"])
            ws.cell(row=row, column=1, value=r["FullNames"])
            ws.cell(row=row, column=2, value=r["Phone"])
            ws.cell(row=row, column=3, value=r["FieldOfficer"])
            ws.cell(row=row, column=4, value=r["Portfolio"])
            ws.cell(row=row, column=5, value=r["ArrearsOLB"])
            ws.cell(row=row, column=6, value=r["ArrearsAmount"])
            ws.cell(row=row, column=7, value=int(r["DaysInArrears"]) if pd.notna(r["DaysInArrears"]) else 0)
            ws.cell(row=row, column=8, value=tier)
            for c in range(1, 9):
                cell = ws.cell(row=row, column=c)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
                cell.fill = PatternFill("solid", fgColor=color)
            ws.cell(row=row, column=5).number_format = MONEY_FMT
            ws.cell(row=row, column=6).number_format = MONEY_FMT
            row += 1

        if oi < len(officers_in_order) - 1:
            row += 1  # blank separator row between officers

    # Legend
    row += 1
    ws.cell(row=row, column=1, value="Risk Tiers:").font = Font(name="Arial", size=9, bold=True)
    legend_items = [("1-3 days: Early", "FFF2CC"), ("4-7 days: Elevated", "FCE4D6"),
                    ("8-14 days: High", "F8CBAD"), ("15+ days: Critical", "F4A6A6")]
    for i, (label, color) in enumerate(legend_items):
        cell = ws.cell(row=row, column=2 + i, value=label)
        cell.font = Font(name="Arial", size=9)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center")
    return row + 2


def write_report(output_path: str, joined: pd.DataFrame, portfolio_summary: pd.DataFrame,
                  officer_portfolio_summary: pd.DataFrame, exec_summary: dict, validation: dict,
                  arrears_list: pd.DataFrame):
    wb = Workbook()

    # --- Executive Summary ---
    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = "New Loans in Arrears - Executive Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = SUBTITLE_FONT
    row = 4
    pct_keys = {"Count-Based Arrears Rate", "OLB-Based Arrears Rate"}
    money_keys = {"Total New OLB", "Total OLB in Arrears"}
    for label, value in exec_summary.items():
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        c = ws.cell(row=row, column=2, value=value)
        c.font = BODY_FONT
        if label in pct_keys:
            c.number_format = PCT_FMT
        elif label in money_keys:
            c.number_format = MONEY_FMT
        row += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    # --- Portfolio Summary (combined: totals, per-officer, arrears list) ---
    ws = wb.create_sheet("Portfolio Summary")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "New Loans in Arrears Report"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:I1")
    ws["A2"] = f"Generated {datetime.now():%d %b %Y}"
    ws["A2"].font = SUBTITLE_FONT
    ws.merge_cells("A2:I2")

    row = _write_portfolio_summary_block(ws, 4, portfolio_summary)
    row = _write_officer_breakdown_block(ws, row, officer_portfolio_summary)
    _write_arrears_list_block(ws, row, arrears_list)

    widths = {"A": 26, "B": 15, "C": 18, "D": 12, "E": 14, "F": 17, "G": 13, "H": 14, "I": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A1"

    # --- Validation ---
    ws = wb.create_sheet("Validation")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Data Validation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now():%d %b %Y}"
    ws["A2"].font = SUBTITLE_FONT
    row = 4
    row = _write_table_header(ws, row, ["Check", "Value", "Status"])
    for label, value in validation.items():
        ws.cell(row=row, column=1, value=label).font = BODY_FONT
        vcell = ws.cell(row=row, column=2, value=value)
        vcell.font = BODY_FONT
        vcell.alignment = Alignment(horizontal="center")
        if label in VALIDATION_FLAG_KEYS:
            status = "OK" if not value else "REVIEW"
            scell = ws.cell(row=row, column=3, value=status)
            scell.font = Font(name="Arial", bold=True, color="2E7D32" if status == "OK" else "C62828")
            scell.fill = PatternFill("solid", fgColor="E8F5E9" if status == "OK" else "FDECEA")
            scell.alignment = Alignment(horizontal="center")
        for c in range(1, 4):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.freeze_panes = "A5"

    # --- Raw joined data (audit trail) ---
    ws = wb.create_sheet("Data")
    ws.sheet_view.showGridLines = False
    export_cols = ["LoanId", "FullNames", "FieldOfficer", "Portfolio", "Branch",
                   "Disbursement Date", "OLB", "InArrears", "ArrearsOLB",
                   "DaysInArrears", "PortfolioMismatch"]
    for c in export_cols:
        if c not in joined.columns:
            joined[c] = ""
    data_df = joined[export_cols].copy()
    data_df["InArrears"] = data_df["InArrears"].map({1: "Yes", 0: "No"})

    headers = [c.replace("_", " ") for c in export_cols]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for i, (_, r) in enumerate(data_df.iterrows(), start=2):
        in_arrears = r["InArrears"] == "Yes"
        row_fill = PatternFill("solid", fgColor="FDECEA") if in_arrears else None
        for j, col in enumerate(export_cols, start=1):
            val = r[col]
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if col not in ("FullNames", "FieldOfficer") else "left")
            if col == "OLB" or col == "ArrearsOLB":
                cell.number_format = MONEY_FMT
            elif col == "Disbursement Date":
                cell.number_format = DATE_FMT
            if row_fill:
                cell.fill = row_fill

    ws3_widths = {"A": 11, "B": 26, "C": 18, "D": 10, "E": 9, "F": 14, "G": 12,
                  "H": 10, "I": 13, "J": 11, "K": 14}
    for col, w in ws3_widths.items():
        ws.column_dimensions[col].width = w

    last_row = ws.max_row
    if last_row > 1:
        tab = Table(displayName="DataTable", ref=f"A1:K{last_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(tab)
    ws.freeze_panes = "A2"

    wb.save(output_path)


def run(disb_path: str, arr_path: str, output_path: str, window_days: Optional[int] = None):
    disb = load_disbursement(disb_path)
    arr = load_arrears(arr_path)
    joined = build_joined_dataset(disb, arr, window_days=window_days)

    validation = compute_validation(disb, arr, joined)
    portfolio_summary = compute_portfolio_summary(joined)
    officer_summary = compute_officer_summary(joined)
    officer_portfolio_summary = compute_officer_portfolio_summary(joined)
    exec_summary = compute_executive_summary(joined, portfolio_summary, officer_summary)
    arrears_list = loans_in_arrears_list(joined)

    # Determine output file path
    if os.path.isdir(output_path):
        output_file_path = os.path.join(output_path, "new_loans_in_arrears_report.xlsx")
    else:
        output_file_path = output_path

    # Ensure parent directories exist
    parent_dir = os.path.dirname(output_file_path)
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)

    write_report(output_file_path, joined, portfolio_summary, officer_portfolio_summary,
                 exec_summary, validation, arrears_list)

    # Calculate stats for Kotlin compatibility
    total_disbursed = len(disb)
    matched_count = int(joined["InArrears"].sum())
    match_pct = float(matched_count / total_disbursed * 100) if total_disbursed else 0.0

    if "DaysInArrears" in joined.columns:
        max_days_val = float(joined["DaysInArrears"].max())
        if pd.isna(max_days_val):
            max_days_val = 0.0
    else:
        max_days_val = 0.0

    total_arr_amt = float(joined["ArrearsOLB"].sum())

    return {
        "success": True,
        "output_file": output_file_path,
        "stats": {
            "total_disbursed": total_disbursed,
            "matched": matched_count,
            "match_pct": match_pct,
            "max_days_in_arrears": max_days_val,
            "total_arrears_amount": total_arr_amt,
        },
        "validation": validation,
        "executive_summary": exec_summary,
    }
