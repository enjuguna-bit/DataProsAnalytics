"""
RO Performance Global Processor
Adapted from "RO perfomance global.py" for Android/Chaquopy (no tkinter).
"""

import hashlib
import os
import re
from collections import defaultdict
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt


WEIGHTS = {
    "income": 0.35,
    "disbursement": 0.30,
    "collection_rate": 0.25,
    "loan_count": 0.05,
    "portfolio_quality": 0.05,
}

THRESHOLDS = {
    "excellent": 85,
    "good": 70,
    "average": 50,
}

REGIONS = {
    "NAIROBI": ["KASARANI", "RUIRU", "THIKA", "MLOLONGO", "RUAI", "NGONG", "KITENGELA"],
    "RIFT VALLEY": [
        "NAKURU",
        "ELDAMA RAVINE",
        "NYAHURURU",
        "MOLO",
        "NJORO",
        "NAIVASHA",
        "GILGIL",
        "SUBUKIA",
        "KABARNET",
        "MARIGAT",
        "ELDORET",
        "KAPSABET",
        "KITALE",
        "KAPENGURIA",
        "MARALAL",
    ],
    "WESTERN": ["KAKAMEGA", "BUNGOMA", "WEBUYE", "MUMIAS", "BUSIA", "MALABA"],
    "NYANZA": ["KISUMU", "SIAYA", "HOMABAY", "MIGORI", "OYUGIS", "KISII", "KEHANCHA"],
    "COAST": [
        "MOMBASA",
        "MALINDI",
        "KILIFI",
        "MTWAPA",
        "UKUNDA",
        "KWALE",
        "VOI",
        "WUNDANYI",
        "TAVETA",
        "MARIAKANI",
        "LIKONI",
        "BAMBURI",
        "KISAUNI",
        "CHANGAMWE",
    ],
    "EASTERN": ["EMBU", "MERU", "CHUKA", "RUNYENJES", "NKUBU", "MAUA", "KERUGOYA", "KARATINA", "NYERI", "CHOGORIA", "NANYUKI", "ISIOLO"],
    "CENTRAL": ["KIAMBU", "LIMURU", "TALA", "THIKA", "MURANGA", "KENOL"],
    "NORTH EASTERN": ["GARISSA", "WAJIR", "MANDERA"],
    "SOUTHERN": ["KAJIADO", "LOITOKTOK", "KIBWEZI", "MAKINDU", "MTITO ANDEI"],
}

ZONES = {
    "ZONE A": ["NAIROBI", "CENTRAL"],
    "ZONE B": ["RIFT VALLEY", "WESTERN"],
    "ZONE C": ["NYANZA", "COAST"],
    "ZONE D": ["EASTERN", "NORTH EASTERN", "SOUTHERN"],
}


def _read(path: str) -> pd.DataFrame:
    if path.lower().endswith(".csv"):
        df = pd.read_csv(path, dtype=str)
    else:
        df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _clean_string(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text.upper()


def _norm(value: str) -> str:
    return _clean_string(value).replace(" ", "").replace("_", "").replace("-", "")


def _find_column(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    norm_map = {_norm(col): col for col in df.columns}
    for candidate in candidates:
        key = _norm(candidate)
        if key in norm_map:
            return norm_map[key]
    for candidate in candidates:
        key = _norm(candidate)
        for col in df.columns:
            col_key = _norm(col)
            if key in col_key or col_key in key:
                return col
    return None


def _to_float(value: object) -> float:
    if value is None or pd.isna(value):
        return 0.0
    cleaned = re.sub(r"[^\d\.\-]", "", str(value))
    if cleaned in ("", "-", "."):
        return 0.0
    try:
        return float(cleaned)
    except Exception:
        return 0.0


def _determine_region_zone(branch: str) -> Tuple[str, str]:
    branch_clean = _clean_string(branch)

    region = "UNKNOWN"
    for reg, branches in REGIONS.items():
        if any(branch_clean == _clean_string(b) for b in branches):
            region = reg
            break

    zone = "UNKNOWN"
    for zn, regions in ZONES.items():
        if region in regions:
            zone = zn
            break
    return region, zone


def _branch_code(branch: str) -> str:
    clean = re.sub(r"[^A-Z]", "", _clean_string(branch))
    prefix = clean[:3] if clean else "BRN"
    hash_part = hashlib.md5(branch.encode("utf-8")).hexdigest()[:3].upper()
    return f"{prefix}{hash_part}"


def _key(name: str, branch: str) -> str:
    return f"{_clean_string(name)}::{_clean_string(branch)}"


def _performance_category(score: float) -> str:
    if score >= THRESHOLDS["excellent"]:
        return "Excellent"
    if score >= THRESHOLDS["good"]:
        return "Good"
    if score >= THRESHOLDS["average"]:
        return "Average"
    return "Poor"


def _calculate_scores(df: pd.DataFrame) -> pd.DataFrame:
    max_income = max(float(df["Income"].max()), 1.0)
    max_disbursed = max(float(df["Disbursed_Amount"].max()), 1.0)
    max_loans = max(float(df["Loan_Count"].max()), 1.0)

    income_score = (df["Income"] / max_income) * 100
    disbursed_score = (df["Disbursed_Amount"] / max_disbursed) * 100
    collection_score = df["Collection_Rate"]
    loan_score = (df["Loan_Count"] / max_loans) * 100
    quality_score = 100 - df["Delinquency_Rate"]

    score = (
        income_score * WEIGHTS["income"]
        + disbursed_score * WEIGHTS["disbursement"]
        + collection_score * WEIGHTS["collection_rate"]
        + loan_score * WEIGHTS["loan_count"]
        + quality_score * WEIGHTS["portfolio_quality"]
    )
    df["Position_Score"] = score.round(2)
    return df


def _dense_rank(series: pd.Series, ascending: bool = False) -> pd.Series:
    return series.rank(method="dense", ascending=ascending).astype(int)


def _build_chart(df: pd.DataFrame, chart_path: str) -> str:
    try:
        if df.empty:
            fig, ax = plt.subplots(figsize=(8, 4))
            ax.text(0.5, 0.5, "No data", ha="center", va="center")
            ax.axis("off")
            fig.savefig(chart_path, dpi=150, bbox_inches="tight")
            plt.close(fig)
            return chart_path

        top_ro = df.nlargest(15, "Position_Score")
        branch_avg = (
            df.groupby("Branch", as_index=False)["Position_Score"].mean().sort_values("Position_Score", ascending=False).head(10)
        )

        fig, axes = plt.subplots(1, 2, figsize=(14, 6))

        axes[0].barh(
            top_ro["Full_Name"][::-1],
            top_ro["Position_Score"].values[::-1],
            color="#1A237E",
            edgecolor="white",
        )
        axes[0].set_title("Top 15 ROs by Position Score", fontweight="bold")
        axes[0].set_xlabel("Score")

        axes[1].bar(
            branch_avg["Branch"],
            branch_avg["Position_Score"],
            color="#00C853",
            edgecolor="white",
        )
        axes[1].set_title("Top 10 Branches by Avg Position Score", fontweight="bold")
        axes[1].tick_params(axis="x", rotation=45)
        axes[1].set_ylabel("Avg Score")

        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _auto_width(ws):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)


def _write_dataframe(ws, df: pd.DataFrame):
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=str(col))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        for col_idx, col in enumerate(df.columns, 1):
            value = row[col]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if isinstance(value, (int, float, np.floating)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    _auto_width(ws)


def _prepare_results(df: pd.DataFrame) -> Dict[str, object]:
    branch_summary = (
        df.groupby("Branch")
        .agg(
            Number_of_ROs=("Full_Name", "count"),
            Income=("Income", "sum"),
            Disbursed_Amount=("Disbursed_Amount", "sum"),
            Avg_Collection_Rate=("Collection_Rate", "mean"),
            Loan_Count=("Loan_Count", "sum"),
            Avg_Position_Score=("Position_Score", "mean"),
            Best_Global_Rank=("Global_Rank", "min"),
        )
        .round(2)
        .sort_values("Avg_Position_Score", ascending=False)
    )

    region_summary = (
        df.groupby("Region")
        .agg(
            Number_of_ROs=("Full_Name", "count"),
            Income=("Income", "sum"),
            Disbursed_Amount=("Disbursed_Amount", "sum"),
            Avg_Collection_Rate=("Collection_Rate", "mean"),
            Avg_Position_Score=("Position_Score", "mean"),
        )
        .round(2)
        .sort_values("Avg_Position_Score", ascending=False)
    )

    zone_summary = (
        df.groupby("Zone")
        .agg(
            Number_of_ROs=("Full_Name", "count"),
            Income=("Income", "sum"),
            Disbursed_Amount=("Disbursed_Amount", "sum"),
            Avg_Collection_Rate=("Collection_Rate", "mean"),
            Avg_Position_Score=("Position_Score", "mean"),
        )
        .round(2)
        .sort_values("Avg_Position_Score", ascending=False)
    )

    return {
        "detailed_data": df,
        "branch_summary": branch_summary,
        "region_summary": region_summary,
        "zone_summary": zone_summary,
    }


def _export_excel(output_file: str, chart_path: str, results: Dict[str, object], global_stats: Dict[str, object]):
    wb = Workbook()

    ws_detail = wb.active
    ws_detail.title = "Detailed_RO_Performance"
    _write_dataframe(ws_detail, results["detailed_data"])

    ws_branch = wb.create_sheet("Branch_Summary")
    _write_dataframe(ws_branch, results["branch_summary"].reset_index())

    ws_region = wb.create_sheet("Region_Summary")
    _write_dataframe(ws_region, results["region_summary"].reset_index())

    ws_zone = wb.create_sheet("Zone_Summary")
    _write_dataframe(ws_zone, results["zone_summary"].reset_index())

    top_20 = results["detailed_data"].head(20)
    ws_top = wb.create_sheet("Top_20_Performers_Global")
    _write_dataframe(ws_top, top_20)

    bottom_20 = results["detailed_data"].tail(20)
    ws_bottom = wb.create_sheet("Bottom_20_Performers")
    _write_dataframe(ws_bottom, bottom_20)

    distribution = (
        results["detailed_data"]["Performance_Category"].value_counts().rename_axis("Category").reset_index(name="Count")
    )
    ws_dist = wb.create_sheet("Performance_Distribution")
    _write_dataframe(ws_dist, distribution)

    summary = pd.DataFrame(
        {
            "Metric": [
                "Total ROs Analyzed",
                "Total Branches",
                "Total Regions",
                "Total Zones",
                "Total Income",
                "Total Disbursed",
                "Average Collection Rate",
                "Average Position Score",
                "Top Performer",
                "Processing Timestamp",
            ],
            "Value": [
                global_stats["total_ros"],
                global_stats["branch_count"],
                global_stats["region_count"],
                global_stats["zone_count"],
                f"KES {global_stats['total_income']:,.2f}",
                f"KES {global_stats['total_disbursed']:,.2f}",
                f"{global_stats['avg_collection_rate']:.2f}%",
                f"{global_stats['avg_position_score']:.2f}",
                global_stats["top_performer"],
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ],
        }
    )
    ws_summary = wb.create_sheet("Executive_Summary")
    _write_dataframe(ws_summary, summary)

    if chart_path and os.path.exists(chart_path):
        ws_chart = wb.create_sheet("Chart")
        ws_chart.add_image(XLImage(chart_path), "A1")

    wb.save(output_file)


def run(income_path: str, collection_rate_path: str, disbursement_path: str, output_dir: str) -> dict:
    os.makedirs(output_dir, exist_ok=True)

    income_df = _read(income_path)
    cr_df = _read(collection_rate_path)
    disb_df = _read(disbursement_path)

    income_name_col = _find_column(income_df, ["Full Name", "Name"])
    income_branch_col = _find_column(income_df, ["Branch Name", "Branch"])
    income_amount_col = _find_column(income_df, ["Total Amount", "Income", "Amount"])
    income_id_col = _find_column(income_df, ["#ID", "ID", "RO ID"])

    cr_name_col = _find_column(cr_df, ["Full Name", "Name"])
    cr_branch_col = _find_column(cr_df, ["Branch Name", "Branch"])
    cr_rate_col = _find_column(cr_df, ["Collection Rate (%)", "Collection Rate", "CR %"])
    cr_total_col = _find_column(cr_df, ["Total Amount", "Amount"])
    cr_paid_col = _find_column(cr_df, ["Paid Amount", "Collected", "Paid"])
    cr_outstanding_col = _find_column(cr_df, ["Outstanding Amount", "Outstanding", "Uncollected"])

    disb_name_col = _find_column(disb_df, ["Full Name", "Name"])
    disb_branch_col = _find_column(disb_df, ["Branch Name", "Branch"])
    disb_amount_col = _find_column(disb_df, ["Disbursed Amount", "Disbursement", "Amount"])
    disb_loan_col = _find_column(disb_df, ["Loan Count", "Loans", "LoanCount"])

    missing = []
    if income_name_col is None:
        missing.append("Income file: Full Name")
    if income_branch_col is None:
        missing.append("Income file: Branch Name")
    if cr_name_col is None:
        missing.append("Collection file: Full Name")
    if cr_branch_col is None:
        missing.append("Collection file: Branch Name")
    if disb_name_col is None:
        missing.append("Disbursement file: Full Name")
    if disb_branch_col is None:
        missing.append("Disbursement file: Branch Name")
    if missing:
        raise ValueError("Missing required columns: " + "; ".join(missing))

    profiles: Dict[str, Dict[str, object]] = {}

    for _, row in income_df.iterrows():
        name = _clean_string(row.get(income_name_col, ""))
        branch = _clean_string(row.get(income_branch_col, ""))
        if not name or not branch:
            continue

        key = _key(name, branch)
        ro_id = str(row.get(income_id_col, "")).strip() if income_id_col else ""
        if not ro_id or ro_id.lower() == "nan":
            ro_id = hashlib.md5(f"{name}_{branch}".encode("utf-8")).hexdigest()[:8].upper()

        region, zone = _determine_region_zone(branch)
        profiles[key] = {
            "ID": ro_id,
            "Full_Name": name,
            "Branch": branch,
            "Branch_Code": _branch_code(branch),
            "Region": region,
            "Zone": zone,
            "Income": _to_float(row.get(income_amount_col, 0)) if income_amount_col else 0.0,
            "Disbursed_Amount": 0.0,
            "Collection_Rate": 0.0,
            "Loan_Count": 0,
            "Active_Loans": 0,
            "Portfolio_Size": 0.0,
            "Delinquency_Rate": 0.0,
        }

    for _, row in cr_df.iterrows():
        name = _clean_string(row.get(cr_name_col, ""))
        branch = _clean_string(row.get(cr_branch_col, ""))
        key = _key(name, branch)
        if key not in profiles:
            continue

        total_amount = _to_float(row.get(cr_total_col, 0)) if cr_total_col else 0.0
        paid_amount = _to_float(row.get(cr_paid_col, 0)) if cr_paid_col else 0.0
        outstanding = _to_float(row.get(cr_outstanding_col, 0)) if cr_outstanding_col else 0.0

        if cr_rate_col:
            collection_rate = _to_float(row.get(cr_rate_col, 0))
        else:
            denom = paid_amount + outstanding if (paid_amount + outstanding) > 0 else total_amount
            collection_rate = (paid_amount / denom * 100) if denom > 0 else 0.0

        profile = profiles[key]
        profile["Collection_Rate"] = collection_rate
        profile["Portfolio_Size"] = total_amount if total_amount > 0 else (paid_amount + outstanding)
        if profile["Portfolio_Size"] > 0:
            profile["Delinquency_Rate"] = (outstanding / profile["Portfolio_Size"]) * 100
        else:
            profile["Delinquency_Rate"] = 0.0

    for _, row in disb_df.iterrows():
        name = _clean_string(row.get(disb_name_col, ""))
        branch = _clean_string(row.get(disb_branch_col, ""))
        key = _key(name, branch)
        if key not in profiles:
            continue
        profiles[key]["Disbursed_Amount"] = _to_float(row.get(disb_amount_col, 0)) if disb_amount_col else 0.0
        profiles[key]["Loan_Count"] = int(round(_to_float(row.get(disb_loan_col, 0)) if disb_loan_col else 0.0))
        profiles[key]["Active_Loans"] = profiles[key]["Loan_Count"]

    if not profiles:
        raise ValueError("No valid RO records were built from the provided files")

    df = pd.DataFrame(list(profiles.values()))
    df = _calculate_scores(df)

    df = df.sort_values(
        by=["Position_Score", "Collection_Rate", "Income"],
        ascending=[False, False, False],
    ).reset_index(drop=True)

    df["Global_Rank"] = _dense_rank(-df["Position_Score"], ascending=True)
    n = len(df)
    df["Percentile"] = ((1 - ((df["Global_Rank"] - 1) / max(n, 1))) * 100).round(2)

    df["Branch_Rank"] = (
        df.groupby("Branch")["Position_Score"].rank(method="dense", ascending=False).astype(int)
    )
    df["Region_Rank"] = (
        df.groupby("Region")["Position_Score"].rank(method="dense", ascending=False).astype(int)
    )
    df["Zone_Rank"] = (
        df.groupby("Zone")["Position_Score"].rank(method="dense", ascending=False).astype(int)
    )
    df["Experience_Months"] = 0
    df["Target_Achievement"] = 0.0
    df["Performance_Category"] = df["Position_Score"].apply(_performance_category)
    df["Last_Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    ordered_cols = [
        "ID",
        "Full_Name",
        "Branch",
        "Branch_Code",
        "Region",
        "Zone",
        "Income",
        "Disbursed_Amount",
        "Collection_Rate",
        "Loan_Count",
        "Active_Loans",
        "Portfolio_Size",
        "Delinquency_Rate",
        "Position_Score",
        "Global_Rank",
        "Branch_Rank",
        "Region_Rank",
        "Zone_Rank",
        "Percentile",
        "Experience_Months",
        "Target_Achievement",
        "Performance_Category",
        "Last_Updated",
    ]
    df = df[ordered_cols]

    results = _prepare_results(df)

    total_income = float(df["Income"].sum())
    total_disbursed = float(df["Disbursed_Amount"].sum())
    total_loans = int(df["Loan_Count"].sum())
    avg_collection_rate = float(df["Collection_Rate"].mean()) if len(df) > 0 else 0.0
    avg_score = float(df["Position_Score"].mean()) if len(df) > 0 else 0.0
    top_performer = str(df.iloc[0]["Full_Name"]) if len(df) > 0 else "N/A"

    global_stats = {
        "total_ros": int(len(df)),
        "branch_count": int(df["Branch"].nunique()),
        "region_count": int(df["Region"].nunique()),
        "zone_count": int(df["Zone"].nunique()),
        "total_income": total_income,
        "total_disbursed": total_disbursed,
        "total_loans": total_loans,
        "avg_collection_rate": avg_collection_rate,
        "avg_position_score": avg_score,
        "top_performer": top_performer,
    }

    chart_path = os.path.join(output_dir, "ro_chart.png")
    chart_file = _build_chart(df, chart_path)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"RO_Analysis_Premium_{timestamp}.xlsx")
    _export_excel(output_file, chart_file, results, global_stats)

    # Backward compatibility keys are retained for existing app cards.
    total_portfolio_size = float(df["Portfolio_Size"].sum()) if len(df) > 0 else 0.0
    avg_delinquency = float(df["Delinquency_Rate"].mean()) if len(df) > 0 else 0.0

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            # Previous keys used in app
            "total_clients": int(total_loans),
            "total_ros": int(len(df)),
            "portfolio_balance": round(total_portfolio_size, 2),
            "portfolio_par1_pct": round(avg_delinquency, 2),
            # Global script-aligned keys
            "ro_count": int(len(df)),
            "branch_count": int(df["Branch"].nunique()),
            "region_count": int(df["Region"].nunique()),
            "zone_count": int(df["Zone"].nunique()),
            "total_income": round(total_income, 2),
            "total_disbursed": round(total_disbursed, 2),
            "avg_collection_rate": round(avg_collection_rate, 2),
            "avg_position_score": round(avg_score, 2),
            "top_performer": top_performer,
        },
    }
