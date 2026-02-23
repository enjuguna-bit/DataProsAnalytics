import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
import numpy as np
import os
import sys
import argparse
import builtins
import json
import logging
import subprocess
from pathlib import Path
from textwrap import fill

# =========================================================
# CONFIGURATION
# =========================================================
CHART_DPI = 200
plt.style.use('ggplot')

# =========================================================
# 1. FUNCTION: File Selection (GUI)
# =========================================================
def select_file():
    """Opens a dialog to select the input file."""
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError:
        print("âŒ GUI libraries not available. Use command-line arguments instead.")
        sys.exit(1)

    root = tk.Tk()
    root.withdraw()  # Hide the main window
    file_path = filedialog.askopenfilename(
        title="Select Loan Data File",
        filetypes=[("Excel Files", "*.xlsx"), ("CSV Files", "*.csv")]
    )
    return file_path

# =========================================================
# 2. FUNCTION: Load and Clean Data
# =========================================================
def load_and_clean_data(filepath):
    """Loads Excel/CSV, standardizes columns, and cleans data types."""
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"The file '{filepath}' was not found.")

    try:
        if filepath.lower().endswith('.csv'):
            df = pd.read_csv(filepath, dtype=str)
        else:
            df = pd.read_excel(filepath, dtype=str)
    except Exception as e:
        raise Exception(f"Error reading file: {e}")

    if df.empty:
        raise ValueError("The file is empty.")

    # Standardize headers
    df.columns = [str(c).strip() for c in df.columns]

    # Mapping logic for column variations
    colmap = {
        "fullnames": "FullNames",
        "phonenumber": "PhoneNumber",
        "installmentno": "InstallmentNo",
        "amount due": "AmountDue",
        "arrears": "Arrears",
        "loanbalance": "LoanBalance",
        "fieldofficer": "FieldOfficer",
        "principal": "FundedAmount",
        "funded amount": "FundedAmount",
        "loan amount": "FundedAmount",
        "disbursed amount": "FundedAmount"
    }

    # Apply mapping for any variations
    renamed_cols = {}
    for original_col in df.columns:
        col_lower = original_col.lower()
        for key, target in colmap.items():
            if key in col_lower:
                renamed_cols[original_col] = target
                break
    
    if renamed_cols:
        df = df.rename(columns=renamed_cols)

    # Handle missing essential columns
    essential_columns = ["PhoneNumber", "FullNames", "Arrears"]
    for col in essential_columns:
        if col not in df.columns:
            print(f"âš ï¸  Warning: Missing essential column '{col}'. Creating with default values.")
            df[col] = "" if col in ["PhoneNumber", "FullNames"] else 0

    if "FundedAmount" not in df.columns:
        df["FundedAmount"] = 0
    if "FieldOfficer" not in df.columns:
        df["FieldOfficer"] = "Unassigned"
    if "LoanBalance" not in df.columns:
        df["LoanBalance"] = 0
    if "InstallmentNo" not in df.columns:
        df["InstallmentNo"] = 1

    # Numeric conversion with error handling
    numeric_cols = ["InstallmentNo", "AmountDue", "Arrears", "LoanBalance", "FundedAmount"]
    for c in numeric_cols:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
            if c == "InstallmentNo":
                df[c] = df[c].astype('Int64')
            else:
                df[c] = df[c].astype(float)

    # Clean phone numbers
    if "PhoneNumber" in df.columns:
        df["PhoneNumber"] = df["PhoneNumber"].astype(str).str.replace(r"\D", "", regex=True)
        df["PhoneNumber"] = df["PhoneNumber"].replace(["", "nan", "None"], "Unknown")

    # Remove completely empty rows
    df = df.dropna(how='all')

    print(f"âœ… Loaded {len(df)} records with {df['PhoneNumber'].nunique()} unique customers.")
    return df

def load_config(path: str = "premium_config.json"):
    """Loads a JSON config file containing scoring weights and thresholds."""
    try:
        p = Path(path)
        if p.exists():
            with p.open() as f:
                cfg = json.load(f)
                print(f"âœ… Loaded config from {path}")
                return cfg
        else:
            print(f"âš ï¸  Config file {path} not found. Using defaults.")
            return {}
    except Exception as e:
        print(f"âš ï¸  Failed to load config {path}: {e}")
        return {}

# =========================================================
# 3. FUNCTION: Calculate Risk
# =========================================================
def calculate_risk_metrics(df, config: dict = None):
    """Calculates risk scores and categories with configuration-driven parameters."""
    
    if df.empty:
        raise ValueError("No data available for risk calculation.")

    required_cols = ["PhoneNumber", "InstallmentNo", "Arrears", "LoanBalance"]
    missing_cols = [col for col in required_cols if col not in df.columns]
    if missing_cols:
        raise KeyError(f"Missing required columns: {missing_cols}")

    cfg = config or {}
    scoring_cfg = cfg.get('scoring', {})
    weights_cfg = cfg.get('weights', {})

    # Configuration with defaults
    arrears_quantile = scoring_cfg.get('arrears_quantile', 0.95)
    balance_quantile = scoring_cfg.get('balance_quantile', 0.95)
    missed_cap = scoring_cfg.get('missed_cap', 12)

    arrears_weight = weights_cfg.get('arrears', 0.50)
    missed_weight = weights_cfg.get('missed', 0.35)
    balance_weight = weights_cfg.get('balance', 0.15)

    missed_threshold = scoring_cfg.get('missed_threshold', 4)
    arrears_threshold = scoring_cfg.get('arrears_threshold', 5000)

    q_low = scoring_cfg.get('percentile_low', 0.40)
    q_high = scoring_cfg.get('percentile_high', 0.75)

    # Sort and get latest status
    df_sorted = df.sort_values(by=["PhoneNumber", "InstallmentNo"])
    latest_status = df_sorted.groupby("PhoneNumber").tail(1).copy()

    # Calculate missed installments
    df["MissedFlag"] = (df["Arrears"] > 0).astype(int)
    missed_inst = (
        df.groupby("PhoneNumber")["MissedFlag"]
        .sum()
        .reset_index()
        .rename(columns={"MissedFlag": "MissedInstallments"})
    )

    customer_risk = latest_status.merge(missed_inst, on="PhoneNumber", how="left")
    customer_risk["MissedInstallments"] = customer_risk["MissedInstallments"].fillna(0).astype(int)

    # Calculate caps safely
    if len(customer_risk) > 0:
        arrears_cap = customer_risk["Arrears"].quantile(arrears_quantile)
        balance_cap = customer_risk["LoanBalance"].quantile(balance_quantile)
    else:
        arrears_cap = balance_cap = 0

    arrears_capped = np.minimum(customer_risk["Arrears"], arrears_cap)
    balance_capped = np.minimum(customer_risk["LoanBalance"], balance_cap)
    missed_capped = np.minimum(customer_risk["MissedInstallments"], missed_cap)

    # Calculate risk score
    customer_risk["RiskScore"] = (
        arrears_capped * arrears_weight +
        missed_capped * missed_weight +
        balance_capped * balance_weight
    )

    # Determine percentile thresholds
    if len(customer_risk) > 0:
        q40 = customer_risk["RiskScore"].quantile(q_low)
        q75 = customer_risk["RiskScore"].quantile(q_high)
    else:
        q40 = q75 = 0

    def determine_category(row):
        if row["MissedInstallments"] >= missed_threshold:
            return "High Risk"
        if row["Arrears"] > arrears_threshold:
            return "High Risk"
        score = row["RiskScore"]
        if score >= q75:
            return "High Risk"
        elif score >= q40:
            return "Medium Risk"
        else:
            return "Low Risk"

    customer_risk["RiskCategory"] = customer_risk.apply(determine_category, axis=1)
    customer_risk = customer_risk.sort_values(by="RiskScore", ascending=False).reset_index(drop=True)

    # Log results
    print(f"âœ… Risk calculation complete: {len(customer_risk)} customers categorized")
    if len(customer_risk) > 0:
        high = (customer_risk['RiskCategory'] == 'High Risk').sum()
        medium = (customer_risk['RiskCategory'] == 'Medium Risk').sum()
        low = (customer_risk['RiskCategory'] == 'Low Risk').sum()
        print(f"   High Risk: {high} | Medium Risk: {medium} | Low Risk: {low}")

    return customer_risk

# =========================================================
# 4. FUNCTION: Build Report Tables
# =========================================================
def build_printable_risk_sheet(customer_risk):
    """Builds the main risk ranking report."""
    if customer_risk.empty:
        return pd.DataFrame()

    # Columns to exclude from report
    exclude_patterns = [
        "loanid", "loan id", 
        "amountpaid", "amount paid", 
        "loanamount", "loan amount", 
        "productname", "product name", 
        "unittitle", "unit title", 
        "duedate", "due date"
    ]

    cols_to_drop = [col for col in customer_risk.columns 
                    if col.lower() in exclude_patterns]
    
    if cols_to_drop:
        customer_risk = customer_risk.drop(columns=cols_to_drop)

    officer_blocks_list = []

    def make_spacer():
        return pd.DataFrame([{c: None for c in customer_risk.columns}])

    def make_row(officer, name_val, arrears=None, balance=None, missed=None, score=None, cat=None):
        data = {c: None for c in customer_risk.columns}
        data["FieldOfficer"] = officer
        data["FullNames"] = name_val
        if "Arrears" in data: data["Arrears"] = arrears
        if "LoanBalance" in data: data["LoanBalance"] = balance
        if "MissedInstallments" in data: data["MissedInstallments"] = missed
        if "RiskScore" in data: data["RiskScore"] = score
        if "RiskCategory" in data: data["RiskCategory"] = cat
        return pd.DataFrame([data])

    # Build officer portfolios
    for officer_name, officer_df in customer_risk.groupby("FieldOfficer"):
        officer_blocks_list.append(make_row(officer_name, f"{officer_name} PORTFOLIO"))

        for cat in ["High Risk", "Medium Risk", "Low Risk"]:
            block = officer_df[officer_df["RiskCategory"] == cat].copy()
            if block.empty:
                continue

            block = block.sort_values("RiskScore", ascending=False)
            officer_blocks_list.append(block)

            subtotal_row = make_row(
                officer_name, f"{cat} SUBTOTAL",
                arrears=block["Arrears"].sum() if "Arrears" in block.columns else 0,
                balance=block["LoanBalance"].sum() if "LoanBalance" in block.columns else 0,
                missed=block["MissedInstallments"].sum() if "MissedInstallments" in block.columns else 0,
                score=block["RiskScore"].sum() if "RiskScore" in block.columns else 0,
                cat=cat
            )
            officer_blocks_list.append(subtotal_row)
        
        officer_blocks_list.append(make_spacer())

    if not officer_blocks_list:
        return pd.DataFrame()

    return pd.concat(officer_blocks_list, ignore_index=True)

def build_early_arrears_report(df, customer_risk):
    """Builds early arrears report for loans at risk in first 2 installments."""
    if customer_risk.empty:
        return pd.DataFrame(columns=[
            "FieldOfficer", "FullNames", "PhoneNumber", "InstallmentNo",
            "AmountDue", "Arrears", "ArrearsPercentage", "SeverityIndicator",
            "DaysInArrears", "RiskCategory", "AmountDisbursed"
        ])

    early_mask = (customer_risk["InstallmentNo"] <= 2) & (customer_risk["Arrears"] > 0)
    early_df = customer_risk.loc[early_mask].copy()

    if early_df.empty:
        return pd.DataFrame(columns=[
            "FieldOfficer", "FullNames", "PhoneNumber", "InstallmentNo",
            "AmountDue", "Arrears", "ArrearsPercentage", "SeverityIndicator",
            "DaysInArrears", "RiskCategory", "AmountDisbursed"
        ])

    early_df = early_df.sort_values(by=["FieldOfficer", "Arrears"], ascending=[True, False])

    # Calculate metrics
    if "AmountDue" in early_df.columns and early_df["AmountDue"].sum() > 0:
        early_df["ArrearsPercentage"] = (early_df["Arrears"] / early_df["AmountDue"] * 100).round(2)
    else:
        early_df["ArrearsPercentage"] = 0

    early_df["DaysInArrears"] = early_df["InstallmentNo"] * 7

    if "FundedAmount" in early_df.columns:
        early_df["AmountDisbursed"] = early_df["FundedAmount"].fillna(0)
    else:
        early_df["AmountDisbursed"] = 0

    def get_severity_color(percentage):
        if percentage > 50:
            return "CRITICAL"
        elif percentage >= 25:
            return "MEDIUM"
        else:
            return "LOW"

    early_df["SeverityIndicator"] = early_df["ArrearsPercentage"].apply(get_severity_color)

    # Build report with subtotals
    cols = [
        "FieldOfficer", "FullNames", "PhoneNumber", "InstallmentNo",
        "AmountDue", "Arrears", "ArrearsPercentage", "SeverityIndicator",
        "DaysInArrears", "RiskCategory", "AmountDisbursed"
    ]

    available_cols = [col for col in cols if col in early_df.columns]
    blocks = []

    def create_row(data_dict):
        return pd.DataFrame([data_dict])

    for officer, block in early_df.groupby("FieldOfficer"):
        if block.empty:
            continue

        header = {col: None for col in available_cols}
        header["FieldOfficer"] = officer
        blocks.append(create_row(header))
        blocks.append(block[available_cols])

        sub = {col: None for col in available_cols}
        sub["FieldOfficer"] = officer
        sub["FullNames"] = "SUBTOTAL"
        sub["AmountDue"] = block["AmountDue"].sum() if "AmountDue" in block.columns else 0
        sub["Arrears"] = block["Arrears"].sum() if "Arrears" in block.columns else 0

        if "AmountDue" in block.columns and block["AmountDue"].sum() > 0:
            sub["ArrearsPercentage"] = (block["Arrears"].sum() / block["AmountDue"].sum() * 100).round(2)
        else:
            sub["ArrearsPercentage"] = 0

        sub["SeverityIndicator"] = "SUBTOTAL"
        sub["DaysInArrears"] = block["DaysInArrears"].mean().round(1) if "DaysInArrears" in block.columns else 0
        sub["RiskCategory"] = "SUBTOTAL"
        sub["AmountDisbursed"] = block["AmountDisbursed"].sum() if "AmountDisbursed" in block.columns else 0
        blocks.append(create_row(sub))
        blocks.append(create_row({col: None for col in available_cols}))

    if not blocks:
        return pd.DataFrame(columns=available_cols)

    result = pd.concat(blocks, ignore_index=True)

    # Grand total
    grand_total = {col: None for col in available_cols}
    grand_total["FieldOfficer"] = "TOTAL"
    grand_total["FullNames"] = "GRAND TOTAL"
    grand_total["AmountDue"] = early_df["AmountDue"].sum() if "AmountDue" in early_df.columns else 0
    grand_total["Arrears"] = early_df["Arrears"].sum() if "Arrears" in early_df.columns else 0

    if "AmountDue" in early_df.columns and early_df["AmountDue"].sum() > 0:
        grand_total["ArrearsPercentage"] = (early_df["Arrears"].sum() / early_df["AmountDue"].sum() * 100).round(2)
    else:
        grand_total["ArrearsPercentage"] = 0

    grand_total["SeverityIndicator"] = "GRAND TOTAL"
    grand_total["DaysInArrears"] = early_df["DaysInArrears"].mean().round(1) if "DaysInArrears" in early_df.columns else 0
    grand_total["RiskCategory"] = "GRAND TOTAL"
    grand_total["AmountDisbursed"] = early_df["AmountDisbursed"].sum() if "AmountDisbursed" in early_df.columns else 0

    result = pd.concat([result, create_row(grand_total)], ignore_index=True)
    return result

# =========================================================
# 5. FUNCTION: Charts
# =========================================================

def make_top10_labels(customer_risk, wrap_width: int = 30):
    """Returns DataFrame with wrapped labels for Top 10."""
    if customer_risk.empty:
        return pd.DataFrame(columns=["DisplayName", "DisplayNameWrapped", "RiskScore"])

    top10 = customer_risk.sort_values("RiskScore", ascending=False).head(10).copy()
    top10['DisplayName'] = (
        top10['FullNames'].astype(str) + " (" + 
        top10['FieldOfficer'].fillna('Unassigned').astype(str) + ")"
    )
    top10['DisplayNameWrapped'] = top10['DisplayName'].apply(lambda s: fill(str(s), width=wrap_width))
    return top10[["DisplayName", "DisplayNameWrapped", "RiskScore"]]


def generate_charts(customer_risk, officer_arrears, arrears_by_inst, wb, config: dict = None):
    """Generate charts with error handling."""
    
    if config is None:
        config = {}
    
    charts_cfg = config.get('charts', {})
    wrap_width = charts_cfg.get('top10_wrap_width', 30)

    try:
        # 1. Arrears by Officer
        if not officer_arrears.empty and len(officer_arrears) > 0:
            plt.figure(figsize=(8, 5))
            sorted_df = officer_arrears.sort_values("Arrears", ascending=True)

            if len(sorted_df) > 0:
                colors = plt.cm.Oranges(np.linspace(0.4, 1, len(sorted_df)))
                bars = plt.bar(sorted_df["FieldOfficer"], sorted_df["Arrears"], color=colors)
                plt.title("Total Arrears by Field Officer", fontsize=12, fontweight='bold')
                plt.xticks(rotation=45, ha="right")

                for bar in bars:
                    height = bar.get_height()
                    plt.text(bar.get_x() + bar.get_width()/2, height, f"{int(height):,}",
                            ha='center', va='bottom', fontsize=8)
                plt.tight_layout()
                plt.savefig("chart_arrears.png", dpi=CHART_DPI)
                plt.close()

        # 2. Arrears by Installment
        if not arrears_by_inst.empty and len(arrears_by_inst) > 0:
            plt.figure(figsize=(8, 5))
            plt.bar(arrears_by_inst["InstallmentNo"], arrears_by_inst["Arrears"], color='#e74c3c')
            plt.title("Worst Arrears by Loan Age", fontsize=12, fontweight='bold')
            plt.xlabel("Installment Number")
            plt.ylabel("Arrears Amount")
            plt.tight_layout()
            plt.savefig("chart_age.png", dpi=CHART_DPI)
            plt.close()

        # 3. Top 10 High Risk
        if not customer_risk.empty and len(customer_risk) >= 10:
            top10 = make_top10_labels(customer_risk, wrap_width=wrap_width)
            if not top10.empty:
                plt.figure(figsize=(8, 5))
                colors = plt.cm.Reds(np.linspace(1, 0.5, len(top10)))
                plt.barh(top10['DisplayNameWrapped'], top10['RiskScore'], color=colors)
                plt.title("Top 10 Highest Risk Customers", fontsize=12, fontweight='bold')
                plt.gca().invert_yaxis()
                plt.tight_layout()
                plt.savefig("chart_top10.png", dpi=CHART_DPI)
                plt.close()

        # 4. Risk Distribution Pie Chart - Enhanced with Officer Breakdown
        if not customer_risk.empty and "RiskCategory" in customer_risk.columns:
            risk_counts = customer_risk["RiskCategory"].value_counts().reindex(
                ["High Risk", "Medium Risk", "Low Risk"]
            ).fillna(0)

            if risk_counts.sum() > 0:
                total_customers = len(customer_risk)
                colors = ['#e74c3c', '#f39c12', '#27ae60']
                
                # Check if we have high risk data to show detailed breakdown
                high_risk_data = customer_risk[customer_risk["RiskCategory"] == "High Risk"]
                
                if not high_risk_data.empty and len(high_risk_data) > 0:
                    # Enhanced chart with officer breakdown
                    total_high_risk_arrears = high_risk_data["Arrears"].sum()
                    total_high_risk_customers = len(high_risk_data)

                    # Group by FieldOfficer for high risk data
                    officer_high_risk_summary = high_risk_data.groupby("FieldOfficer").agg(
                        TotalArrears=("Arrears", "sum"),
                        CustomerCount=("PhoneNumber", "nunique")
                    ).sort_values("TotalArrears", ascending=False)

                    if not officer_high_risk_summary.empty:
                        fig, ax = plt.subplots(figsize=(12, 8))

                        # Create subplots: main pie on left, breakdown info on right
                        gs = fig.add_gridspec(1, 2, width_ratios=[1.2, 0.8])
                        ax_pie = fig.add_subplot(gs[0])
                        ax_info = fig.add_subplot(gs[1])

                        # Main pie chart - Overall Risk Categories
                        wedges, texts, autotexts = ax_pie.pie(
                            risk_counts,
                            labels=risk_counts.index,
                            autopct=lambda p: f'{p:.1f}%\n({int(p*total_customers/100):,} customers)',
                            colors=colors,
                            startangle=90,
                            wedgeprops={'linewidth': 2, 'edgecolor': 'white'},
                            textprops={'fontsize': 10, 'fontweight': 'bold'}
                        )

                        # Style pie chart text
                        for autotext in autotexts:
                            autotext.set_color('white')
                            autotext.set_fontsize(9)
                            autotext.set_fontweight('bold')

                        ax_pie.set_title("Portfolio Risk Composition",
                                        fontsize=14, fontweight='bold', pad=20)

                        # Create breakdown info box on the right
                        ax_info.axis('off')
                        info_text = "HIGH RISK BREAKDOWN\n" + "="*25 + "\n\n"
                        
                        for officer, row in officer_high_risk_summary.iterrows():
                            arrears_percent = (row["TotalArrears"] / total_high_risk_arrears * 100)
                            customer_percent = (row["CustomerCount"] / total_high_risk_customers * 100)
                            info_text += f"{officer}:\n"
                            info_text += f"  {row['TotalArrears']:,.0f} ({arrears_percent:.1f}%)\n"
                            info_text += f"  {row['CustomerCount']} customers ({customer_percent:.1f}%)\n\n"

                        # Add portfolio summary at bottom
                        info_text += "\nPORTFOLIO SUMMARY\n" + "="*25 + "\n"
                        info_text += f"Total Customers: {len(customer_risk)}\n"
                        info_text += f"Total Arrears: {customer_risk['Arrears'].sum():,.0f}\n"
                        info_text += f"Customers in Arrears: {(customer_risk['Arrears'] > 0).sum()}\n"
                        info_text += f"Arrears Rate: {((customer_risk['Arrears'] > 0).mean() * 100):.1f}%"

                        ax_info.text(0.05, 0.95, info_text, transform=ax_info.transAxes,
                                   fontsize=9, verticalalignment='top', fontfamily='monospace',
                                   bbox=dict(boxstyle='round,pad=0.5', facecolor='lightgray', alpha=0.8))

                        # Add legend for risk categories at bottom right
                        legend_text = "RISK BREAKDOWN\n" + "="*20 + "\n"
                        
                        for i, (risk, count) in enumerate(risk_counts.items()):
                            if count > 0:
                                risk_arrears = customer_risk[customer_risk["RiskCategory"] == risk]["Arrears"].sum()
                                risk_arrears_percent = (risk_arrears / customer_risk["Arrears"].sum() * 100)
                                legend_text += f"â— {risk}: {count} customers\n"
                                legend_text += f"  {risk_arrears:,.0f} arrears ({risk_arrears_percent:.1f}%)\n\n"

                        ax_info.text(0.05, 0.35, legend_text, transform=ax_info.transAxes,
                                   fontsize=8, verticalalignment='top',
                                   bbox=dict(boxstyle='round,pad=0.5', facecolor='white', alpha=0.9))

                        plt.tight_layout()
                        plt.savefig("chart_pie.png", dpi=CHART_DPI, bbox_inches='tight')
                        plt.close()
                else:
                    # Simple pie chart if no high risk data
                    fig, ax = plt.subplots(figsize=(10, 8))
                    wedges, texts, autotexts = ax.pie(
                        risk_counts,
                        labels=risk_counts.index,
                        autopct=lambda p: f'{p:.1f}%\n({int(p*total_customers/100):,} customers)',
                        colors=colors,
                        startangle=90,
                        wedgeprops={'linewidth': 2, 'edgecolor': 'white'},
                        textprops={'fontsize': 10, 'fontweight': 'bold'}
                    )

                    for autotext in autotexts:
                        autotext.set_color('white')
                        autotext.set_fontsize(9)
                        autotext.set_fontweight('bold')

                    ax.set_title("Portfolio Risk Composition",
                                fontsize=14, fontweight='bold', pad=20)

                    plt.tight_layout()
                    plt.savefig("chart_pie.png", dpi=CHART_DPI, bbox_inches='tight')
                    plt.close()

        # Insert charts into workbook
        chart_files = {
            "chart_arrears.png": ("Officer Summary", "E2"),
            "chart_age.png": ("Arrears by Installment", "E2"),
            "chart_top10.png": ("Customer Risk Ranking", "J2"),
            "chart_pie.png": ("Portfolio Summary", "F2"),
            "chart_high_risk_by_officer.png": ("High Risk Summary", "F2")
        }

        for chart_file, (sheet_name, cell) in chart_files.items():
            if os.path.exists(chart_file) and sheet_name in wb.sheetnames:
                try:
                    wb[sheet_name].add_image(Image(chart_file), cell)
                except Exception as e:
                    print(f"âš ï¸  Could not insert chart {chart_file}: {e}")

    except Exception as e:
        print(f"âš ï¸  Chart generation warning: {e}")

# =========================================================
# 6. FUNCTION: Excel Formatting
# =========================================================
def format_excel_workbook(filepath):
    """Applies formatting, currency, and conditional rules."""
    try:
        wb = load_workbook(filepath)
        currency_fmt = '#,##0.00'

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Auto-fit columns
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = (max_length + 2) * 1.1
                if adjusted_width > 50:
                    adjusted_width = 50
                ws.column_dimensions[column_letter].width = adjusted_width

            # Currency formatting
            for column in ws.columns:
                header_val = ws[f"{column[0].column_letter}1"].value
                if header_val and any(x in str(header_val) for x in ["Arrears", "Balance", "Amount", "Funded", "Principal"]):
                    for cell in column[1:]:
                        try:
                            cell.number_format = currency_fmt
                        except Exception:
                            pass

        # Conditional formatting for risk categories
        for sheet_name in ["Customer Risk Ranking", "Early Arrears by Officer"]:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                risk_col_idx = None
                
                for i, cell in enumerate(ws[1], 1):
                    if cell.value == "RiskCategory":
                        risk_col_idx = i
                        break

                if risk_col_idx:
                    col_char = ws.cell(1, risk_col_idx).column_letter
                    rng = f"{col_char}2:{col_char}{ws.max_row}"
                    
                    red_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
                    yellow_fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
                    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
                    
                    ws.conditional_formatting.add(rng, 
                        FormulaRule(formula=[f'${col_char}2="High Risk"'], fill=red_fill))
                    ws.conditional_formatting.add(rng, 
                        FormulaRule(formula=[f'${col_char}2="Medium Risk"'], fill=yellow_fill))
                    ws.conditional_formatting.add(rng, 
                        FormulaRule(formula=[f'${col_char}2="Low Risk"'], fill=green_fill))

        # Bold special rows
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                txt = " ".join([str(c.value).upper() for c in row[:2] if c.value])
                if any(x in txt for x in ["SUBTOTAL", "TOTAL", "PORTFOLIO"]):
                    for cell in row:
                        cell.font = Font(bold=True)
                        if "SUBTOTAL" in txt or "TOTAL" in txt:
                            cell.border = Border(
                                top=Side(style='thin'),
                                bottom=Side(style='thin')
                            )

        wb.save(filepath)
        print("âœ… Excel formatting applied")

    except Exception as e:
        print(f"âš ï¸  Warning: Excel formatting failed: {e}")

def open_output_location(path: str) -> bool:
    """Opens the output file location in file explorer."""
    p = Path(path)
    if not p.exists():
        print(f"âš ï¸  Output file does not exist: {path}")
        return False
    
    try:
        if sys.platform.startswith('win'):
            subprocess.run(['explorer', '/select,', str(p)], check=False)
        elif sys.platform == 'darwin':
            subprocess.run(['open', '-R', str(p)], check=False)
        else:
            subprocess.run(['xdg-open', str(p.parent)], check=False)
        return True
    except Exception as e:
        print(f"âš ï¸  Could not open output location: {e}")
        return False

# =========================================================
# 7. MAIN EXECUTION
# =========================================================
def main():
    parser = argparse.ArgumentParser(description="Loan Risk Analysis Report Generator")
    parser.add_argument("-i", "--input", help="Input Excel/CSV file")
    parser.add_argument("-c", "--config", help="Config JSON path", default="premium_config.json")
    parser.add_argument("-o", "--output", help="Output file path (optional)")
    parser.add_argument("-v", "--verbose", action="store_true", help="Enable debug logging")
    parser.add_argument("--no-charts", action="store_true", help="Skip chart generation")
    parser.add_argument("--safe-mode", action="store_true", help="Safe mode - skip charts and formatting")
    parser.add_argument("--open-output", action="store_true", help="Open output folder when done")
    args = parser.parse_args()

    # Setup logging
    log_level = logging.DEBUG if args.verbose else logging.INFO
    logging.basicConfig(level=log_level, format="%(asctime)s - %(levelname)s - %(message)s")

    # Load config
    config = load_config(args.config)

    # Select input file
    input_file = args.input if args.input else select_file()
    if not input_file:
        print("âŒ No file selected. Exiting.")
        return

    # Determine output file
    base_name = os.path.splitext(os.path.basename(input_file))[0]
    output_file = args.output if args.output else os.path.join(
        os.path.dirname(input_file), f"{base_name}_Risk_Report.xlsx"
    )

    print(f"\n--- Starting Analysis for {os.path.basename(input_file)} ---\n")
    print(f"ðŸ” Debug: Input file: {input_file}")
    print(f"ðŸ” Debug: Output file: {output_file}")

    try:
        # Load and process data
        print("ðŸ“Š Step 1: Loading and cleaning data...")
        print(f"ðŸ” Debug: Attempting to load: {input_file}")
        
        if not os.path.exists(input_file):
            print(f"âŒ ERROR: Input file does not exist: {input_file}")
            return
            
        file_size = os.path.getsize(input_file)
        print(f"ðŸ” Debug: File size: {file_size:,} bytes")
        
        df = load_and_clean_data(input_file)
        print(f"âœ… Loaded {len(df)} rows")
        print(f"ðŸ” Debug: Columns found: {list(df.columns)}")
        
        # Check for essential data
        if df.empty:
            print("âŒ ERROR: No data loaded from file")
            return
            
        print("ðŸ“Š Step 2: Calculating risk metrics...")
        customer_risk = calculate_risk_metrics(df, config)
        print(f"âœ… Risk calculated for {len(customer_risk)} customers")

        if customer_risk.empty:
            print("âŒ Error: Risk calculation failed. Please check your data.")
            print("ðŸ” Debug: Checking data quality...")
            print(f"   - Total rows: {len(df)}")
            print(f"   - Phone numbers: {df['PhoneNumber'].nunique() if 'PhoneNumber' in df.columns else 'N/A'}")
            print(f"   - Arrears > 0: {(df['Arrears'] > 0).sum() if 'Arrears' in df.columns else 'N/A'}")
            return

        # Build summaries
        print("ðŸ“Š Step 3: Building summaries...")
        officer_arrears = customer_risk.groupby("FieldOfficer")["Arrears"].sum().reset_index().sort_values("Arrears", ascending=False)
        print(f"ðŸ” Debug: Found {len(officer_arrears)} officers")

        summary_dict = {
            "Total Customers": customer_risk["PhoneNumber"].nunique(),
            "Customers in Arrears": int((customer_risk["Arrears"] > 0).sum()),
            "Portfolio Arrears (Sum)": float(customer_risk["Arrears"].sum()),
            "Average Arrears per Customer": round(float(customer_risk["Arrears"].mean()), 2),
            "Loans in Arrears %": round(((customer_risk["Arrears"] > 0).mean() * 100.0), 2),
            "High Risk Customers": int((customer_risk["RiskCategory"] == "High Risk").sum()),
            "Medium Risk Customers": int((customer_risk["RiskCategory"] == "Medium Risk").sum()),
            "Low Risk Customers": int((customer_risk["RiskCategory"] == "Low Risk").sum()),
        }
        portfolio_summary = pd.DataFrame([summary_dict])

        # Officer matrix
        officer_matrix = pd.pivot_table(
            customer_risk, index="FieldOfficer", columns="RiskCategory",
            values="Arrears", aggfunc="sum", fill_value=0
        )

        for c in ["High Risk", "Medium Risk", "Low Risk"]:
            if c not in officer_matrix.columns:
                officer_matrix[c] = 0

        officer_matrix["TotalArrears"] = officer_matrix.sum(axis=1)
        officer_matrix["HighRisk%"] = (
            officer_matrix["High Risk"] / officer_matrix["TotalArrears"].replace(0, 1) * 100
        ).fillna(0)
        officer_matrix = officer_matrix.sort_values("TotalArrears", ascending=False).reset_index()

        # Add grand total
        grand_total_vals = officer_matrix.select_dtypes(include=[np.number]).sum().to_dict()
        grand_total_row = {"FieldOfficer": "TOTAL", **grand_total_vals}
        if grand_total_row["TotalArrears"] > 0:
            grand_total_row["HighRisk%"] = (
                grand_total_row["High Risk"] / grand_total_row["TotalArrears"]
            ) * 100
        else:
            grand_total_row["HighRisk%"] = 0

        officer_matrix = pd.concat([officer_matrix, pd.DataFrame([grand_total_row])], ignore_index=True)

        # Arrears by installment
        arrears_by_inst = df.groupby("InstallmentNo")["Arrears"].max().reset_index().sort_values("InstallmentNo")

        # Build reports
        printable_risk = build_printable_risk_sheet(customer_risk)
        early_arrears = build_early_arrears_report(df, customer_risk)

        # Write to Excel
        print("â³ Generating Excel Report...")
        os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)

        try:
            with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
                printable_risk.to_excel(writer, sheet_name="Customer Risk Ranking", index=False)
                officer_arrears.to_excel(writer, sheet_name="Officer Summary", index=False)
                officer_matrix.to_excel(writer, sheet_name="Officer Risk Breakdown", index=False)
                portfolio_summary.to_excel(writer, sheet_name="Portfolio Summary", index=False)
                arrears_by_inst.to_excel(writer, sheet_name="Arrears by Installment", index=False)
                early_arrears.to_excel(writer, sheet_name="Early Arrears by Officer", index=False)

                # Top 10 Report
                top10_df = customer_risk.sort_values("RiskScore", ascending=False).head(10).copy()
                top10_df['DisplayName'] = (
                    top10_df['FullNames'].astype(str) + " (" +
                    top10_df['FieldOfficer'].fillna('Unassigned').astype(str) + ")"
                )
                top10_df = top10_df[[
                    "DisplayName", "FullNames", "FieldOfficer", "PhoneNumber",
                    "Arrears", "LoanBalance", "RiskScore", "RiskCategory"
                ]]
                top10_df.to_excel(writer, sheet_name="Top 10 High Risk", index=False)

            print(f"âœ… Excel file created: {output_file}")

        except PermissionError:
            print(f"âŒ CRITICAL ERROR: Could not save to '{output_file}'.")
            print("   >> Please CLOSE the Excel file if it's open and try again.")
            return
        except Exception as e:
            print(f"âŒ Error writing Excel file: {e}")
            return

        # Generate charts and format
        if not args.safe_mode and not args.no_charts:
            try:
                print("ðŸ“Š Step 4: Generating charts...")
                wb = load_workbook(output_file)
                generate_charts(customer_risk, officer_arrears, arrears_by_inst, wb, config)
                wb.save(output_file)
                print("âœ… Charts generated and added to report")
            except Exception as e:
                print(f"âš ï¸  Warning: Chart generation failed: {e}")

        # Format workbook
        if not args.safe_mode:
            try:
                print("ðŸ“Š Step 5: Applying Excel formatting...")
                format_excel_workbook(output_file)
            except Exception as e:
                print(f"âš ï¸  Warning: Excel formatting failed: {e}")
        else:
            print("â­ï¸  Skipping formatting (safe mode)")

        # Cleanup temporary images
        for f in ["chart_arrears.png", "chart_age.png", "chart_top10.png", "chart_pie.png", "chart_high_risk_by_officer.png"]:
            if os.path.exists(f):
                try:
                    os.remove(f)
                    print(f"ðŸ§¹ Cleaned up: {f}")
                except Exception as e:
                    print(f"âš ï¸  Could not remove {f}: {e}")

        print(f"\nâœ… Success! Analysis saved to: {output_file}")
        print("\nðŸ“Š REPORT SUMMARY:")
        print(f"   Total Customers: {summary_dict['Total Customers']:,}")
        print(f"   Total Arrears: KES {summary_dict['Portfolio Arrears (Sum)']:,.2f}")
        print(f"   High Risk Customers: {summary_dict['High Risk Customers']:,}")
        print(f"   Arrears Rate: {summary_dict['Loans in Arrears %']:.1f}%")

        # Always open output location
        print("\nðŸ“‚ Opening output location...")
        if open_output_location(output_file):
            print(f"âœ… Output folder opened successfully")
        else:
            print("âš ï¸  Could not open the output folder automatically.")
            print(f"   Please navigate to: {os.path.dirname(os.path.abspath(output_file))}")

        # Also open if explicitly requested (redundant but maintains compatibility)
        if args.open_output:
            if open_output_location(output_file):
                print(f"ðŸ“‚ Opened output location")
            else:
                print("âš ï¸  Could not open the output folder. See the path above.")

    except FileNotFoundError as e:
        print(f"âŒ File Error: {e}")
        sys.exit(1)
    except ValueError as e:
        print(f"âŒ Data Error: {e}")
        sys.exit(1)
    except KeyError as e:
        print(f"âŒ Column Error: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"âŒ Unexpected error: {e}")
        if args.verbose:
            import traceback
            traceback.print_exc()
        sys.exit(1)

def _build_preview_chart(officer_arrears: pd.DataFrame, output_dir: str) -> str:
    """Create a stable chart image path for the Android preview card."""
    try:
        if officer_arrears is None or officer_arrears.empty:
            return ""

        chart_df = (
            officer_arrears.sort_values("Arrears", ascending=False)
            .head(10)
            .sort_values("Arrears", ascending=True)
        )
        if chart_df.empty:
            return ""

        chart_path = os.path.join(output_dir, "mtd_chart.png")
        plt.figure(figsize=(8, 5))
        plt.barh(chart_df["FieldOfficer"].astype(str), chart_df["Arrears"], color="#1A237E")
        plt.title("Top Field Officers by Arrears", fontsize=12, fontweight="bold")
        plt.xlabel("Arrears")
        plt.tight_layout()
        plt.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close()
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def run(input_path: str, output_dir: str) -> dict:
    """
    App entrypoint used by Chaquopy.
    Preserves the original MTD_unpaid_dues.py processing logic and output structure.
    """
    temp_charts = [
        "chart_arrears.png",
        "chart_age.png",
        "chart_top10.png",
        "chart_pie.png",
        "chart_high_risk_by_officer.png",
    ]

    os.makedirs(output_dir, exist_ok=True)
    if not input_path:
        raise ValueError("No input file provided.")
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"The file '{input_path}' was not found.")

    original_print = builtins.print
    builtins.print = lambda *args, **kwargs: None

    try:
        config = load_config()
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        output_file = os.path.join(output_dir, f"{base_name}_Risk_Report.xlsx")

        df = load_and_clean_data(input_path)
        if df.empty:
            raise ValueError("No data loaded from file.")

        customer_risk = calculate_risk_metrics(df, config)
        if customer_risk.empty:
            raise ValueError("Risk calculation returned no rows.")

        officer_arrears = (
            customer_risk.groupby("FieldOfficer")["Arrears"]
            .sum()
            .reset_index()
            .sort_values("Arrears", ascending=False)
        )

        summary_dict = {
            "Total Customers": int(customer_risk["PhoneNumber"].nunique()),
            "Customers in Arrears": int((customer_risk["Arrears"] > 0).sum()),
            "Portfolio Arrears (Sum)": float(customer_risk["Arrears"].sum()),
            "Average Arrears per Customer": round(float(customer_risk["Arrears"].mean()), 2),
            "Loans in Arrears %": round(((customer_risk["Arrears"] > 0).mean() * 100.0), 2),
            "High Risk Customers": int((customer_risk["RiskCategory"] == "High Risk").sum()),
            "Medium Risk Customers": int((customer_risk["RiskCategory"] == "Medium Risk").sum()),
            "Low Risk Customers": int((customer_risk["RiskCategory"] == "Low Risk").sum()),
        }
        portfolio_summary = pd.DataFrame([summary_dict])

        officer_matrix = pd.pivot_table(
            customer_risk,
            index="FieldOfficer",
            columns="RiskCategory",
            values="Arrears",
            aggfunc="sum",
            fill_value=0,
        )
        for c in ["High Risk", "Medium Risk", "Low Risk"]:
            if c not in officer_matrix.columns:
                officer_matrix[c] = 0

        officer_matrix["TotalArrears"] = officer_matrix.sum(axis=1)
        officer_matrix["HighRisk%"] = (
            officer_matrix["High Risk"] / officer_matrix["TotalArrears"].replace(0, 1) * 100
        ).fillna(0)
        officer_matrix = officer_matrix.sort_values("TotalArrears", ascending=False).reset_index()

        grand_total_vals = officer_matrix.select_dtypes(include=[np.number]).sum().to_dict()
        grand_total_row = {"FieldOfficer": "TOTAL", **grand_total_vals}
        if grand_total_row["TotalArrears"] > 0:
            grand_total_row["HighRisk%"] = (
                grand_total_row["High Risk"] / grand_total_row["TotalArrears"]
            ) * 100
        else:
            grand_total_row["HighRisk%"] = 0
        officer_matrix = pd.concat([officer_matrix, pd.DataFrame([grand_total_row])], ignore_index=True)

        arrears_by_inst = (
            df.groupby("InstallmentNo")["Arrears"]
            .max()
            .reset_index()
            .sort_values("InstallmentNo")
        )

        printable_risk = build_printable_risk_sheet(customer_risk)
        early_arrears = build_early_arrears_report(df, customer_risk)

        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            printable_risk.to_excel(writer, sheet_name="Customer Risk Ranking", index=False)
            officer_arrears.to_excel(writer, sheet_name="Officer Summary", index=False)
            officer_matrix.to_excel(writer, sheet_name="Officer Risk Breakdown", index=False)
            portfolio_summary.to_excel(writer, sheet_name="Portfolio Summary", index=False)
            arrears_by_inst.to_excel(writer, sheet_name="Arrears by Installment", index=False)
            early_arrears.to_excel(writer, sheet_name="Early Arrears by Officer", index=False)

            top10_df = customer_risk.sort_values("RiskScore", ascending=False).head(10).copy()
            top10_df["DisplayName"] = (
                top10_df["FullNames"].astype(str)
                + " ("
                + top10_df["FieldOfficer"].fillna("Unassigned").astype(str)
                + ")"
            )
            top10_df = top10_df[
                [
                    "DisplayName",
                    "FullNames",
                    "FieldOfficer",
                    "PhoneNumber",
                    "Arrears",
                    "LoanBalance",
                    "RiskScore",
                    "RiskCategory",
                ]
            ]
            top10_df.to_excel(writer, sheet_name="Top 10 High Risk", index=False)

        wb = load_workbook(output_file)
        generate_charts(customer_risk, officer_arrears, arrears_by_inst, wb, config)
        wb.save(output_file)
        format_excel_workbook(output_file)

        chart_file = _build_preview_chart(officer_arrears, output_dir)

        total_due = (
            float(pd.to_numeric(df["AmountDue"], errors="coerce").fillna(0).sum())
            if "AmountDue" in df.columns
            else 0.0
        )
        total_paid = (
            float(pd.to_numeric(df["AmountPaid"], errors="coerce").fillna(0).sum())
            if "AmountPaid" in df.columns
            else 0.0
        )
        total_unpaid = (
            float(pd.to_numeric(df["Arrears"], errors="coerce").fillna(0).sum())
            if "Arrears" in df.columns
            else max(total_due - total_paid, 0.0)
        )
        if total_due > 0:
            if total_paid > 0:
                collection_rate_pct = (total_paid / total_due) * 100.0
            else:
                collection_rate_pct = max(0.0, min(100.0, (1.0 - (total_unpaid / total_due)) * 100.0))
        else:
            collection_rate_pct = 0.0

        return {
            "success": True,
            "output_file": output_file,
            "chart_file": chart_file,
            "stats": {
                "total_due": round(total_due, 2),
                "total_paid": round(total_paid, 2),
                "total_unpaid": round(total_unpaid, 2),
                "collection_rate_pct": round(collection_rate_pct, 1),
                "total_customers": summary_dict["Total Customers"],
                "customers_in_arrears": summary_dict["Customers in Arrears"],
                "portfolio_arrears_sum": round(summary_dict["Portfolio Arrears (Sum)"], 2),
                "average_arrears_per_customer": summary_dict["Average Arrears per Customer"],
                "loans_in_arrears_pct": summary_dict["Loans in Arrears %"],
                "high_risk_customers": summary_dict["High Risk Customers"],
                "medium_risk_customers": summary_dict["Medium Risk Customers"],
                "low_risk_customers": summary_dict["Low Risk Customers"],
            },
        }
    finally:
        builtins.print = original_print
        for chart_file in temp_charts:
            try:
                if os.path.exists(chart_file):
                    os.remove(chart_file)
            except Exception:
                pass

if __name__ == "__main__":
    main()
