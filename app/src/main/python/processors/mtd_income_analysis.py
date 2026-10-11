"""
MTD Income Analysis Processor
Produces a clean 4-column report:
  Full Name | Total Amount | TargetD{N} | D/S

  Total Amount  = actual MTD income collected (from input file)
  TargetD{N}    = daily_target × day_number  (entered by user)
  D/S           = Total Amount - TargetD{N}  (negative = behind, shown in red)
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List, Tuple

import matplotlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from processors.tabular_io import read_tabular_file


# ── Helpers ───────────────────────────────────────────────────────────────────

def _read_input(path: str) -> pd.DataFrame:
    return read_tabular_file(path, dtype=str)


def _normalize_header(name: str) -> str:
    clean = str(name or "").strip().lower()
    clean = clean.replace("%", "pct")
    clean = re.sub(r"[^a-z0-9]+", "", clean)
    return clean


def _pick_column(columns: List[str], candidates: List[str]) -> str:
    lookup = {_normalize_header(c): c for c in columns}
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in lookup:
            return lookup[key]
    # partial match fallback
    for candidate in candidates:
        key = _normalize_header(candidate)
        for col in columns:
            col_key = _normalize_header(col)
            if key in col_key or col_key in key:
                return col
    return ""


def _to_number(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(r"[^\d\.\-]", "", regex=True)
        .replace("", "0")
        .pipe(pd.to_numeric, errors="coerce")
        .fillna(0.0)
    )


def _parse_day(day_input: str) -> int:
    text = str(day_input or "").strip()
    digits = re.sub(r"[^0-9]", "", text)
    if not digits:
        return 1
    value = int(digits)
    return max(1, min(31, value))


def _parse_daily_targets(daily_targets_input: str, names: List[str]) -> Dict[str, float]:
    """
    Parse user-supplied targets from text like:
        JOASH MUTISYA=3500
        DEFAULT=5000
    Returns a dict of {UPPER_NAME: daily_amount}.
    """
    targets: Dict[str, float] = {n.strip().upper(): 0.0 for n in names}
    default_target = 0.0

    if not daily_targets_input or not str(daily_targets_input).strip():
        return targets

    lookup = {n.strip().upper(): n.strip().upper() for n in names}

    for raw_line in str(daily_targets_input).splitlines():
        line = raw_line.strip()
        if not line or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip().upper()
        try:
            amount = max(0.0, float(str(value).strip().replace(",", "")))
        except Exception:
            continue

        if key in {"DEFAULT", "ALL", "*"}:
            default_target = amount
            continue
        if key in lookup:
            targets[lookup[key]] = amount

    if default_target > 0:
        for name_key in targets:
            if targets[name_key] <= 0:
                targets[name_key] = default_target

    return targets


# ── Public: list portfolios ────────────────────────────────────────────────────

def get_portfolios(input_path: str) -> List[str]:
    """Return sorted list of officer/portfolio names from the input file."""
    try:
        df = _read_input(input_path)
        if df is None or df.empty:
            return []
        df.columns = [str(c).strip() for c in df.columns]
        name_col = _pick_column(
            list(df.columns),
            ["Full Name", "FullName", "Name", "SalesRep", "Officer", "Field Officer"],
        )
        if not name_col:
            return []
        names = (
            df[name_col]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.replace(r"\s+", " ", regex=True)
            .str.upper()
        )
        return sorted(names[names != ""].drop_duplicates().tolist())
    except Exception:
        return []


# ── Core: build the analysis dataframe ───────────────────────────────────────

def _build_report(
    df: pd.DataFrame,
    day_input: str,
    daily_targets_input: str,
) -> Tuple[pd.DataFrame, str, dict]:
    """
    Returns (report_df, target_col_name, stats_dict).

    report_df columns:
        Full Name | Total Amount | TargetD{N} | D/S
    plus a grand-total row at the bottom (Full Name = "").
    """
    if df is None or df.empty:
        raise ValueError("Input file is empty.")

    day_number = _parse_day(day_input)
    target_col_name = f"TargetD{day_number}"

    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    cols = list(df.columns)

    name_col = _pick_column(cols, ["Full Name", "FullName", "Name", "SalesRep", "Officer", "Field Officer"])
    income_col = _pick_column(cols, ["Total Amount", "TotalAmount", "Amount", "Paid Amount", "PaidAmount", "Income"])

    if not name_col:
        raise ValueError("Cannot find a Name column. Expected 'Full Name', 'Name', or 'SalesRep'.")
    if not income_col:
        raise ValueError("Cannot find an income column. Expected 'Total Amount' or 'Paid Amount'.")

    names_series = (
        df[name_col]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.replace(r"\s+", " ", regex=True)
        .str.upper()
    )
    mask = names_series != ""
    if not mask.any():
        raise ValueError("No valid officer names found in the file.")

    base = df.loc[mask].copy()
    names_list: List[str] = names_series.loc[mask].tolist()
    income_list = _to_number(base[income_col]).round(2).tolist()

    # Resolve daily targets (entered by user)
    daily_targets = _parse_daily_targets(daily_targets_input, names_list)

    target_d_list = [daily_targets.get(n.strip().upper(), 0.0) * day_number for n in names_list]
    ds_list = [inc - tgt for inc, tgt in zip(income_list, target_d_list)]

    rows = []
    for name, income, target_d, ds in zip(names_list, income_list, target_d_list, ds_list):
        rows.append({
            "Full Name": name,
            "Total Amount": income,
            target_col_name: target_d,
            "D/S": ds,
        })

    report_df = pd.DataFrame(rows)

    # Sort alphabetically by Full Name to match expected report order
    report_df = report_df.sort_values("Full Name", ascending=True).reset_index(drop=True)

    # Grand total row (blank name, as shown in screenshot)
    total_row = pd.DataFrame([{
        "Full Name": "",
        "Total Amount": report_df["Total Amount"].sum(),
        target_col_name: report_df[target_col_name].sum(),
        "D/S": report_df["D/S"].sum(),
    }])
    final_df = pd.concat([report_df, total_row], ignore_index=True)

    # Stats for the Android UI card
    total_actual = float(report_df["Total Amount"].sum())
    total_target = float(report_df[target_col_name].sum())
    total_ds = float(report_df["D/S"].sum())
    behind_count = int((report_df["D/S"] < 0).sum())
    on_track_count = int((report_df["D/S"] >= 0).sum())

    stats = {
        "day_number": day_number,
        "total_portfolios": len(report_df),
        "total_actual_income": round(total_actual, 2),
        "total_target_income": round(total_target, 2),
        "total_income_gap": round(total_ds, 2),
        "on_track_count": on_track_count,
        "behind_count": behind_count,
    }

    return final_df, target_col_name, stats


# ── Chart ─────────────────────────────────────────────────────────────────────

def _build_chart(df: pd.DataFrame, target_col: str, chart_path: str) -> str:
    try:
        body = df[df["Full Name"] != ""].copy()
        if body.empty:
            return ""

        body = body.sort_values("D/S", ascending=True)
        names = body["Full Name"].tolist()
        actuals = body["Total Amount"].tolist()
        targets = body[target_col].tolist()

        x = list(range(len(names)))
        fig, ax = plt.subplots(figsize=(max(8, len(names) * 1.2), 5))
        ax.bar([i - 0.2 for i in x], actuals, width=0.4, label="Actual Income", color="#1E88E5")
        ax.bar([i + 0.2 for i in x], targets, width=0.4, label=target_col, color="#FB8C00")
        ax.set_xticks(x)
        ax.set_xticklabels(names, rotation=45, ha="right", fontsize=8)
        ax.set_ylabel("Amount (KES)")
        ax.set_title(f"MTD Income Analysis — Actual vs {target_col}", fontweight="bold")
        ax.legend()
        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


# ── Excel writer ──────────────────────────────────────────────────────────────

def _write_excel(df: pd.DataFrame, target_col: str, output_file: str, chart_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Analysis"

    columns = ["Full Name", "Total Amount", target_col, "D/S"]

    thin = Side(style="thin", color="D0D0D0")
    thick = Side(style="medium", color="888888")
    border_normal = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_total = Border(left=thin, right=thin, top=thick, bottom=thin)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Header
    header_fill = PatternFill("solid", fgColor="0D47A1")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col_idx, header in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center if col_idx > 1 else align_left
        cell.border = border_normal

    # Data rows
    red_font = Font(color="C62828", bold=True)
    black_bold = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="F0F0F0")

    num_rows = len(df)
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        is_total_row = str(row.get("Full Name", "")) == ""
        this_border = border_total if is_total_row else border_normal

        for col_idx, col_name in enumerate(columns, start=1):
            value = row.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = this_border

            if is_total_row:
                cell.fill = total_fill
            else:
                cell.fill = PatternFill("solid", fgColor="FFFFFF")

            # Alignment
            if col_name == "Full Name":
                cell.alignment = align_left
            else:
                cell.alignment = align_right

            # Number format for amounts
            if col_name in ("Total Amount", target_col):
                cell.number_format = "#,##0"
            elif col_name == "D/S":
                cell.number_format = "#,##0;[Red]-#,##0"

            # Font: D/S red when negative, total row bold, normal otherwise
            ds_value = row.get("D/S", 0)
            try:
                ds_float = float(ds_value)
            except Exception:
                ds_float = 0.0

            if col_name == "D/S" and ds_float < 0:
                cell.font = red_font
            elif is_total_row:
                if col_name == "D/S" and ds_float < 0:
                    cell.font = red_font
                else:
                    cell.font = black_bold
            else:
                cell.font = Font()

    # Column widths
    col_widths = {"Full Name": 22, "Total Amount": 16, target_col: 16, "D/S": 14}
    for col_idx, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 14)

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # Chart sheet
    if chart_path and os.path.exists(chart_path):
        from openpyxl.drawing.image import Image as XLImage
        chart_ws = wb.create_sheet("Chart")
        chart_ws.add_image(XLImage(chart_path), "A1")

    wb.save(output_file)


# ── Public entry point ────────────────────────────────────────────────────────

def run(input_path: str, output_dir: str, day_input: str = "1", daily_targets_input: str = "") -> dict:
    os.makedirs(output_dir, exist_ok=True)

    source_df = _read_input(input_path)
    final_df, target_col_name, stats = _build_report(
        source_df,
        day_input=day_input,
        daily_targets_input=daily_targets_input,
    )

    chart_path = os.path.join(output_dir, "mtd_income_analysis_chart.png")
    chart_file = _build_chart(final_df, target_col_name, chart_path)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"MTD_Income_Analysis_{timestamp}.xlsx")
    _write_excel(final_df, target_col_name, output_file, chart_file)

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": stats,
    }
