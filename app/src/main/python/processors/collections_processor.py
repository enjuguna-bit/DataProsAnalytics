"""
Collections Processor
Script-parity implementation of Arreas_collected.py for Android/Chaquopy.
"""

import os
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt


VALID_BUCKETS = ["1-15", "16-30", "31-180"]
REQUIRED_COLUMNS_SOD = ["LoanId", "SalesRep", "CustomerName", "ArrearsAmount", "DaysInArrears"]
REQUIRED_COLUMNS_CUR = ["LoanId", "ArrearsAmount"]


def _get_bucket(days: int) -> str:
    """Categorize arrears into age buckets."""
    try:
        d = int(float(days))
    except Exception:
        return "Other"
    if 1 <= d <= 15:
        return "1-15"
    if 16 <= d <= 30:
        return "16-30"
    if 31 <= d <= 180:
        return "31-180"
    return "Other"


def _normalize_officer_names(df: pd.DataFrame) -> pd.DataFrame:
    """Clean and normalize officer names."""
    if "SalesRep" not in df.columns:
        return df

    df["SalesRep"] = df["SalesRep"].astype(str).str.strip()
    name_corrections = {
        "Brian Wanj:": "Brian Wanjau",
        "Brian Wanjau:": "Brian Wanjau",
        "Brian W.": "Brian Wanjau",
    }
    for wrong_name, correct_name in name_corrections.items():
        df["SalesRep"] = df["SalesRep"].str.replace(wrong_name, correct_name, regex=False)

    df["SalesRep"] = df["SalesRep"].str.title()
    return df


def _read_input(file_path: str) -> pd.DataFrame:
    """Load and clean data from CSV/Excel."""
    try:
        if file_path.lower().endswith(".csv"):
            encodings = ["utf-8", "latin1", "cp1252", "iso-8859-1"]
            for encoding in encodings:
                try:
                    df = pd.read_csv(file_path, encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        df.columns = df.columns.str.strip().str.title().str.replace(" ", "", regex=False)

        column_mapping = {
            "Loanid": "LoanId",
            "LoanId": "LoanId",
            "Loan_Id": "LoanId",
            "Salesrep": "SalesRep",
            "SalesRep": "SalesRep",
            "Sales_Rep": "SalesRep",
            "Officer": "SalesRep",
            "Agent": "SalesRep",
            "FieldOfficer": "SalesRep",
            "Field_Officer": "SalesRep",
            "Fieldofficer": "SalesRep",
            "Arrearsamount": "ArrearsAmount",
            "ArrearsAmount": "ArrearsAmount",
            "Arrears_Amount": "ArrearsAmount",
            "Arrears": "ArrearsAmount",
            "Amount": "ArrearsAmount",
            "Balance": "ArrearsAmount",
            "Daysinarrears": "DaysInArrears",
            "DaysInArrears": "DaysInArrears",
            "Days_In_Arrears": "DaysInArrears",
            "Days": "DaysInArrears",
            "Age": "DaysInArrears",
            "Daysinarr": "DaysInArrears",
            "Customername": "CustomerName",
            "CustomerName": "CustomerName",
            "Clientname": "CustomerName",
            "ClientName": "CustomerName",
            "Name": "CustomerName",
            "Borrowername": "CustomerName",
            "BorrowerName": "CustomerName",
            "Fullnames": "CustomerName",
            "FullNames": "CustomerName",
            "Fullname": "CustomerName",
            "FullName": "CustomerName",
            "Client": "CustomerName",
        }
        df.rename(columns=column_mapping, inplace=True)
        return df
    except Exception as exc:
        raise Exception(f"Error loading {os.path.basename(file_path)}: {exc}") from exc


def _validate_dataframes(df_sod: pd.DataFrame, df_cur: pd.DataFrame) -> Tuple[bool, str]:
    """Validate required columns exist in both files."""
    missing_sod = [col for col in REQUIRED_COLUMNS_SOD if col not in df_sod.columns]
    missing_cur = [col for col in REQUIRED_COLUMNS_CUR if col not in df_cur.columns]

    if missing_sod:
        readable = {
            "LoanId": "LoanId or Loan ID",
            "SalesRep": "SalesRep or Sales Rep",
            "CustomerName": "Customer Name, Client Name, or Name",
            "ArrearsAmount": "Arrears Amount",
            "DaysInArrears": "DaysInArrears or Days In Arrears",
        }
        missing_readable = [readable.get(col, col) for col in missing_sod]
        return False, f"Start-of-Day file missing columns: {', '.join(missing_readable)}"

    if missing_cur:
        readable = {
            "LoanId": "LoanId or Loan ID",
            "ArrearsAmount": "Arrears Amount",
        }
        missing_readable = [readable.get(col, col) for col in missing_cur]
        return False, f"Current file missing columns: {', '.join(missing_readable)}"

    return True, "Validation successful"


def _parse_targets_from_text(targets_input: str, officers: List[str]) -> Dict[str, float]:
    """
    Parse manual targets from text input.
    Expected format per line: OfficerName=TargetAmount
    """
    if not targets_input or not targets_input.strip():
        return {}

    lookup = {str(o).strip().lower(): str(o) for o in officers}
    parsed: Dict[str, float] = {}

    for raw_line in targets_input.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if "=" in line:
            officer_name, amount_text = line.split("=", 1)
            officer_key = officer_name.strip().lower()
            if officer_key not in lookup:
                continue

            amount_text = amount_text.strip()
            if amount_text == "":
                amount = 0.0
            else:
                try:
                    amount = float(amount_text.replace(",", ""))
                except Exception:
                    amount = 0.0

            parsed[lookup[officer_key]] = round(max(0.0, amount), 2)
        else:
            officer_key = line.lower()
            if officer_key in lookup:
                parsed[lookup[officer_key]] = 0.0

    for officer in officers:
        if officer not in parsed:
            parsed[officer] = 0.0

    return parsed


def _collect_officer_targets(
    officers: List[str], df_sod: pd.DataFrame, targets_input: str = ""
) -> Dict[str, float]:
    """
    Android-safe target collection.
    If a target column exists, use it; otherwise default every officer to 0.
    """
    if not officers:
        return {}

    manual_targets = _parse_targets_from_text(targets_input, officers)
    if manual_targets:
        return manual_targets

    target_candidates = ["Target", "CollectionTarget", "OfficerTarget", "DailyTarget"]
    target_col = next((c for c in target_candidates if c in df_sod.columns), None)

    targets: Dict[str, float] = {}
    if target_col:
        target_df = df_sod[["SalesRep", target_col]].copy()
        target_df[target_col] = pd.to_numeric(target_df[target_col], errors="coerce").fillna(0.0)
        mapped = target_df.groupby("SalesRep")[target_col].max().to_dict()
        for officer in officers:
            targets[officer] = round(float(mapped.get(officer, 0.0)), 2)
    else:
        for officer in officers:
            targets[officer] = 0.0

    return targets


def _process_data(
    sod_path: str, cur_path: str, targets_input: str = ""
) -> Optional[Tuple[pd.DataFrame, pd.DataFrame, List[str], Dict[str, float]]]:
    """Process arrears data and return collections, merged frame, officers, and targets."""
    try:
        df_sod = _read_input(sod_path)
        df_cur = _read_input(cur_path)
        df_sod = _normalize_officer_names(df_sod)

        ok, message = _validate_dataframes(df_sod, df_cur)
        if not ok:
            raise ValueError(message)

        df_sod = df_sod.rename(columns={"ArrearsAmount": "Arrears_SOD", "DaysInArrears": "Age_SOD"})
        df_cur = df_cur.rename(columns={"ArrearsAmount": "Arrears_CUR"})

        df_sod["Arrears_SOD"] = pd.to_numeric(df_sod["Arrears_SOD"], errors="coerce")
        df_sod["Age_SOD"] = pd.to_numeric(df_sod["Age_SOD"], errors="coerce")
        df_cur["Arrears_CUR"] = pd.to_numeric(df_cur["Arrears_CUR"], errors="coerce")

        df_sod = df_sod.dropna(subset=["LoanId", "SalesRep", "Arrears_SOD"])
        df_cur = df_cur.dropna(subset=["LoanId", "Arrears_CUR"])

        unique_officers = sorted(df_sod["SalesRep"].astype(str).unique().tolist())
        officer_targets = _collect_officer_targets(unique_officers, df_sod, targets_input)

        df_merged = pd.merge(
            df_sod[["LoanId", "SalesRep", "CustomerName", "Arrears_SOD", "Age_SOD"]],
            df_cur[["LoanId", "Arrears_CUR"]],
            on="LoanId",
            how="left",
            validate="one_to_one",
        )

        # Missing in current file = paid off loan.
        df_merged["Arrears_CUR"] = df_merged["Arrears_CUR"].fillna(0)
        df_merged["Collected"] = df_merged["Arrears_SOD"] - df_merged["Arrears_CUR"]

        df_collected = df_merged[df_merged["Collected"] > 0].copy()
        df_collected["Bucket"] = df_collected["Age_SOD"].apply(_get_bucket)
        df_collected = df_collected[df_collected["Bucket"].isin(VALID_BUCKETS)]
        df_collected = df_collected.sort_values(by=["SalesRep", "Bucket"])

        return df_collected, df_merged, unique_officers, officer_targets
    except Exception as exc:
        raise Exception(f"Error processing data: {exc}") from exc


def _create_formatted_table(
    df_collected: pd.DataFrame,
    df_merged: pd.DataFrame,
    officers: List[str],
    officer_targets: Dict[str, float],
) -> pd.DataFrame:
    """Create pivot-style summary table with targets and achievement fields."""
    pivot = pd.pivot_table(
        df_collected,
        values="Collected",
        index="SalesRep",
        columns="Bucket",
        aggfunc="sum",
        fill_value=0,
        margins=False,
    )

    for officer in officers:
        if officer not in pivot.index:
            new_row = pd.Series({bucket: 0.0 for bucket in VALID_BUCKETS}, name=officer)
            pivot = pd.concat([pivot, pd.DataFrame([new_row])])

    for bucket in VALID_BUCKETS:
        if bucket not in pivot.columns:
            pivot[bucket] = 0.0

    pivot = pivot[VALID_BUCKETS]
    pivot["Grand Total"] = pivot.sum(axis=1)
    pivot = pivot.sort_index(ascending=True)

    if officer_targets:
        for officer in pivot.index:
            if officer not in officer_targets:
                officer_targets[officer] = 0.0

        pivot["Target"] = pivot.index.map(officer_targets)
        pivot["Remaining"] = pivot["Target"] - pivot["Grand Total"]
        pivot["Achievement"] = np.where(
            pivot["Target"] > 0, (pivot["Grand Total"] / pivot["Target"] * 100), 0
        )

    grand_row = pd.DataFrame(pivot.sum(axis=0)).T
    grand_row.index = ["GRAND TOTAL"]

    if "Target" in pivot.columns:
        grand_row["Target"] = pivot["Target"].sum()
        grand_row["Remaining"] = grand_row["Target"] - grand_row["Grand Total"]
        if grand_row["Target"].iloc[0] > 0:
            grand_row["Achievement"] = grand_row["Grand Total"] / grand_row["Target"] * 100
        else:
            grand_row["Achievement"] = 0

    final_df = pd.concat([pivot, grand_row])

    numeric_cols = final_df.select_dtypes(include=[np.number]).columns
    final_df[numeric_cols] = final_df[numeric_cols].round(2)

    for col in numeric_cols:
        if col == "Achievement":
            final_df[col] = final_df[col].apply(lambda x: f"{x:,.2f}%" if pd.notnull(x) else "")
        else:
            final_df[col] = final_df[col].apply(
                lambda x: f"{x:,.2f}" if pd.notnull(x) and isinstance(x, (int, float, np.floating)) else ""
            )

    return final_df


def _create_excel_report(
    df_collected: pd.DataFrame,
    df_merged: pd.DataFrame,
    final_df: pd.DataFrame,
    output_path: str,
    officer_targets: Dict[str, float],
):
    """Create report workbook with summary, detailed, raw, and statistics sheets."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_sheet = "Summary Report"
        summary_start_row = 4
        final_df.to_excel(writer, sheet_name=summary_sheet, startrow=summary_start_row)

        start_row = summary_start_row + len(final_df) + 3
        detailed_cols = ["SalesRep", "CustomerName", "Arrears_SOD", "Age_SOD", "Arrears_CUR", "Collected", "Bucket"]
        df_detailed_out = df_collected[detailed_cols].copy()

        total_counts = df_merged.groupby("SalesRep")["LoanId"].nunique()
        collected_counts = df_collected.groupby("SalesRep")["LoanId"].nunique()

        detailed_rows = []
        for rep_name, group in df_collected.groupby("SalesRep"):
            total_sod = int(total_counts.get(rep_name, 0))
            total_coll = int(collected_counts.get(rep_name, 0))
            summary_row = {
                "SalesRep": rep_name,
                "CustomerName": f"({total_coll:,} accounts collected out of {total_sod:,})",
                "Arrears_SOD": None,
                "Age_SOD": None,
                "Arrears_CUR": None,
                "Collected": None,
                "Bucket": None,
            }
            detailed_rows.append(summary_row)
            for _, row in group.iterrows():
                cust_row = row[["CustomerName", "Arrears_SOD", "Age_SOD", "Arrears_CUR", "Collected", "Bucket"]].to_dict()
                cust_row["SalesRep"] = ""
                detailed_rows.append(cust_row)

        df_detailed_out = pd.DataFrame(detailed_rows)
        worksheet = writer.sheets[summary_sheet]
        worksheet.cell(row=start_row, column=1, value="Detailed Collections")
        df_detailed_out.to_excel(writer, sheet_name=summary_sheet, startrow=start_row, index=False)

        df_merged.to_excel(writer, sheet_name="Raw Data", index=False)

        total_collected = df_collected["Collected"].sum() if not df_collected.empty else 0
        avg_collected = df_collected["Collected"].mean() if not df_collected.empty else 0
        officer_count = len(df_collected["SalesRep"].unique()) if not df_collected.empty else 0

        stats_data = {
            "Metric": [
                "Total Loans Collected",
                "Total Amount Collected",
                "Average per Loan",
                "Number of Officers",
                "Total Loans Processed",
            ],
            "Value": [
                len(df_collected),
                f"{total_collected:,.2f}",
                f"{avg_collected:,.2f}",
                officer_count,
                len(df_merged),
            ],
        }

        if officer_targets:
            total_target = sum(officer_targets.values())
            overall_achievement = (total_collected / total_target * 100) if total_target > 0 else 0
            stats_data["Metric"].extend(["Total Target", "Overall Achievement %", "Total Remaining"])
            stats_data["Value"].extend(
                [
                    f"{total_target:,.2f}",
                    f"{overall_achievement:.1f}%" if total_target > 0 else "N/A",
                    f"{total_target - total_collected:,.2f}" if total_target > 0 else "N/A",
                ]
            )

        pd.DataFrame(stats_data).to_excel(writer, sheet_name="Statistics", index=False)

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                    except Exception:
                        pass
                ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _format_excel_report(output_path: str):
    """Apply professional formatting to Summary Report and Detailed Collections."""
    wb = load_workbook(output_path)
    if "Summary Report" not in wb.sheetnames:
        wb.save(output_path)
        return

    ws = wb["Summary Report"]
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = "ARREARS COLLECTION REPORT"
    title_cell.font = Font(bold=True, size=16, color="000000")
    title_cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    title_cell.alignment = center_align

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    date_cell = ws.cell(row=2, column=1)
    date_cell.value = f"Generated on: {datetime.now().strftime('%d-%b-%Y %H:%M:%S')}"
    date_cell.font = Font(italic=True, size=10, color="555555")
    date_cell.alignment = right_align

    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")
    total_font = Font(bold=True)
    section_title_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    row_fill_odd = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    row_fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    thick_top_border = Border(top=Side(style="thick", color="008080"), bottom=Side(style="double", color="008080"))

    summary_header_row = 5
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=summary_header_row, column=col)
        if cell.value:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

    summary_end_row = summary_header_row
    for r in range(summary_header_row + 1, ws.max_row + 1):
        cell = ws.cell(row=r, column=1)
        if cell.value == "Detailed Collections":
            break
        if cell.value == "GRAND TOTAL":
            summary_end_row = r
            break
        summary_end_row = r

    ws.auto_filter.ref = f"A{summary_header_row}:{get_column_letter(ws.max_column)}{summary_end_row}"

    for r in range(summary_header_row + 1, summary_end_row + 1):
        cell_val = str(ws.cell(row=r, column=1).value or "")
        is_total_row = "GRAND TOTAL" in cell_val
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if is_total_row:
                cell.fill = total_fill
                cell.font = total_font
                cell.border = thick_top_border
            else:
                cell.fill = row_fill_odd if (r - summary_header_row) % 2 != 0 else row_fill_even

            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c > 1 and cell.value is not None:
                try:
                    val_str = str(cell.value).replace(",", "").replace("%", "")
                    float(val_str)
                    cell.alignment = right_align
                except Exception:
                    pass

    detailed_start_row = None
    for row in range(summary_end_row + 1, ws.max_row + 1):
        cell = ws.cell(row=row, column=1)
        if cell.value == "Detailed Collections":
            detailed_start_row = row
            break

    if detailed_start_row:
        ws.merge_cells(start_row=detailed_start_row, start_column=1, end_row=detailed_start_row, end_column=8)
        title_cell = ws.cell(row=detailed_start_row, column=1)
        title_cell.value = "Detailed Collections"
        title_cell.font = Font(bold=True, size=12, color="000000")
        title_cell.fill = section_title_fill
        title_cell.alignment = Alignment(horizontal="left", vertical="center")

        table_header_row = detailed_start_row + 1
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=table_header_row, column=col)
            if cell.value:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

        current_row = table_header_row + 1
        while current_row <= ws.max_row:
            if ws.cell(row=current_row, column=1).value is None:
                current_row += 1
                continue

            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.border = thin_border
                cell.fill = row_fill_odd if (current_row - table_header_row) % 2 != 0 else row_fill_even

                header_val = ws.cell(row=table_header_row, column=col).value
                cell_val_b = str(ws.cell(row=current_row, column=2).value or "")
                is_ratio_row = "collected out of" in cell_val_b

                if is_ratio_row:
                    cell.font = Font(bold=True, italic=True)
                    if col == 2:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = center_align
                elif col == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                elif header_val and col > 1:
                    if "Age" in str(header_val) or "Days" in str(header_val):
                        cell.number_format = "0"
                        cell.alignment = center_align
                    elif isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
                        cell.alignment = right_align
            current_row += 1

        ach_col_letter = None
        for col in range(1, ws.max_column + 1):
            if "Achievement" in str(ws.cell(row=summary_header_row, column=col).value):
                ach_col_letter = get_column_letter(col)
                break

        if ach_col_letter:
            ws.conditional_formatting.add(
                f"{ach_col_letter}{summary_header_row + 1}:{ach_col_letter}{summary_end_row - 1}",
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="percentile",
                    end_value=90,
                    end_color="63BE7B",
                ),
            )

        coll_col_letter = None
        for col in range(1, ws.max_column + 1):
            if "Collected" in str(ws.cell(row=table_header_row, column=col).value):
                coll_col_letter = get_column_letter(col)
                break

        if coll_col_letter:
            ws.conditional_formatting.add(
                f"{coll_col_letter}{table_header_row + 1}:{coll_col_letter}{ws.max_row}",
                DataBarRule(start_type="min", end_type="max", color="008080"),
            )

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                lines = str(cell.value).split("\n")
                current_max = max(len(line) for line in lines)
                if current_max > max_length:
                    max_length = current_max
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 40)
        if column_letter == "A":
            adjusted_width = min(adjusted_width, 18)
        ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(output_path)


def _build_preview_chart(df_collected: pd.DataFrame, chart_path: str) -> str:
    """Create chart PNG for Android card preview."""
    try:
        rep_collected = (
            df_collected.groupby("SalesRep")["Collected"].sum().sort_values(ascending=False).head(10)
        )
        fig, ax = plt.subplots(figsize=(9, 5))
        ax.bar(rep_collected.index, rep_collected.values, color="#00C853", edgecolor="white")
        ax.set_title("Collections by Officer (Top 10)", fontweight="bold")
        ax.set_ylabel("KES Collected")
        ax.tick_params(axis="x", rotation=45)
        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def get_officers(sod_path: str) -> List[str]:
    """Return unique normalized officer names from the SOD file."""
    try:
        df_sod = _read_input(sod_path)
        df_sod = _normalize_officer_names(df_sod)
        if "SalesRep" not in df_sod.columns:
            return []
        return sorted(
            df_sod["SalesRep"]
            .astype(str)
            .str.strip()
            .replace("", np.nan)
            .dropna()
            .unique()
            .tolist()
        )
    except Exception:
        return []


def run(sod_path: str, current_path: str, output_dir: str, targets_input: str = "") -> dict:
    """Compute arrears collections and generate report workbook."""
    os.makedirs(output_dir, exist_ok=True)

    result = _process_data(sod_path, current_path, targets_input)
    if result is None:
        raise ValueError("Unable to process collections data")

    df_collected, df_merged, officers, officer_targets = result
    if df_collected.empty:
        raise ValueError("No collections found in the data")

    final_df = _create_formatted_table(df_collected, df_merged, officers, officer_targets)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"Arrears_Collection_Report_{timestamp}.xlsx")
    _create_excel_report(df_collected, df_merged, final_df, output_file, officer_targets)
    _format_excel_report(output_file)

    chart_file = _build_preview_chart(df_collected, os.path.join(output_dir, "collections_chart.png"))

    bucket_totals = {
        "1-15": float(df_collected[df_collected["Bucket"] == "1-15"]["Collected"].sum()),
        "16-30": float(df_collected[df_collected["Bucket"] == "16-30"]["Collected"].sum()),
        "31-180": float(df_collected[df_collected["Bucket"] == "31-180"]["Collected"].sum()),
    }
    total_collected = float(df_collected["Collected"].sum())
    total_target = float(sum(officer_targets.values())) if officer_targets else 0.0
    overall_achievement = (total_collected / total_target * 100.0) if total_target > 0 else 0.0

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            "total_collected": round(total_collected, 2),
            "bucket_1_15": round(bucket_totals.get("1-15", 0), 2),
            "bucket_16_30": round(bucket_totals.get("16-30", 0), 2),
            "bucket_31_180": round(bucket_totals.get("31-180", 0), 2),
            "unique_officers": int(len(officers)),
            "total_target": round(total_target, 2),
            "overall_achievement": round(overall_achievement, 2),
            "remaining_target": round(total_target - total_collected, 2) if total_target > 0 else 0.0,
        },
    }
