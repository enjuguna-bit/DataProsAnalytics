"""
Dues Processor
Premium dues report adapted from arrange dues for Android/Chaquopy.
"""

import logging
import os
import warnings
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import matplotlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

warnings.filterwarnings("ignore")

logger = logging.getLogger(__name__)
if not logger.handlers:
    logger.setLevel(logging.INFO)
    _stream = logging.StreamHandler()
    _stream.setFormatter(logging.Formatter("%(asctime)s - %(levelname)s - %(message)s"))
    logger.addHandler(_stream)


@dataclass
class ReportStats:
    """Container for dues report statistics."""

    total_clients: int = 0
    total_groups: int = 0
    total_amount_due: float = 0.0
    total_arrears: float = 0.0
    total_amount_paid: float = 0.0
    total_loan_balance: float = 0.0
    clients_with_arrears: int = 0
    average_arrears: float = 0.0
    completion_rate: float = 0.0
    total_branches: int = 0


class PremiumLoanReportGenerator:
    """Premium report pipeline without desktop UI dependencies."""

    def __init__(self, config_path: str = None):
        self.df = None
        self.final_df = None
        self.groups: List[str] = []
        self.grand_totals: Dict[str, float] = {}
        self.stats = ReportStats()
        self.config = self.load_config(config_path)

        self.COLORS = {
            "primary": "366092",
            "secondary": "4F81BD",
            "success": "00B050",
            "warning": "FFC000",
            "danger": "C00000",
            "info": "0070C0",
            "header_bg": "1F4E78",
            "subtotal_bg": "F2F2F2",
            "total_bg": "FFF2CC",
        }

    @staticmethod
    def normalize_column_name(col_name: str) -> str:
        normalized = str(col_name).strip().lower()
        normalized = normalized.replace("_", " ")
        while "  " in normalized:
            normalized = normalized.replace("  ", " ")
        return normalized

    def find_column_match(self, target_col: str, available_cols: List[str]) -> Optional[str]:
        target_normalized = self.normalize_column_name(target_col)

        for col in available_cols:
            if self.normalize_column_name(col) == target_normalized:
                return col

        target_no_space = target_normalized.replace(" ", "")
        for col in available_cols:
            col_no_space = self.normalize_column_name(col).replace(" ", "")
            if col_no_space == target_no_space:
                return col

        for col in available_cols:
            col_normalized = self.normalize_column_name(col)
            if col_normalized.replace(" ", "").replace("_", "") == target_no_space:
                return col

        return None

    def load_config(self, config_path: str) -> Dict:
        default_config = {
            "required_columns": ["FullNames", "FieldOfficer", "Amount Due"],
            "column_mapping": {
                "FullNames": "Client Name",
                "PhoneNumber": "Phone",
                "InstallmentNo": "Installment",
                "Amount Due": "Amount Due",
                "AmountDue": "Amount Due",
                "Arrears": "Arrears",
                "AmountPaid": "Amount Paid",
                "Amount Paid": "Amount Paid",
                "LoanBalance": "Loan Balance",
                "Loan Balance": "Loan Balance",
                "FieldOfficer": "Field Officer",
                "Field Officer": "Field Officer",
                "LoanID": "Loan ID",
                "Loan ID": "Loan ID",
                "DueDate": "Due Date",
                "Due Date": "Due Date",
                "Branch": "Branch",
            },
            "currency": "KES",
            "date_format": "%Y-%m-%d",
            "phone_format": "ke",
            "max_file_size_mb": 100,
            "auto_open_excel": False,
        }

        # Config file support is intentionally omitted in the mobile flow unless needed.
        if config_path and os.path.exists(config_path):
            try:
                import json

                with open(config_path, "r", encoding="utf-8") as f:
                    user_config = json.load(f)
                default_config.update(user_config)
            except Exception as exc:
                logger.warning("Could not load config %s: %s", config_path, exc)

        return default_config

    def validate_file(self, file_path: str) -> Tuple[bool, str]:
        try:
            if not os.path.exists(file_path):
                return False, "File does not exist"

            max_size = int(self.config["max_file_size_mb"]) * 1024 * 1024
            if os.path.getsize(file_path) > max_size:
                return False, f"File size exceeds {self.config['max_file_size_mb']}MB limit"

            valid_extensions = [".csv", ".xlsx", ".xls"]
            file_ext = os.path.splitext(file_path)[1].lower()
            if file_ext not in valid_extensions:
                return False, f"Invalid file type. Supported: {', '.join(valid_extensions)}"

            return True, "File validated successfully"
        except Exception as exc:
            return False, f"Validation error: {exc}"

    def load_data(self, file_path: str) -> bool:
        try:
            ok, msg = self.validate_file(file_path)
            if not ok:
                logger.error(msg)
                return False

            file_ext = os.path.splitext(file_path)[1].lower()
            if file_ext == ".csv":
                self.df = None
                for encoding in ("utf-8", "latin-1", "iso-8859-1", "cp1252"):
                    try:
                        self.df = pd.read_csv(file_path, encoding=encoding)
                        logger.info("Loaded CSV with encoding: %s", encoding)
                        break
                    except UnicodeDecodeError:
                        continue
            else:
                self.df = pd.read_excel(file_path)

            if self.df is None or self.df.empty:
                logger.error("Input file is empty or unreadable")
                return False

            logger.info("Loaded %s", os.path.basename(file_path))
            logger.info("Records: %s", f"{len(self.df):,}")
            logger.info("Columns: %s", list(self.df.columns))
            return True
        except pd.errors.EmptyDataError:
            logger.error("The file is empty")
            return False
        except Exception as exc:
            logger.error("Failed to load file: %s", exc)
            return False

    def clean_numeric_columns(self):
        if self.df is None:
            return

        numeric_patterns = ("amount", "arrears", "balance", "due", "paid", "installment")
        for col in self.df.columns:
            col_l = str(col).lower()
            if any(pattern in col_l for pattern in numeric_patterns):
                self.df[col] = (
                    self.df[col]
                    .astype(str)
                    .str.replace(r"[^\d\.\-]", "", regex=True)
                    .replace(["", "nan", "None", "none", "null", "NaN"], "0")
                )
                self.df[col] = pd.to_numeric(self.df[col], errors="coerce").fillna(0)

    def format_phone_numbers(self):
        if self.df is None:
            return

        def normalize_phone(phone):
            if pd.isna(phone) or str(phone).strip() in ("", "nan", "None", "none", "N/A", "n/a"):
                return ""

            phone_str = str(phone).strip()
            if phone_str.startswith("+"):
                digits = phone_str[0] + "".join(c for c in phone_str[1:] if c.isdigit())
            else:
                digits = "".join(c for c in phone_str if c.isdigit())

            if not digits:
                return ""

            if digits.startswith("+254") and len(digits) == 13:
                return f"{digits[0:4]} {digits[4:7]} {digits[7:10]} {digits[10:]}"
            if digits.startswith("254") and len(digits) == 12:
                return f"+{digits[0:3]} {digits[3:6]} {digits[6:9]} {digits[9:]}"
            if digits.startswith("07") and len(digits) == 10:
                return f"+254 {digits[1:4]} {digits[4:7]} {digits[7:]}"
            if len(digits) == 9 and digits[0] == "7":
                return f"+254 {digits[0:3]} {digits[3:6]} {digits[6:]}"

            return phone_str

        phone_cols = [col for col in self.df.columns if "phone" in str(col).lower()]
        for col in phone_cols:
            self.df[col] = self.df[col].apply(normalize_phone)

    def clean_dates(self):
        if self.df is None:
            return

        date_columns = [col for col in self.df.columns if "date" in str(col).lower()]
        for col in date_columns:
            parsed = pd.to_datetime(self.df[col], errors="coerce")
            self.df[col] = parsed.dt.strftime(self.config["date_format"]).fillna("")

    def calculate_additional_metrics(self):
        if self.df is None:
            return

        if "Installment" in self.df.columns and "Total Installments" in self.df.columns:
            denom = pd.to_numeric(self.df["Total Installments"], errors="coerce").replace(0, pd.NA)
            numer = pd.to_numeric(self.df["Installment"], errors="coerce")
            self.df["Completion %"] = ((numer / denom) * 100).fillna(0).round(1)

        if "Due Date" in self.df.columns:
            parsed = pd.to_datetime(self.df["Due Date"], errors="coerce")
            today = pd.Timestamp.now().normalize()
            overdue = (today - parsed).dt.days.fillna(0)
            self.df["Days Overdue"] = overdue.apply(lambda x: max(0, int(x)))

    def clean_and_prepare_data(self) -> bool:
        try:
            logger.info("Starting dues cleaning and preparation")
            self.df.columns = [str(col).strip().replace('"', "").replace("'", "") for col in self.df.columns]

            required_columns = [
                "Client Name",
                "Phone",
                "Installment",
                "Amount Due",
                "Arrears",
                "Amount Paid",
                "Loan Balance",
                "Field Officer",
                "Loan ID",
                "Due Date",
                "Branch",
            ]

            column_variations = {
                "Client Name": ["FullNames", "Name", "Client Name", "Client", "Customer Name", "Customer", "ClientName"],
                "Phone": ["Phone", "PhoneNumber", "Phone Number", "Mobile", "Telephone", "Contact", "PhoneNo", "Phone No"],
                "Installment": ["Installment", "InstallmentNo", "Installment No", "Inst No", "InstNo", "Inst", "Installment Number"],
                "Amount Due": ["Amount Due", "AmountDue", "Due Amount", "Due", "Total Due"],
                "Arrears": ["Arrears", "Arrear", "Outstanding", "Balance Due"],
                "Amount Paid": ["Amount Paid", "Paid", "AmountPaid", "Payment", "Amount Received"],
                "Loan Balance": ["Loan Balance", "LoanBalance", "Balance", "Remaining Balance", "Outstanding Balance"],
                "Field Officer": ["Field Officer", "FieldOfficer", "Officer", "Agent", "Loan Officer", "Account Manager", "SalesRep", "RO"],
                "Loan ID": ["Loan ID", "LoanID", "Loan No", "Loan Number", "Loan", "Account Number", "LoanId"],
                "Due Date": ["Due Date", "DueDate", "Payment Date", "Date Due", "Next Payment", "Payment Due"],
                "Branch": ["Branch", "Office", "Region"],
            }

            new_df = pd.DataFrame(index=self.df.index)
            for req_col in required_columns:
                variations = column_variations.get(req_col, [req_col])
                source_col = None

                for var in variations:
                    if var in self.df.columns:
                        source_col = var
                        break

                if source_col is None:
                    for col in self.df.columns:
                        if any(str(v).lower() == str(col).lower() for v in variations):
                            source_col = col
                            break

                if source_col is None:
                    for col in self.df.columns:
                        if any(str(v).lower() in str(col).lower() for v in variations):
                            source_col = col
                            break

                if source_col and source_col in self.df.columns:
                    new_df[req_col] = self.df[source_col]
                else:
                    if req_col in ("Installment",):
                        new_df[req_col] = 0
                    elif req_col in ("Amount Due", "Arrears", "Amount Paid", "Loan Balance"):
                        new_df[req_col] = 0.0
                    else:
                        new_df[req_col] = ""
                    logger.warning("No source mapping for column: %s", req_col)

            self.df = new_df
            self.clean_numeric_columns()
            self.format_phone_numbers()
            self.clean_dates()
            self.calculate_additional_metrics()

            for col in ("Installment", "Amount Due", "Arrears", "Amount Paid", "Loan Balance"):
                self.df[col] = pd.to_numeric(self.df[col], errors="coerce").fillna(0)

            key_columns = [col for col in ("Client Name", "Loan ID") if col in self.df.columns]
            if len(key_columns) == 2:
                before = len(self.df)
                self.df = self.df.drop_duplicates(subset=key_columns, keep="first")
                removed = before - len(self.df)
                if removed > 0:
                    logger.info("Removed duplicate rows: %s", removed)

            logger.info("Cleaning complete: shape=%s", self.df.shape)
            return True
        except Exception as exc:
            logger.error("Error in clean_and_prepare_data: %s", exc)
            return False

    def _find_field_officer_col(self) -> Optional[str]:
        if self.df is None:
            return None
        for col in self.df.columns:
            col_l = str(col).lower()
            if "field" in col_l or "officer" in col_l:
                return col
        return None

    def extract_groups_and_sort(self) -> bool:
        try:
            numeric_cols = ["Installment", "Amount Due", "Arrears", "Amount Paid", "Loan Balance"]
            self.grand_totals = {col: 0.0 for col in numeric_cols}

            field_officer_col = self._find_field_officer_col()
            if field_officer_col:
                self.df[field_officer_col] = self.df[field_officer_col].replace("", pd.NA).fillna("Unassigned")
                self.groups = sorted(self.df[field_officer_col].dropna().astype(str).unique().tolist())

                sort_columns = [field_officer_col, "Installment", "Client Name"]
                available = [col for col in sort_columns if col in self.df.columns]
                if "Installment" in self.df.columns:
                    self.df["Installment"] = pd.to_numeric(self.df["Installment"], errors="coerce").fillna(0)
                if available:
                    self.df = self.df.sort_values(by=available, ascending=[True] * len(available), kind="mergesort")
            else:
                self.groups = ["All Clients"]
                self.df["Field Officer"] = "All Clients"

            logger.info("Field-officer groups: %s", len(self.groups))
            return True
        except Exception as exc:
            logger.error("Error in extract_groups_and_sort: %s", exc)
            return False

    def create_subtotal_row(self, group_data: pd.DataFrame, group_name: str, numeric_cols: List[str]) -> Dict:
        subtotal_row = {col: "" for col in self.df.columns}
        subtotal_row["Client Name"] = f"Subtotal {group_name}"
        for col in numeric_cols:
            if col in group_data.columns:
                subtotal = float(group_data[col].sum())
                subtotal_row[col] = subtotal
                self.grand_totals[col] = float(self.grand_totals.get(col, 0.0)) + subtotal
        return subtotal_row

    def create_grand_total_row(self, numeric_cols: List[str]) -> Dict:
        total_row = {col: "" for col in self.df.columns}
        total_row["Client Name"] = "GRAND TOTALS"
        for col in numeric_cols:
            total_row[col] = float(self.grand_totals.get(col, 0.0))
        return total_row

    def calculate_statistics(self):
        if self.df is None or self.df.empty:
            return

        self.stats.total_clients = int(len(self.df))
        self.stats.total_groups = int(len(self.groups))
        self.stats.total_branches = int(self.df["Branch"].replace("", pd.NA).dropna().nunique()) if "Branch" in self.df.columns else 0

        self.stats.total_amount_due = float(self.df["Amount Due"].sum()) if "Amount Due" in self.df.columns else 0.0
        self.stats.total_arrears = float(self.df["Arrears"].sum()) if "Arrears" in self.df.columns else 0.0
        self.stats.total_amount_paid = float(self.df["Amount Paid"].sum()) if "Amount Paid" in self.df.columns else 0.0
        self.stats.total_loan_balance = float(self.df["Loan Balance"].sum()) if "Loan Balance" in self.df.columns else 0.0

        if "Arrears" in self.df.columns:
            self.stats.clients_with_arrears = int((self.df["Arrears"] > 0).sum())
            if self.stats.clients_with_arrears > 0:
                self.stats.average_arrears = float(self.stats.total_arrears / self.stats.clients_with_arrears)

        if "Completion %" in self.df.columns and len(self.df["Completion %"]) > 0:
            self.stats.completion_rate = float(pd.to_numeric(self.df["Completion %"], errors="coerce").fillna(0).mean())

    def add_premium_summary(self):
        summary_rows = []
        summary_header = {col: "" for col in self.final_df.columns}
        summary_header["Client Name"] = "REPORT SUMMARY & ANALYTICS"
        summary_rows.append(summary_header)

        arrears_pct = (
            (self.stats.clients_with_arrears / self.stats.total_clients * 100.0)
            if self.stats.total_clients > 0
            else 0.0
        )
        stats_lines = [
            f"Total Clients: {self.stats.total_clients:,}",
            f"Field Officers: {self.stats.total_groups}",
            f"Branches: {self.stats.total_branches}",
            f"Clients with Arrears: {self.stats.clients_with_arrears:,} ({arrears_pct:.1f}%)",
            f"Total Amount Due: {self.config['currency']} {self.stats.total_amount_due:,.2f}",
            f"Total Arrears: {self.config['currency']} {self.stats.total_arrears:,.2f}",
            f"Average Arrears: {self.config['currency']} {self.stats.average_arrears:,.2f}",
            f"Total Amount Paid: {self.config['currency']} {self.stats.total_amount_paid:,.2f}",
            f"Total Loan Balance: {self.config['currency']} {self.stats.total_loan_balance:,.2f}",
        ]

        for line in stats_lines:
            row = {col: "" for col in self.final_df.columns}
            row["Client Name"] = f"  - {line}"
            summary_rows.append(row)

        if len(self.groups) > 1 and "Arrears" in self.df.columns:
            perf_header = {col: "" for col in self.final_df.columns}
            perf_header["Client Name"] = "PERFORMANCE BY FIELD OFFICER"
            summary_rows.append(perf_header)

            field_officer_col = self._find_field_officer_col()
            if field_officer_col:
                for group in self.groups[:5]:
                    group_data = self.df[self.df[field_officer_col] == group]
                    arrears_count = int((group_data["Arrears"] > 0).sum())
                    total_arrears = float(group_data["Arrears"].sum())
                    perf_row = {col: "" for col in self.final_df.columns}
                    perf_row["Client Name"] = (
                        f"  * {group}: {arrears_count} clients in arrears "
                        f"({self.config['currency']} {total_arrears:,.2f})"
                    )
                    summary_rows.append(perf_row)

        summary_df = pd.DataFrame(summary_rows)
        self.final_df = pd.concat([self.final_df, summary_df], ignore_index=True)

    def generate_premium_report(self) -> bool:
        try:
            if not self.extract_groups_and_sort():
                return False

            final_rows = []
            numeric_cols = [c for c in ("Installment", "Amount Due", "Arrears", "Amount Paid", "Loan Balance") if c in self.df.columns]
            self.grand_totals = {col: 0.0 for col in numeric_cols}

            field_officer_col = self._find_field_officer_col()
            for idx, group_name in enumerate(self.groups, start=1):
                if group_name == "All Clients" or field_officer_col is None:
                    group_data = self.df.copy()
                else:
                    group_data = self.df[self.df[field_officer_col] == group_name].copy()

                if group_data.empty:
                    continue

                header_row = {col: "" for col in self.df.columns}
                header_row["Client Name"] = f"--- {str(group_name).upper()} ({len(group_data)} clients) ---"
                final_rows.append(header_row)

                for _, row in group_data.iterrows():
                    final_rows.append(row.to_dict())

                final_rows.append(self.create_subtotal_row(group_data, str(group_name), numeric_cols))
                if idx < len(self.groups):
                    final_rows.append({col: "" for col in self.df.columns})

            if final_rows:
                final_rows.append(self.create_grand_total_row(numeric_cols))

            self.final_df = pd.DataFrame(final_rows, columns=self.df.columns)
            self.calculate_statistics()
            self.add_premium_summary()
            return True
        except Exception as exc:
            logger.error("Error generating premium report: %s", exc)
            return False

    @staticmethod
    def _ensure_named_style(wb: Workbook, style: NamedStyle):
        style_names = set()
        for ns in wb.named_styles:
            if isinstance(ns, str):
                style_names.add(ns)
            else:
                style_names.add(ns.name)
        if style.name not in style_names:
            wb.add_named_style(style)

    def apply_premium_formatting(self, wb: Workbook, ws):
        header_style = NamedStyle(name="header_style")
        header_style.font = Font(bold=True, color="FFFFFF", size=12)
        header_style.fill = PatternFill(start_color=self.COLORS["header_bg"], fill_type="solid")
        header_style.alignment = Alignment(horizontal="center", vertical="center")
        self._ensure_named_style(wb, header_style)

        group_header_style = NamedStyle(name="group_header_style")
        group_header_style.font = Font(bold=True, color="366092", size=11)
        self._ensure_named_style(wb, group_header_style)

        subtotal_style = NamedStyle(name="subtotal_style")
        subtotal_style.font = Font(bold=True, color="804000")
        subtotal_style.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
        self._ensure_named_style(wb, subtotal_style)

        total_style = NamedStyle(name="total_style")
        total_style.font = Font(bold=True, color="FF0000", size=11)
        total_style.fill = PatternFill(start_color="FFFF00", fill_type="solid")
        self._ensure_named_style(wb, total_style)

        arrears_style = NamedStyle(name="arrears_style")
        arrears_style.font = Font(color="FF0000", bold=True)
        self._ensure_named_style(wb, arrears_style)

        headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
        installment_col = None
        arrears_cols = []
        amount_like_cols = []

        for i, header in enumerate(headers, start=1):
            header_s = str(header) if header is not None else ""
            if "Installment" in header_s and installment_col is None:
                installment_col = i
            if "Arrears" in header_s:
                arrears_cols.append(i)
            if any(word in header_s for word in ("Amount", "Arrears", "Balance", "Due", "Paid")):
                amount_like_cols.append(i)

        for row_idx in range(1, ws.max_row + 1):
            row_cells = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, min_col=1, max_col=ws.max_column))[0]
            if row_idx == 1:
                for cell in row_cells:
                    cell.style = "header_style"
                continue

            first_val = ws.cell(row=row_idx, column=1).value
            first_text = str(first_val) if first_val is not None else ""
            if first_text.startswith("---"):
                for cell in row_cells:
                    cell.style = "group_header_style"
            elif first_text.startswith("Subtotal"):
                for cell in row_cells:
                    cell.style = "subtotal_style"
            elif "GRAND TOTALS" in first_text:
                for cell in row_cells:
                    cell.style = "total_style"

            if installment_col is not None:
                inst_val = ws.cell(row=row_idx, column=installment_col).value
                try:
                    if float(inst_val) == 1.0:
                        for cell in row_cells:
                            cell.font = Font(color="FF0000", bold=True)
                except Exception:
                    pass

            for col_idx in arrears_cols:
                val = ws.cell(row=row_idx, column=col_idx).value
                if isinstance(val, (int, float)) and val > 0:
                    ws.cell(row=row_idx, column=col_idx).style = "arrears_style"

        for col_idx in amount_like_cols:
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, ws.max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        if arrears_cols:
            red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            arrears_rule = CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill)
            for col_idx in arrears_cols:
                col_letter = get_column_letter(col_idx)
                ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{ws.max_row}", arrears_rule)

    def add_charts_sheet(self, wb: Workbook):
        try:
            ws_charts = wb.create_sheet(title="Charts & Analytics")
            field_officer_col = self._find_field_officer_col()
            if field_officer_col is None or "Arrears" not in self.df.columns:
                return

            summary_data = []
            for group in self.groups:
                group_df = self.df[self.df[field_officer_col] == group]
                summary_data.append(
                    {
                        "Field Officer": group,
                        "Clients": len(group_df),
                        "Total Arrears": float(group_df["Arrears"].sum()),
                        "Average Arrears": float(group_df["Arrears"].mean()) if len(group_df) > 0 else 0.0,
                    }
                )

            if not summary_data:
                return

            summary_df = pd.DataFrame(summary_data).sort_values("Total Arrears", ascending=False)
            ws_charts.append(["Field Officer", "Clients", "Total Arrears", "Average Arrears"])
            for _, row in summary_df.iterrows():
                ws_charts.append(list(row))

            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Arrears by Field Officer"
            chart.y_axis.title = "Amount"
            chart.x_axis.title = "Field Officer"

            data_ref = Reference(ws_charts, min_col=3, min_row=1, max_row=len(summary_df) + 1, max_col=3)
            cats_ref = Reference(ws_charts, min_col=1, min_row=2, max_row=len(summary_df) + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.height = 8
            chart.width = 16
            ws_charts.add_chart(chart, "F2")

            for col_idx in range(1, ws_charts.max_column + 1):
                max_len = 0
                for row_idx in range(1, ws_charts.max_row + 1):
                    val = ws_charts.cell(row=row_idx, column=col_idx).value
                    if val is not None:
                        max_len = max(max_len, len(str(val)))
                ws_charts.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 30)
        except Exception as exc:
            logger.warning("Could not create charts sheet: %s", exc)

    def save_excel_report(self, output_path: str) -> bool:
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Loan Report"

            headers = list(self.final_df.columns)
            for c_idx, header in enumerate(headers, start=1):
                ws.cell(row=1, column=c_idx, value=str(header))

            for r_idx, row in enumerate(self.final_df.itertuples(index=False), start=2):
                for c_idx, value in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value)

            self.apply_premium_formatting(wb, ws)

            if len(self.groups) > 1:
                self.add_charts_sheet(wb)

            wb.save(output_path)
            logger.info("Saved premium dues report: %s", output_path)
            return True
        except Exception as exc:
            logger.error("Error saving Excel report: %s", exc)
            return False

    def save_csv_report(self, output_path: str) -> bool:
        try:
            self.final_df.to_csv(output_path, index=False, encoding="utf-8-sig")
            logger.info("Saved CSV dues report: %s", output_path)
            return True
        except Exception as exc:
            logger.error("Error saving CSV report: %s", exc)
            return False

    def save_report(self, output_path: str) -> bool:
        ext = os.path.splitext(output_path)[1].lower()
        if ext == ".xlsx":
            return self.save_excel_report(output_path)
        if ext == ".csv":
            return self.save_csv_report(output_path)
        logger.error("Unsupported output extension: %s", ext)
        return False

    def build_preview_chart(self, chart_path: str) -> str:
        """Generate PNG preview used by Android chart card."""
        try:
            if self.df is None or self.df.empty:
                return ""

            field_officer_col = self._find_field_officer_col()
            metric_col = "Arrears" if "Arrears" in self.df.columns else "Amount Due"
            if field_officer_col is None or metric_col not in self.df.columns:
                return ""

            chart_df = (
                self.df.groupby(field_officer_col)[metric_col]
                .sum()
                .sort_values(ascending=False)
                .head(10)
                .sort_values(ascending=True)
            )

            fig, ax = plt.subplots(figsize=(8, 5))
            ax.barh(chart_df.index.astype(str), chart_df.values, color="#1A237E")
            ax.set_title(f"Top Field Officers by {metric_col}", fontweight="bold")
            ax.set_xlabel("KES")
            fig.tight_layout()
            fig.savefig(chart_path, dpi=150, bbox_inches="tight")
            plt.close(fig)
            return chart_path
        except Exception as exc:
            logger.warning("Could not create preview chart: %s", exc)
            plt.close("all")
            return ""


def run(input_path: str, output_dir: str) -> dict:
    """App entrypoint: generate premium dues report and return JSON-safe dict."""
    os.makedirs(output_dir, exist_ok=True)

    generator = PremiumLoanReportGenerator()
    if not generator.load_data(input_path):
        raise ValueError("Failed to load dues input file")
    if not generator.clean_and_prepare_data():
        raise ValueError("Failed to clean dues data")
    if not generator.generate_premium_report():
        raise ValueError("Failed to build dues report")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"Premium_Loan_Report_{timestamp}.xlsx")
    if not generator.save_report(output_file):
        raise ValueError("Failed to save dues report")

    chart_file = generator.build_preview_chart(os.path.join(output_dir, "dues_chart.png"))

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            "total_clients": int(generator.stats.total_clients),
            "total_due": round(float(generator.stats.total_amount_due), 2),
            "total_arrears": round(float(generator.stats.total_arrears), 2),
            "num_branches": int(generator.stats.total_branches),
            "total_groups": int(generator.stats.total_groups),
            "clients_with_arrears": int(generator.stats.clients_with_arrears),
            "total_amount_paid": round(float(generator.stats.total_amount_paid), 2),
            "total_loan_balance": round(float(generator.stats.total_loan_balance), 2),
            "average_arrears": round(float(generator.stats.average_arrears), 2),
            "completion_rate": round(float(generator.stats.completion_rate), 2),
        },
    }
