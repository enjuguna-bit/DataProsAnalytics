"""
Branch Comparison Processor
Premium branch performance analysis adapted from MTD_parameters_branch_comparison.py.
"""

import os
from datetime import datetime
from typing import Dict, List, Optional

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt


def _read_input(path: str) -> pd.DataFrame:
    if path.lower().endswith(".csv"):
        df = pd.read_csv(path, dtype=str)
    else:
        df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _norm(text: str) -> str:
    return str(text).strip().lower().replace(" ", "").replace("_", "")


def _find_column(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    normalized = {_norm(col): col for col in df.columns}

    for candidate in candidates:
        key = _norm(candidate)
        if key in normalized:
            return normalized[key]

    for candidate in candidates:
        key = _norm(candidate)
        for col in df.columns:
            col_key = _norm(col)
            if key in col_key or col_key in key:
                return col
    return None


def _to_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(
        series.astype(str).str.replace(",", "", regex=False).str.replace(r"[^\d\.\-]", "", regex=True),
        errors="coerce",
    ).fillna(0)


def _normalize_branch(df: pd.DataFrame, branch_col: Optional[str], fallback_name: str) -> pd.DataFrame:
    work = df.copy()
    if branch_col and branch_col in work.columns:
        work["Branch Name"] = work[branch_col]
    else:
        work["Branch Name"] = fallback_name

    work["Branch Name"] = (
        work["Branch Name"].fillna("Unknown").astype(str).str.strip().replace("", "Unknown").str.title()
    )
    return work


def _prepare_income(df: pd.DataFrame) -> pd.DataFrame:
    branch_col = _find_column(df, ["Branch Name", "Branch", "BranchName"])
    income_col = _find_column(df, ["Income (KES)", "Income", "Total Amount", "Amount", "Total"])

    work = _normalize_branch(df, branch_col, "Unknown Income Branch")
    if income_col and income_col in work.columns:
        work["Income (KES)"] = _to_numeric(work[income_col])
    else:
        work["Income (KES)"] = 0

    out = work.groupby("Branch Name", as_index=False)["Income (KES)"].sum()
    return out


def _prepare_cr(df: pd.DataFrame) -> pd.DataFrame:
    branch_col = _find_column(df, ["Branch Name", "Branch", "BranchName"])
    collected_col = _find_column(df, ["Collected", "Paid Amount", "Paid", "Amount Paid"])
    uncollected_col = _find_column(df, ["Uncollected", "Outstanding Amount", "Outstanding", "Arrears"])
    cr_col = _find_column(df, ["CR %", "Collection Rate (%)", "Collection Rate", "CollectionRate"])
    total_col = _find_column(df, ["Total Amount", "Amount Due", "Total", "Due Amount"])

    work = _normalize_branch(df, branch_col, "Unknown CR Branch")

    work["Collected"] = _to_numeric(work[collected_col]) if collected_col in work.columns else 0
    work["Uncollected"] = _to_numeric(work[uncollected_col]) if uncollected_col in work.columns else 0
    work["CR %"] = _to_numeric(work[cr_col]) if cr_col in work.columns else np.nan

    if total_col and total_col in work.columns:
        total = _to_numeric(work[total_col])
    else:
        total = work["Collected"] + work["Uncollected"]

    # If one side missing, backfill from total.
    if collected_col is None and uncollected_col is not None:
        work["Collected"] = (total - work["Uncollected"]).clip(lower=0)
    if uncollected_col is None and collected_col is not None:
        work["Uncollected"] = (total - work["Collected"]).clip(lower=0)

    derived_cr = np.where((work["Collected"] + work["Uncollected"]) > 0,
                          work["Collected"] / (work["Collected"] + work["Uncollected"]) * 100, 0)
    work["CR %"] = work["CR %"].fillna(pd.Series(derived_cr))

    out = (
        work.groupby("Branch Name", as_index=False)
        .agg(
            Collected=("Collected", "sum"),
            Uncollected=("Uncollected", "sum"),
        )
    )
    out["CR %"] = np.where(
        (out["Collected"] + out["Uncollected"]) > 0,
        out["Collected"] / (out["Collected"] + out["Uncollected"]) * 100,
        0,
    )
    out["CR %"] = out["CR %"].round(2)
    return out


def _prepare_disbursement(df: pd.DataFrame) -> pd.DataFrame:
    branch_col = _find_column(df, ["Branch Name", "Branch", "BranchName"])
    disb_col = _find_column(df, ["Disbursement", "Disbursed Amount", "Disbursed", "Amount"])
    loan_col = _find_column(df, ["Loan Count", "Loans", "LoanCount", "Accounts"])

    work = _normalize_branch(df, branch_col, "Unknown Disbursement Branch")
    work["Disbursement"] = _to_numeric(work[disb_col]) if disb_col in work.columns else 0
    work["Loan Count"] = _to_numeric(work[loan_col]) if loan_col in work.columns else 0
    work["Loan Count"] = work["Loan Count"].round(0).astype(int)

    out = (
        work.groupby("Branch Name", as_index=False)
        .agg(
            Disbursement=("Disbursement", "sum"),
            Loan_Count_Num=("Loan Count", "sum"),
        )
    )
    out["Loan Count"] = out["Loan_Count_Num"].round(0).astype(int)
    out.drop(columns=["Loan_Count_Num"], inplace=True)
    return out


def _build_chart(data: pd.DataFrame, chart_path: str) -> str:
    try:
        top = data.nlargest(12, "Performance Score")
        if top.empty:
            fig, ax = plt.subplots(figsize=(8, 4))
            ax.text(0.5, 0.5, "No data", ha="center", va="center")
            ax.axis("off")
            fig.savefig(chart_path, dpi=150, bbox_inches="tight")
            plt.close(fig)
            return chart_path

        fig, axes = plt.subplots(1, 2, figsize=(14, 6))

        axes[0].barh(
            top["Branch Name"][::-1],
            top["Performance Score"].values[::-1],
            color="#1A237E",
            edgecolor="white",
        )
        axes[0].set_title("Top Branches by Performance Score", fontweight="bold")
        axes[0].set_xlabel("Performance Score")

        scatter = axes[1].scatter(
            top["Income (KES)"],
            top["CR %"],
            c=top["Disbursement"],
            cmap="viridis",
            s=90,
            alpha=0.85,
            edgecolors="black",
            linewidth=0.4,
        )
        axes[1].set_title("Income vs CR % (Color = Disbursement)", fontweight="bold")
        axes[1].set_xlabel("Income (KES)")
        axes[1].set_ylabel("CR %")
        axes[1].grid(True, alpha=0.3)
        fig.colorbar(scatter, ax=axes[1], label="Disbursement")

        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _apply_borders(cell):
    thin = Side(style="thin", color="000000")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _auto_width(ws):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 3, 14), 40)


def _write_excel(output_file: str, data: pd.DataFrame, chart_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Branch Performance"

    columns = [
        "Rank",
        "Branch Name",
        "Income (KES)",
        "CR %",
        "Collected",
        "Uncollected",
        "Disbursement",
        "Loan Count",
        "Avg Loan Size",
        "Performance Score",
    ]

    ws.merge_cells("A1:J1")
    title = ws["A1"]
    title.value = "PREMIUM BRANCH PERFORMANCE REPORT"
    title.font = Font(bold=True, size=15, color="FFFFFF")
    title.fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    title.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:J2")
    subtitle = ws["A2"]
    subtitle.value = f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    subtitle.font = Font(bold=True, size=10, color="000000")
    subtitle.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    subtitle.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, name in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=name)
        cell.font = Font(bold=True, size=11, color="000000")
        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        _apply_borders(cell)

    for row_idx, (_, row) in enumerate(data.iterrows(), 4):
        values = [
            int(row["Rank"]),
            str(row["Branch Name"]),
            float(row["Income (KES)"]),
            float(row["CR %"]),
            float(row["Collected"]),
            float(row["Uncollected"]),
            float(row["Disbursement"]),
            int(row["Loan Count"]),
            float(row["Avg Loan Size"]),
            float(row["Performance Score"]),
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            _apply_borders(cell)

            if row_idx % 2 == 0 and col_idx in (1, 2, 5, 6, 9, 10):
                cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

            if col_idx in (3, 5, 6, 7):
                cell.number_format = "#,##0"
            elif col_idx == 4:
                cell.number_format = "0.00"
            elif col_idx == 8:
                cell.number_format = "#,##0"
            elif col_idx == 9:
                cell.number_format = "#,##0.00"
            elif col_idx == 10:
                cell.number_format = "0.0"

    end_row = len(data) + 3
    ws.auto_filter.ref = f"A3:J{end_row}"
    ws.freeze_panes = "A4"

    ws.conditional_formatting.add(
        f"C4:C{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color="FF0000",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFF00",
            end_type="max",
            end_color="00B050",
        ),
    )
    ws.conditional_formatting.add(
        f"D4:D{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color="FF0000",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFF00",
            end_type="max",
            end_color="00B050",
        ),
    )
    ws.conditional_formatting.add(
        f"G4:G{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color="FF0000",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFF00",
            end_type="max",
            end_color="00B050",
        ),
    )
    ws.conditional_formatting.add(
        f"H4:H{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color="FF0000",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFF00",
            end_type="max",
            end_color="00B050",
        ),
    )

    ws_summary = wb.create_sheet("Summary")
    summary_rows = [
        ["SUMMARY STATISTICS", "VALUE"],
        ["", ""],
        ["Total Branches Analyzed", len(data)],
        ["Total Income (KES)", float(data["Income (KES)"].sum())],
        ["Average Income (KES)", float(data["Income (KES)"].mean())],
        ["Total Disbursement (KES)", float(data["Disbursement"].sum())],
        ["Average Disbursement (KES)", float(data["Disbursement"].mean())],
        ["Total Loan Count", int(data["Loan Count"].sum())],
        ["Average Loan Size (KES)", float(data["Avg Loan Size"].mean())],
        ["Average CR %", float(data["CR %"].mean())],
        ["Average Performance Score", float(data["Performance Score"].mean())],
    ]

    for row_idx, values in enumerate(summary_rows, 1):
        for col_idx, value in enumerate(values, 1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=value)
            _apply_borders(cell)
            if row_idx == 1:
                cell.fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center")
            elif row_idx == 3 and col_idx == 1:
                cell.font = Font(bold=True)
            else:
                if col_idx == 1:
                    cell.font = Font(bold=True)

    top_cr_start = len(summary_rows) + 2
    ws_summary.cell(row=top_cr_start, column=1, value="TOP 5 BRANCHES BY CR %").font = Font(bold=True)
    top_cr = data.nlargest(5, "CR %")[["Branch Name", "CR %"]]
    for i, row in enumerate(top_cr.itertuples(index=False), top_cr_start + 1):
        ws_summary.cell(row=i, column=1, value=row[0])
        ws_summary.cell(row=i, column=2, value=float(row[1]))
        ws_summary.cell(row=i, column=2).number_format = "0.00"

    top_income_start = top_cr_start + len(top_cr) + 2
    ws_summary.cell(row=top_income_start, column=1, value="TOP 5 BRANCHES BY INCOME").font = Font(bold=True)
    top_income = data.nlargest(5, "Income (KES)")[["Branch Name", "Income (KES)"]]
    for i, row in enumerate(top_income.itertuples(index=False), top_income_start + 1):
        ws_summary.cell(row=i, column=1, value=row[0])
        ws_summary.cell(row=i, column=2, value=float(row[1]))
        ws_summary.cell(row=i, column=2).number_format = "#,##0"

    _auto_width(ws)
    _auto_width(ws_summary)

    if chart_path and os.path.exists(chart_path):
        ws_chart = wb.create_sheet("Chart")
        ws_chart.add_image(XLImage(chart_path), "A1")

    wb.save(output_file)


def run(income_path: str, cr_path: str, disb_path: str, output_dir: str) -> dict:
    os.makedirs(output_dir, exist_ok=True)

    income_df = _prepare_income(_read_input(income_path))
    cr_df = _prepare_cr(_read_input(cr_path))
    disb_df = _prepare_disbursement(_read_input(disb_path))

    merged = income_df.merge(cr_df, on="Branch Name", how="outer").merge(disb_df, on="Branch Name", how="outer")
    merged = merged.fillna(0)

    merged["Loan Count"] = merged["Loan Count"].astype(int)
    merged["Avg Loan Size"] = np.where(
        merged["Loan Count"] > 0,
        merged["Disbursement"] / merged["Loan Count"],
        0,
    )
    merged["Avg Loan Size"] = merged["Avg Loan Size"].round(2)

    income_max = merged["Income (KES)"].max()
    disb_max = merged["Disbursement"].max()
    income_norm = merged["Income (KES)"] / income_max if income_max > 0 else 0
    disb_norm = merged["Disbursement"] / disb_max if disb_max > 0 else 0
    cr_norm = merged["CR %"] / 100

    merged["Performance Score"] = ((cr_norm * 0.4) + (income_norm * 0.3) + (disb_norm * 0.3)) * 100
    merged["Performance Score"] = merged["Performance Score"].round(2)

    merged = merged.sort_values("Performance Score", ascending=False).reset_index(drop=True)
    merged.insert(0, "Rank", range(1, len(merged) + 1))

    chart_path = os.path.join(output_dir, "branch_chart.png")
    chart_file = _build_chart(merged, chart_path)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"Branch_Comparison_Premium_{timestamp}.xlsx")
    _write_excel(output_file, merged, chart_file)

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            "num_branches": int(len(merged)),
            "total_income": round(float(merged["Income (KES)"].sum()), 2),
            "total_credit_risk": round(float(merged["Uncollected"].sum()), 2),
            "total_disbursement": round(float(merged["Disbursement"].sum()), 2),
            "avg_cr_pct": round(float(merged["CR %"].mean()), 2),
            "total_loans": int(merged["Loan Count"].sum()),
            "avg_performance_score": round(float(merged["Performance Score"].mean()), 2),
        },
    }
