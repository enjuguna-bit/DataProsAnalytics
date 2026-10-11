"""
MTD CR Portfolio Processor
Converts raw MTD CR portfolio data into a processed Excel workbook.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List, Tuple

import matplotlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from processors.tabular_io import read_tabular_file


DEFAULT_TARGET_CR = 93

TARGET_CR_BY_ID: Dict[int, int] = {
    288: 95,
    120: 92,
    155: 91,
    69: 93,
    145: 94,
}

TARGET_CR_BY_NAME: Dict[str, int] = {
    "JOASH MUTISYA": 95,
    "BRIAN WANJALA": 92,
    "ESTHER MUMBI": 91,
    "ALEXANDER MWANGI": 93,
    "BENSON KIRUI": 94,
}


def _read_input(path: str) -> pd.DataFrame:
    return read_tabular_file(path, dtype=str)


def _normalize_header(name: str) -> str:
    clean = str(name or "").strip().lower()
    clean = clean.replace("%", "pct")
    clean = re.sub(r"[^a-z0-9]+", "", clean)
    return clean


def _to_number(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(r"[^\d\.\-]", "", regex=True)
        .replace("", "0")
        .pipe(pd.to_numeric, errors="coerce")
        .fillna(0.0)
    )


def _pick_column(columns: List[str], candidates: List[str]) -> str:
    lookup = {_normalize_header(c): c for c in columns}
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in lookup:
            return lookup[key]
    return ""


def _parse_targets_input(targets_input: str) -> Tuple[int, Dict[str, int]]:
    """
    Parse dynamic target CR input lines.
    Supported format:
      DEFAULT=93
      Officer Name=95
    """
    default_target = DEFAULT_TARGET_CR
    overrides: Dict[str, int] = {}

    if not targets_input or not str(targets_input).strip():
        return default_target, overrides

    for raw_line in str(targets_input).splitlines():
        line = raw_line.strip()
        if not line or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip().upper()
        value = value.strip()
        if not value:
            continue

        try:
            parsed = int(round(float(value)))
        except Exception:
            continue
        parsed = max(0, min(100, parsed))
        if parsed == 0:
            continue

        if key in {"DEFAULT", "ALL", "*"}:
            default_target = parsed
        else:
            overrides[key] = parsed

    return default_target, overrides


def _resolve_target_cr(
    ids: pd.Series,
    names: pd.Series,
    target_cr_input: pd.Series | None,
    default_target_cr: int,
    overrides_by_name: Dict[str, int],
) -> pd.Series:
    resolved: List[int] = []
    for idx, name in names.items():
        officer_name = str(name).strip().upper()

        if officer_name in overrides_by_name:
            resolved.append(overrides_by_name[officer_name])
            continue

        if target_cr_input is not None and idx in target_cr_input.index:
            target_val = int(max(0, min(100, round(float(target_cr_input.loc[idx])))))
            if target_val > 0:
                resolved.append(target_val)
                continue

        officer_id = int(ids.loc[idx]) if idx in ids.index else 0
        if officer_id in TARGET_CR_BY_ID:
            resolved.append(TARGET_CR_BY_ID[officer_id])
            continue

        if officer_name in TARGET_CR_BY_NAME:
            resolved.append(TARGET_CR_BY_NAME[officer_name])
            continue

        resolved.append(default_target_cr)

    return pd.Series(resolved, index=names.index, dtype=int)


def _build_output_rows(df: pd.DataFrame, targets_input: str = "") -> Tuple[pd.DataFrame, Dict[str, float]]:
    if df.empty:
        raise ValueError("Input file is empty")

    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    cols = list(df.columns)

    id_col = _pick_column(cols, ["#ID", "ID", "OfficerID"])
    name_col = _pick_column(cols, ["Full Name", "FullName", "Name", "SalesRep", "Officer"])
    total_col = _pick_column(cols, ["Total Amount", "TotalAmount", "AmountDue", "Amount Due"])
    paid_col = _pick_column(cols, ["Paid Amount", "PaidAmount", "AmountPaid", "Amount Paid"])
    outstanding_col = _pick_column(
        cols, ["Outstanding Amount", "OutstandingAmount", "Outstanding", "Arrears", "Balance"]
    )
    cr_col = _pick_column(cols, ["Collection Rate (%)", "CollectionRate", "CR%", "CR"])
    target_cr_col = _pick_column(cols, ["Target CR", "TargetCR", "Target (%)", "CR Target", "Target"])
    target_amount_col = _pick_column(cols, ["TargetAmount", "Target Amount"])

    if not name_col or not total_col or not paid_col:
        raise ValueError("Missing required columns. Need Full Name, Total Amount, and Paid Amount.")

    names = df[name_col].fillna("").astype(str).str.strip().str.replace(r"\s+", " ", regex=True).str.upper()
    filtered = names != ""
    if not filtered.any():
        raise ValueError("No valid rows found in input")

    base = df.loc[filtered].copy()
    names = names.loc[filtered]

    total_amount = _to_number(base[total_col]).round(2)
    paid_amount = _to_number(base[paid_col]).round(2)

    if outstanding_col:
        outstanding_amount = _to_number(base[outstanding_col]).round(2)
    else:
        outstanding_amount = (total_amount - paid_amount).clip(lower=0).round(2)

    if cr_col:
        current_cr = _to_number(base[cr_col]).clip(lower=0, upper=100).round(2)
    else:
        current_cr = ((paid_amount / total_amount.replace(0, pd.NA)) * 100).fillna(0).clip(lower=0, upper=100).round(2)

    if id_col:
        ids = _to_number(base[id_col]).astype(int)
    else:
        ids = pd.Series([0] * len(base), index=base.index, dtype=int)

    targets_text = str(targets_input or "").strip()
    default_target_cr, overrides_by_name = _parse_targets_input(targets_text)

    if targets_text:
        target_cr = names.apply(
            lambda value: int(overrides_by_name.get(str(value).strip().upper(), default_target_cr))
        ).astype(int)
    else:
        target_cr_input = _to_number(base[target_cr_col]).round(0) if target_cr_col else None
        target_cr = _resolve_target_cr(
            ids=ids,
            names=names,
            target_cr_input=target_cr_input,
            default_target_cr=default_target_cr,
            overrides_by_name=overrides_by_name,
        )

    computed_target_amount = (total_amount * (target_cr / 100.0) - paid_amount).clip(lower=0).round(0).astype(int)

    if target_amount_col:
        raw_target_amount = _to_number(base[target_amount_col]).round(0).astype(int)
        target_amount = raw_target_amount.where(raw_target_amount > 0, computed_target_amount)
    else:
        target_amount = computed_target_amount

    out = pd.DataFrame()
    out["Full Name"] = names.values
    out["Total Amount"] = total_amount.values
    out["Paid Amount"] = paid_amount.values
    out["Outstanding"] = outstanding_amount.values
    out["Current CR"] = current_cr.values
    out["Target CR"] = target_cr.values
    out["TargetAmount"] = target_amount.values
    out["CR%"] = out.apply(lambda r: f"{float(r['Current CR']):.2f}/{int(r['Target CR'])}", axis=1)
    out["RowStatus"] = out.apply(
        lambda r: "On/Above Target" if float(r["Current CR"]) >= float(r["Target CR"]) else "Below Target",
        axis=1,
    )

    out = out.sort_values(by=["Current CR", "Outstanding"], ascending=[False, False]).reset_index(drop=True)

    total_amount_sum = float(out["Total Amount"].sum())
    total_paid_sum = float(out["Paid Amount"].sum())
    total_outstanding_sum = float(out["Outstanding"].sum())
    weighted_current_cr = (total_paid_sum / total_amount_sum * 100.0) if total_amount_sum > 0 else 0.0
    weighted_target_cr = (
        (out["Total Amount"] * out["Target CR"]).sum() / total_amount_sum if total_amount_sum > 0 else DEFAULT_TARGET_CR
    )
    total_target_amount = int(out["TargetAmount"].sum())

    total_row = pd.DataFrame(
        [
            {
                "Full Name": "",
                "Total Amount": "",
                "Paid Amount": "",
                "Outstanding": int(round(total_outstanding_sum)),
                "Current CR": round(weighted_current_cr, 2),
                "Target CR": int(round(weighted_target_cr)),
                "TargetAmount": total_target_amount,
                "CR%": f"{weighted_current_cr:.2f}/{int(round(weighted_target_cr))}",
                "RowStatus": "TOTAL",
            }
        ]
    )

    final_df = pd.concat([out, total_row], ignore_index=True)

    stats = {
        "total_portfolios": int(len(out)),
        "total_amount": round(total_amount_sum, 2),
        "total_paid": round(total_paid_sum, 2),
        "total_outstanding": round(total_outstanding_sum, 2),
        "avg_collection_rate": round(weighted_current_cr, 2),
        "avg_target_cr": round(float(out["Target CR"].mean()) if len(out) else float(default_target_cr), 2),
        "total_target_amount": int(total_target_amount),
        "below_target_count": int((out["Current CR"] < out["Target CR"]).sum()),
    }
    return final_df, stats


def _build_preview_chart(df: pd.DataFrame, chart_path: str) -> str:
    try:
        body = df[df["RowStatus"] != "TOTAL"].copy()
        if body.empty:
            return ""
        body = body.sort_values("Current CR", ascending=True)

        fig, ax = plt.subplots(figsize=(9, 5))
        ax.barh(body["Full Name"], body["Current CR"], color="#1E88E5", label="Current CR")
        ax.scatter(body["Target CR"], body["Full Name"], color="#C62828", marker="D", s=40, label="Target CR")
        ax.set_xlim(0, 100)
        ax.set_xlabel("Collection Rate (%)")
        ax.set_title("MTD CR Per Portfolio - Current vs Target", fontweight="bold")
        ax.legend(loc="lower right")
        fig.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _write_workbook(df: pd.DataFrame, output_file: str, chart_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Processed Data"

    headers = ["Full Name", "Total Amount", "Paid Amount", "Outstanding", "CR%", "TargetAmount"]

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    hdr_fill = PatternFill("solid", fgColor="0D47A1")
    hdr_font = Font(bold=True, color="FFFFFF")
    red_font = Font(color="C62828", bold=True)

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center
        cell.border = border

    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        is_total = str(row.get("RowStatus", "")) == "TOTAL"
        row_fill = PatternFill("solid", fgColor="F5F5F5" if is_total else "FFFFFF")

        output_values = [
            row.get("Full Name"),
            row.get("Total Amount"),
            row.get("Paid Amount"),
            row.get("Outstanding"),
            row.get("CR%"),
            row.get("TargetAmount"),
        ]

        for c_idx, value in enumerate(output_values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            cell.fill = row_fill

            if c_idx == 1:
                cell.alignment = left
            else:
                cell.alignment = right

            if c_idx in (2, 3, 4, 6) and value != "":
                cell.number_format = "#,##0"

            if c_idx == 5:
                cell.alignment = center

            if not is_total:
                if str(row.get("RowStatus", "")) == "Below Target" and c_idx in (1, 2, 3, 4, 5, 6):
                    cell.font = red_font
            else:
                cell.font = Font(bold=True)

    for col_idx in range(1, ws.max_column + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 34)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    summary = wb.create_sheet("Summary")
    body = df[df["RowStatus"] != "TOTAL"].copy()
    total_row = df[df["RowStatus"] == "TOTAL"].copy()

    total_amount = float(body["Total Amount"].sum()) if len(body) else 0.0
    total_paid = float(body["Paid Amount"].sum()) if len(body) else 0.0
    total_outstanding = float(body["Outstanding"].sum()) if len(body) else 0.0
    avg_current = (total_paid / total_amount * 100) if total_amount > 0 else 0.0
    avg_target = float(body["Target CR"].mean()) if len(body) else DEFAULT_TARGET_CR
    total_target_amount = int(body["TargetAmount"].sum()) if len(body) else 0
    below_target = int((body["Current CR"] < body["Target CR"]).sum()) if len(body) else 0
    overall_cr_text = str(total_row.iloc[0]["CR%"]) if len(total_row) else f"{avg_current:.2f}/{int(round(avg_target))}"

    rows = [
        ("Metric", "Value"),
        ("Total Portfolios", int(len(body))),
        ("Total Amount", total_amount),
        ("Total Paid", total_paid),
        ("Total Outstanding", total_outstanding),
        ("Overall CR/Target", overall_cr_text),
        ("Average Collection Rate (%)", avg_current),
        ("Average Target CR (%)", avg_target),
        ("Total Target Amount", total_target_amount),
        ("Below Target Count", below_target),
    ]

    for r_idx, (metric, value) in enumerate(rows, start=1):
        c1 = summary.cell(row=r_idx, column=1, value=metric)
        c2 = summary.cell(row=r_idx, column=2, value=value)
        c1.border = border
        c2.border = border

        if r_idx == 1:
            c1.fill = hdr_fill
            c2.fill = hdr_fill
            c1.font = hdr_font
            c2.font = hdr_font
            c1.alignment = center
            c2.alignment = center
            continue

        c1.font = Font(bold=True)
        c1.alignment = left
        c2.alignment = right
        if isinstance(value, (int, float)):
            if "Rate" in metric or "Target CR" in metric:
                c2.number_format = '0.00"%"'
            else:
                c2.number_format = "#,##0"
        else:
            c2.alignment = center

    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 22

    if chart_path and os.path.exists(chart_path):
        from openpyxl.drawing.image import Image as XLImage

        chart_ws = wb.create_sheet("Chart")
        chart_ws.add_image(XLImage(chart_path), "A1")

    wb.save(output_file)


def get_officers(input_path: str) -> List[str]:
    try:
        df = _read_input(input_path)
        if df is None or df.empty:
            return []
        df.columns = [str(c).strip() for c in df.columns]
        name_col = _pick_column(list(df.columns), ["Full Name", "FullName", "Name", "SalesRep", "Officer"])
        if not name_col:
            return []
        names = (
            df[name_col]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.replace(r"\s+", " ", regex=True)
            .str.upper()
        )
        return sorted(names[names != ""].drop_duplicates().tolist())
    except Exception:
        return []


def run(input_path: str, output_dir: str, targets_input: str = "") -> dict:
    os.makedirs(output_dir, exist_ok=True)
    df = _read_input(input_path)
    final_df, stats = _build_output_rows(df, targets_input=targets_input)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"MTD_CR_Per_Portfolio_{timestamp}.xlsx")
    chart_file = _build_preview_chart(final_df, os.path.join(output_dir, "mtd_cr_portfolio_chart.png"))
    _write_workbook(final_df, output_file, chart_file)

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": stats,
    }
