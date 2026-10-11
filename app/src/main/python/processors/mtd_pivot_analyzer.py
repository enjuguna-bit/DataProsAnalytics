"""
MTD Disbursement Pivot Analyzer
Android-safe adaptation of the desktop Tkinter MTD pivot analyzer.
"""

import os
from datetime import datetime
from typing import Dict, List, Optional

import matplotlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from processors.tabular_io import read_tabular_file


DISPLAY_COLUMNS = [
    "Field Officer",
    "New",
    "Repeat",
    "Grand Total",
    "MTD Target",
    "D/S",
    "Today's Target",
]


def _read_input(path: str) -> pd.DataFrame:
    df = read_tabular_file(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _norm(text: str) -> str:
    return str(text).strip().lower().replace("_", "").replace(" ", "")


def _find_column(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    normalized_map = {_norm(c): c for c in df.columns}
    for candidate in candidates:
        key = _norm(candidate)
        if key in normalized_map:
            return normalized_map[key]

    for candidate in candidates:
        key = _norm(candidate)
        for col in df.columns:
            col_key = _norm(col)
            if key in col_key or col_key in key:
                return col
    return None


def _normalize_loan_type(value: object) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return "Other"
    if "new" in text:
        return "New"
    if "repeat" in text:
        return "Repeat"
    return text.title()


def _to_int(value: str) -> int:
    text = str(value or "").strip().replace(",", "")
    if text == "":
        return 0
    try:
        number = int(float(text))
    except Exception:
        return 0
    return max(0, number)


def _parse_targets(text: str, officers: List[str]) -> Dict[str, int]:
    targets = {officer: 0 for officer in officers}
    if not text:
        return targets

    lookup = {officer.strip().lower(): officer for officer in officers}

    for raw_line in str(text).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if "=" in line:
            name, value = line.split("=", 1)
        else:
            name, value = line, "0"
        key = name.strip().lower()
        if key not in lookup:
            continue
        targets[lookup[key]] = _to_int(value)

    return targets


def _build_chart(rows: List[Dict], chart_path: str) -> str:
    try:
        officers = [r["field_officer"] for r in rows if not r["is_grand_total"]]
        grand_totals = [r["grand_total"] for r in rows if not r["is_grand_total"]]
        mtd_targets = [r["mtd_target"] for r in rows if not r["is_grand_total"]]

        if not officers:
            fig, ax = plt.subplots(figsize=(8, 4))
            ax.text(0.5, 0.5, "No data", ha="center", va="center")
            ax.axis("off")
            fig.savefig(chart_path, dpi=150, bbox_inches="tight")
            plt.close(fig)
            return chart_path

        x = list(range(len(officers)))
        fig, ax = plt.subplots(figsize=(max(8, len(officers) * 0.9), 5))
        ax.bar([i - 0.2 for i in x], grand_totals, width=0.4, label="Grand Total", color="#1E88E5")
        ax.bar([i + 0.2 for i in x], mtd_targets, width=0.4, label="MTD Target", color="#FB8C00")
        ax.set_xticks(x)
        ax.set_xticklabels(officers, rotation=45, ha="right", fontsize=8)
        ax.set_title("MTD Disbursement vs Target by Field Officer", fontweight="bold")
        ax.set_ylabel("Count")
        ax.legend()
        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _auto_width(ws):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)


def _write_excel(output_file: str, rows: List[Dict], day_header: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot Analysis"

    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    center = Alignment(horizontal="center", vertical="center")

    header_font = Font(bold=True, color="000000", size=12)
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    grand_font = Font(bold=True, color="FFFFFF")
    grand_fill = PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid")

    positive_fill = PatternFill(start_color="E8F5E8", end_color="E8F5E8", fill_type="solid")
    negative_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    red_dark = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")
    red_medium = PatternFill(start_color="FF3333", end_color="FF3333", fill_type="solid")
    red_light = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    red_faint = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    title = ws.cell(row=1, column=1, value="LOAN PERFORMANCE ANALYSIS REPORT")
    title.font = Font(bold=True, size=16, color="FF0000")
    title.alignment = center

    for idx, header in enumerate(DISPLAY_COLUMNS, 1):
        cell = ws.cell(row=2, column=idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    row_num = 3
    for row in rows:
        values = [
            row["field_officer"],
            row["new_display"],
            row["repeat_count"],
            row["grand_total"],
            row["mtd_display"],
            row["ds_display"],
            row["today_display"],
        ]
        for col_num, value in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = center
            cell.border = thin

            if row["is_grand_total"]:
                cell.fill = grand_fill
                cell.font = grand_font

        if not row["is_grand_total"]:
            if row["ds_value"] is not None:
                ws.cell(row=row_num, column=6).fill = positive_fill if row["ds_value"] >= 0 else negative_fill
            if row["today_target"] is not None and row["today_target"] > 0:
                if row["today_target"] >= 10:
                    ws.cell(row=row_num, column=7).fill = red_dark
                elif row["today_target"] >= 7:
                    ws.cell(row=row_num, column=7).fill = red_medium
                elif row["today_target"] >= 4:
                    ws.cell(row=row_num, column=7).fill = red_light
                else:
                    ws.cell(row=row_num, column=7).fill = red_faint

        row_num += 1

    summary_header_row = row_num + 1
    ws.merge_cells(start_row=summary_header_row, start_column=1, end_row=summary_header_row, end_column=7)
    summary_head = ws.cell(row=summary_header_row, column=1, value="ANALYSIS SUMMARY")
    summary_head.font = Font(bold=True, size=12, color="FFFFFF")
    summary_head.fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
    summary_head.alignment = center

    officers = [r for r in rows if not r["is_grand_total"]]
    positive_performers = sum(1 for r in officers if r["ds_value"] is not None and r["ds_value"] >= 0)
    negative_performers = sum(1 for r in officers if r["ds_value"] is not None and r["ds_value"] < 0)

    info_row = summary_header_row + 2
    ws.cell(row=info_row, column=1, value="Performance Summary").font = Font(bold=True, size=11)
    info_row += 1
    ws.cell(row=info_row, column=1, value=f"Total Officers: {len(officers)}")
    info_row += 1
    good = ws.cell(row=info_row, column=1, value=f"Meeting/Exceeding Target: {positive_performers}")
    good.fill = positive_fill
    info_row += 1
    bad = ws.cell(row=info_row, column=1, value=f"Below Target: {negative_performers}")
    bad.fill = negative_fill

    meta_row = ws.max_row + 2
    ws.cell(row=meta_row, column=1, value="REPORT METADATA").font = Font(bold=True, size=10)
    meta_row += 1
    ws.cell(row=meta_row, column=1, value=f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    meta_row += 1
    ws.cell(row=meta_row, column=1, value=f"Day Header: {day_header}")
    meta_row += 1
    ws.cell(row=meta_row, column=1, value="Generated by: MTD Disbursement Pivot Analyzer")

    ws.freeze_panes = "A3"
    _auto_width(ws)
    wb.save(output_file)


def _prepare_base(input_path: str) -> pd.DataFrame:
    df = _read_input(input_path)
    field_col = _find_column(df, ["FieldOfficer", "Field Officer", "SalesRep", "Officer", "RO"])
    type_col = _find_column(df, ["Loantype", "LoanType", "Loan Type", "Type"])

    if not field_col or not type_col:
        missing = []
        if not field_col:
            missing.append("FieldOfficer")
        if not type_col:
            missing.append("Loantype")
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    work = pd.DataFrame()
    work["FieldOfficer"] = df[field_col].fillna("Unknown").astype(str).str.strip()
    work["FieldOfficer"] = work["FieldOfficer"].replace("", "Unknown")
    work["LoanType"] = df[type_col].apply(_normalize_loan_type)
    return work


def get_officers(input_path: str) -> List[str]:
    work = _prepare_base(input_path)
    return sorted(work["FieldOfficer"].dropna().unique().tolist())


def run(
    input_path: str,
    output_dir: str,
    day_header: str = "D19",
    new_targets_input: str = "",
    mtd_targets_input: str = "",
    daily_targets_input: str = "",
) -> dict:
    os.makedirs(output_dir, exist_ok=True)

    work = _prepare_base(input_path)

    pivot = work.groupby(["FieldOfficer", "LoanType"]).size().unstack(fill_value=0)
    if "New" not in pivot.columns:
        pivot["New"] = 0
    if "Repeat" not in pivot.columns:
        pivot["Repeat"] = 0
    pivot = pivot.sort_index()
    pivot["Grand Total"] = pivot["New"] + pivot["Repeat"]

    officers = pivot.index.tolist()
    new_targets = _parse_targets(new_targets_input, officers)
    mtd_targets = _parse_targets(mtd_targets_input, officers)
    daily_targets = _parse_targets(daily_targets_input, officers)

    rows: List[Dict] = []
    for officer in officers:
        new_count = int(pivot.at[officer, "New"])
        repeat_count = int(pivot.at[officer, "Repeat"])
        grand_total = int(pivot.at[officer, "Grand Total"])
        new_target = int(new_targets.get(officer, 0))
        mtd_target = int(mtd_targets.get(officer, 0))
        daily_target = int(daily_targets.get(officer, 0))

        if mtd_target > 0:
            ds_value = grand_total - mtd_target
            today_target = 0 if ds_value >= 0 else abs(ds_value) + daily_target
            ds_display = f"+{ds_value}" if ds_value >= 0 else str(ds_value)
            mtd_display = mtd_target
            today_display = today_target
        else:
            ds_value = None
            today_target = None
            ds_display = "-"
            mtd_display = "-"
            today_display = "-"

        rows.append(
            {
                "field_officer": officer,
                "new_count": new_count,
                "repeat_count": repeat_count,
                "grand_total": grand_total,
                "new_target": new_target,
                "mtd_target": mtd_target,
                "daily_target": daily_target,
                "new_display": f"{new_count}/{new_target}" if new_target > 0 else f"{new_count}/-",
                "mtd_display": mtd_display,
                "ds_value": ds_value,
                "ds_display": ds_display,
                "today_target": today_target,
                "today_display": today_display,
                "is_grand_total": False,
            }
        )

    total_new = sum(r["new_count"] for r in rows)
    total_repeat = sum(r["repeat_count"] for r in rows)
    total_grand = sum(r["grand_total"] for r in rows)
    total_new_target = sum(r["new_target"] for r in rows)
    total_mtd_target = sum(r["mtd_target"] for r in rows)
    total_daily_target = sum(r["daily_target"] for r in rows)

    if total_mtd_target > 0:
        total_ds = total_grand - total_mtd_target
        total_today = 0 if total_ds >= 0 else abs(total_ds) + total_daily_target
        total_ds_display = f"+{total_ds}" if total_ds >= 0 else str(total_ds)
        total_mtd_display = total_mtd_target
        total_today_display = total_today
    else:
        total_ds = None
        total_today = None
        total_ds_display = "-"
        total_mtd_display = "-"
        total_today_display = "-"

    rows.append(
        {
            "field_officer": "Grand Total",
            "new_count": total_new,
            "repeat_count": total_repeat,
            "grand_total": total_grand,
            "new_target": total_new_target,
            "mtd_target": total_mtd_target,
            "daily_target": total_daily_target,
            "new_display": f"{total_new}/{total_new_target}" if total_new_target > 0 else f"{total_new}/-",
            "mtd_display": total_mtd_display,
            "ds_value": total_ds,
            "ds_display": total_ds_display,
            "today_target": total_today,
            "today_display": total_today_display,
            "is_grand_total": True,
        }
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"MTD_Disbursement_Pivot_Analysis_{timestamp}.xlsx")
    chart_file = _build_chart(rows, os.path.join(output_dir, "mtd_pivot_chart.png"))
    _write_excel(output_file, rows, day_header=day_header or "D19")

    on_track = sum(1 for r in rows if not r["is_grand_total"] and r["ds_value"] is not None and r["ds_value"] >= 0)
    below_target = sum(1 for r in rows if not r["is_grand_total"] and r["ds_value"] is not None and r["ds_value"] < 0)

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            "total_officers": int(len(officers)),
            "total_new": int(total_new),
            "total_repeat": int(total_repeat),
            "total_disbursed": int(total_grand),
            "total_mtd_target": int(total_mtd_target),
            "total_today_target": int(total_today or 0),
            "on_track_count": int(on_track),
            "below_target_count": int(below_target),
            "day_header": str(day_header or "D19"),
        },
    }
