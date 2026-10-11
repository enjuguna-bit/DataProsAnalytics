"""
Dormant Arrangement Processor
Android-safe adaptation of the advanced branch data processor workflow.
"""

import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import matplotlib
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from processors.tabular_io import read_tabular_file


def _read_input(path: str) -> pd.DataFrame:
    return read_tabular_file(path)


def _find_column_by_keywords(df: pd.DataFrame, keywords: List[str]) -> Optional[str]:
    for col in df.columns:
        col_lower = str(col).lower()
        for keyword in keywords:
            if keyword in col_lower:
                return col
    return None


def _normalize_phone_numbers_vectorized(df: pd.DataFrame, phone_cols: List[str], fill_na: bool) -> pd.DataFrame:
    if not phone_cols:
        return df

    out = df.copy()
    for phone_col in phone_cols:
        if phone_col not in out.columns:
            continue

        out[phone_col] = out[phone_col].fillna("").astype(str)
        out[phone_col] = out[phone_col].str.replace(r"\D", "", regex=True)

        mask_07 = out[phone_col].str.startswith("07") & (out[phone_col].str.len() == 10)
        mask_7 = out[phone_col].str.startswith("7") & (out[phone_col].str.len() == 9)
        mask_9digit = (out[phone_col].str.len() == 9) & (~out[phone_col].str.startswith("0"))
        mask_254 = out[phone_col].str.startswith("254") & (out[phone_col].str.len() == 12)

        out.loc[mask_07, phone_col] = "254" + out.loc[mask_07, phone_col].str[1:]
        out.loc[mask_7, phone_col] = "254" + out.loc[mask_7, phone_col]
        out.loc[mask_9digit, phone_col] = "254" + out.loc[mask_9digit, phone_col]

        invalid_mask = ~mask_254 & ~mask_07 & ~mask_7 & ~mask_9digit
        out.loc[invalid_mask & (out[phone_col].str.len() > 0), phone_col] = ""

        if fill_na:
            out.loc[out[phone_col] == "", phone_col] = "N/A"

    return out


def _deduplicate_phone_numbers(df: pd.DataFrame, phone_cols: List[str]) -> Tuple[pd.DataFrame, int]:
    if not phone_cols:
        return df, 0

    clean = df.copy()
    duplicates_found = 0

    for phone_col in phone_cols:
        if phone_col not in clean.columns:
            continue
        mask = (
            clean.duplicated(subset=[phone_col], keep="first")
            & (clean[phone_col] != "")
            & (clean[phone_col] != "N/A")
        )
        duplicates_found += int(mask.sum())
        clean = clean[~mask]

    return clean, duplicates_found


def _perform_quality_checks(df: pd.DataFrame) -> Dict[str, object]:
    checks: Dict[str, object] = {
        "Total Records": len(df),
        "Complete Records": int(df.notna().all(axis=1).sum()),
        "Records with Missing Values": int(df.isna().any(axis=1).sum()),
        "Duplicate Records": int(df.duplicated().sum()),
        "Unique Customers": 0,
        "Average LoanCount": 0,
    }

    phone_cols = [
        col for col in df.columns if any(keyword in str(col).lower() for keyword in ["phone", "mobile", "number"])
    ]
    for phone_col in phone_cols:
        valid_pattern = re.compile(r"^254\d{9}$")
        valid_phones = df[phone_col].apply(
            lambda x: bool(valid_pattern.match(str(x)))
            if pd.notna(x) and str(x).strip() and str(x) != "N/A"
            else False
        ).sum()
        checks[f"Valid {phone_col}"] = int(valid_phones)
        checks[f"Invalid {phone_col}"] = int(df[phone_col].notna().sum() - valid_phones)

    if phone_cols and phone_cols[0] in df.columns:
        checks["Unique Customers"] = int(df[phone_cols[0]].nunique())

    loancount_col = _find_column_by_keywords(df, ["loancount"])
    if loancount_col and loancount_col in df.columns:
        checks["Average LoanCount"] = round(float(pd.to_numeric(df[loancount_col], errors="coerce").mean()), 2)

    date_cols = [
        col for col in df.columns if any(keyword in str(col).lower() for keyword in ["date", "cleared"])
    ]
    for date_col in date_cols:
        parsed = pd.to_datetime(df[date_col], errors="coerce")
        today = pd.Timestamp.now()
        checks["Records with Future Dates"] = int((parsed > today).sum())
        checks["Records with Very Old Dates"] = int((parsed < pd.Timestamp("2000-01-01")).sum())
        valid_dates = parsed.dropna()
        if not valid_dates.empty:
            checks["Earliest Date"] = valid_dates.min().strftime("%Y-%m-%d")
            checks["Latest Date"] = valid_dates.max().strftime("%Y-%m-%d")
        break

    return checks


def _safe_sheet_name(name: str, used: set) -> str:
    raw = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "Sheet"
    raw = raw[:31]
    candidate = raw
    idx = 1
    while candidate in used:
        suffix = f"_{idx}"
        candidate = f"{raw[: 31 - len(suffix)]}{suffix}"
        idx += 1
    used.add(candidate)
    return candidate


def _format_excel_workbook(filepath: str):
    wb = load_workbook(filepath)
    header_font = Font(name="Arial", size=11, bold=True, color="000000")
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    row_fill_even = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment
            cell.border = thin_border

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.alignment = alignment
                cell.border = thin_border
                cell.fill = row_fill_even if row % 2 else row_fill_odd

        ws.freeze_panes = "A2"

        for col in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col)
            max_length = 0
            for cell in ws[col_letter]:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

    wb.save(filepath)


def _process_branch_df(
    branch_df: pd.DataFrame, remove_duplicates: bool, fill_na: bool
) -> Tuple[pd.DataFrame, Dict[str, object], int, List[str], Optional[str], Optional[str], float]:
    # Remove branch and datecreated columns
    branch_cols_remove = [col for col in branch_df.columns if "branch" in str(col).lower()]
    if branch_cols_remove:
        branch_df = branch_df.drop(columns=branch_cols_remove)

    datecreated_cols_remove = [col for col in branch_df.columns if "datecreated" in str(col).lower()]
    if datecreated_cols_remove:
        branch_df = branch_df.drop(columns=datecreated_cols_remove)

    phone_cols = [
        col
        for col in branch_df.columns
        if any(keyword in str(col).lower() for keyword in ["phone", "mobile", "number", "borrowerphone"])
    ]
    if phone_cols:
        branch_df = _normalize_phone_numbers_vectorized(branch_df, phone_cols, fill_na=fill_na)

    duplicates_removed = 0
    if remove_duplicates and phone_cols:
        branch_df, duplicates_removed = _deduplicate_phone_numbers(branch_df, phone_cols)

    ro_col = _find_column_by_keywords(branch_df, ["ro", "relationship", "officer", "portfolio"])
    date_cleared_col = _find_column_by_keywords(branch_df, ["date", "cleared", "dateloancleared"])
    firstname_col = _find_column_by_keywords(branch_df, ["firstname", "first", "name"])

    if date_cleared_col:
        branch_df[date_cleared_col] = pd.to_datetime(branch_df[date_cleared_col], errors="coerce")

    if ro_col:
        officers = sorted([str(x) for x in branch_df[ro_col].dropna().unique()])
        sorted_data = pd.DataFrame()
        for officer in officers:
            officer_data = branch_df[branch_df[ro_col] == officer].copy()
            if date_cleared_col:
                officer_data = officer_data.sort_values(by=date_cleared_col, ascending=False)
            elif firstname_col:
                officer_data = officer_data.sort_values(by=firstname_col, ascending=True)
            sorted_data = pd.concat([sorted_data, officer_data], ignore_index=True)
        if sorted_data.empty:
            if date_cleared_col:
                sorted_data = branch_df.sort_values(by=date_cleared_col, ascending=False)
            elif firstname_col:
                sorted_data = branch_df.sort_values(by=firstname_col, ascending=True)
            else:
                sorted_data = branch_df
    else:
        if date_cleared_col:
            sorted_data = branch_df.sort_values(by=date_cleared_col, ascending=False)
        elif firstname_col:
            sorted_data = branch_df.sort_values(by=firstname_col, ascending=True)
        else:
            sorted_data = branch_df

    quality_checks = _perform_quality_checks(sorted_data)

    loan_balance_col = _find_column_by_keywords(sorted_data, ["loanbalance", "balance"])
    total_balance = 0.0
    if loan_balance_col and loan_balance_col in sorted_data.columns:
        total_balance = float(pd.to_numeric(sorted_data[loan_balance_col], errors="coerce").fillna(0).sum())

    return (
        sorted_data,
        quality_checks,
        duplicates_removed,
        phone_cols,
        ro_col,
        date_cleared_col,
        total_balance,
    )


def _build_chart(summary_df: pd.DataFrame, chart_path: str) -> str:
    try:
        fig, ax = plt.subplots(figsize=(9, 5))
        chart_df = summary_df.sort_values("Records", ascending=False).head(10)
        ax.bar(chart_df["Branch"], chart_df["Records"], color="#f39c12", edgecolor="white")
        ax.set_title("Top Branches by Processed Records", fontweight="bold")
        ax.set_ylabel("Records")
        ax.tick_params(axis="x", rotation=45)
        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _to_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def get_branches(input_path: str) -> List[str]:
    df = _read_input(input_path)
    if df is None or df.empty:
        return []
    branch_col = _find_column_by_keywords(df, ["branch"])
    if not branch_col:
        return []
    return sorted([str(x) for x in df[branch_col].dropna().unique()])


def run(
    input_path: str,
    output_dir: str,
    process_all_branches: bool = True,
    selected_branch: str = "",
    remove_duplicates: bool = True,
    fill_na: bool = True,
    add_formatting: bool = True,
) -> dict:
    os.makedirs(output_dir, exist_ok=True)

    df = _read_input(input_path)
    if df is None or df.empty:
        raise ValueError("Input file is empty")

    branch_col = _find_column_by_keywords(df, ["branch"])
    if not branch_col:
        raise ValueError("No branch column found in the file")

    branches = sorted([str(x) for x in df[branch_col].dropna().unique()])
    if not branches:
        raise ValueError("No branches found in the file")

    process_all = _to_bool(process_all_branches)
    remove_duplicates = _to_bool(remove_duplicates)
    fill_na = _to_bool(fill_na)
    add_formatting = _to_bool(add_formatting)

    if not process_all:
        selected = str(selected_branch or "").strip()
        if not selected:
            raise ValueError("Please select a branch when 'Process all branches' is disabled")

        selected_map = {str(b).strip().lower(): b for b in branches}
        selected_key = selected.lower()
        if selected_key not in selected_map:
            raise ValueError(f"Selected branch not found: {selected}")
        branches = [selected_map[selected_key]]

    used_sheet_names = set()
    processed_branches = []
    summary_rows = []
    quality_rows = []
    total_records = 0
    total_balance = 0.0
    total_duplicates_removed = 0

    for branch in branches:
        branch_data = df[df[branch_col] == branch].copy()
        if branch_data.empty:
            continue

        (
            sorted_data,
            quality_checks,
            duplicates_removed,
            phone_cols,
            ro_col,
            date_cleared_col,
            branch_balance,
        ) = _process_branch_df(branch_data, remove_duplicates=remove_duplicates, fill_na=fill_na)

        total_records += len(sorted_data)
        total_balance += branch_balance
        total_duplicates_removed += duplicates_removed

        summary_rows.append(
            {
                "Branch": branch,
                "Records": len(sorted_data),
                "Duplicates Removed": duplicates_removed,
                "Phone Columns": len(phone_cols),
                "Officer Column": ro_col or "",
                "Date Column": date_cleared_col or "",
                "Loan Balance": round(branch_balance, 2),
            }
        )

        for metric, value in quality_checks.items():
            quality_rows.append({"Branch": branch, "Metric": metric, "Value": value})

        processed_branches.append((branch, sorted_data))

    if not processed_branches:
        raise ValueError("No branch data could be processed")

    summary_df = pd.DataFrame(summary_rows)
    quality_df = pd.DataFrame(quality_rows)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"Dormant_Arrangement_Report_{timestamp}.xlsx")

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        quality_df.to_excel(writer, sheet_name="Quality Checks", index=False)

        for branch, data in processed_branches:
            sheet_name = _safe_sheet_name(str(branch), used_sheet_names)
            data.to_excel(writer, sheet_name=sheet_name, index=False)

    if add_formatting:
        _format_excel_workbook(output_file)

    chart_file = _build_chart(summary_df, os.path.join(output_dir, "dormant_chart.png"))

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            # Backward-compatible keys used by existing DormantFragment.
            "total_dormant": int(total_records),
            "total_balance": round(float(total_balance), 2),
            "num_branches": int(len(processed_branches)),
            "critical_count": int(total_duplicates_removed),
            # Additional explicit keys for branch processor semantics.
            "total_records": int(total_records),
            "total_loan_balance": round(float(total_balance), 2),
            "duplicates_removed": int(total_duplicates_removed),
        },
    }
