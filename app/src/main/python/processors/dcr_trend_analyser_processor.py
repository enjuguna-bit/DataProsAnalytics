"""
DCR Trend Analyser Processor

Compares a morning CR snapshot against a current CR snapshot and exports a
colour-coded Excel workbook that highlights the movement between the two files.
"""

from __future__ import annotations

import json
import os
from datetime import datetime
from typing import Dict, List, Tuple

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from processors.tabular_io import read_tabular_file


DEFAULT_THRESHOLDS = {
    "thresholds": [
        {"min_change": 50, "status": "Most Improved", "comment": "Most improved", "color": "A7F3D0"},
        {"min_change": 25, "status": "Excellent Improvement", "comment": "Excellent improvement", "color": "D1FAE5"},
        {"min_change": 10, "status": "Strong Improvement", "comment": "Strong improvement", "color": "DCFCE7"},
        {"min_change": 1, "status": "Slight Improvement", "comment": "Slight improvement", "color": "FEF3C7"},
        {"min_change": 0, "status": "No Change", "comment": "No change", "color": "E5E7EB"},
        {"min_change": -9999, "status": "Declined", "comment": "Needs attention", "color": "FEE2E2"},
    ],
    "default_color": "FEE2E2",
    "default_comment": "Needs attention",
}


def _threshold_config_path(config_file: str) -> str:
    if os.path.isabs(config_file):
        return config_file
    return os.path.join(os.path.dirname(__file__), config_file)


def load_thresholds(config_file: str = "cr_thresholds.json") -> dict:
    try:
        with open(_threshold_config_path(config_file), "r", encoding="utf-8") as handle:
            data = json.load(handle)
            if isinstance(data, dict):
                thresholds = data.get("thresholds", [])
                if isinstance(thresholds, list):
                    data["thresholds"] = sorted(
                        [item for item in thresholds if isinstance(item, dict)],
                        key=lambda item: float(item.get("min_change", 0)),
                        reverse=True,
                    )
                return data
    except Exception:
        print(f"⚠️ Warning: Could not load {config_file}. Using defaults.")
    return DEFAULT_THRESHOLDS.copy()


THRESHOLDS = load_thresholds()


def _read_input(path: str) -> pd.DataFrame:
    return read_tabular_file(path, dtype=str)


def _normalize_header(name: str) -> str:
    clean = str(name or "").strip().lower()
    clean = clean.replace("%", "pct")
    return "".join(ch for ch in clean if ch.isalnum())


def _pick_column(columns: List[str], candidates: List[str]) -> str:
    lookup = {_normalize_header(col): col for col in columns}
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in lookup:
            return lookup[key]
    for candidate in candidates:
        key = _normalize_header(candidate)
        for col in columns:
            col_key = _normalize_header(col)
            if key in col_key or col_key in key:
                return col
    return ""


def _to_number(series: pd.Series) -> pd.Series:
    return (
        series.fillna("")
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(r"[^\d\.\-]", "", regex=True)
        .replace("", "0")
        .pipe(pd.to_numeric, errors="coerce")
        .fillna(0.0)
    )


def _clean_name(value) -> str:
    return str(value or "").strip().replace("\u00a0", " ").replace("\t", " ")


def clean_cr(value) -> float:
    if isinstance(value, str):
        value = value.strip().replace("%", "").strip()
    try:
        return float(value)
    except Exception:
        return 0.0


def get_trend_decoration(change: float) -> dict:
    thresholds = sorted(
        THRESHOLDS.get("thresholds", []),
        key=lambda item: float(item.get("min_change", 0)),
        reverse=True,
    )
    for threshold in thresholds:
        if change >= float(threshold.get("min_change", 0)):
            color = str(threshold.get("color", THRESHOLDS.get("default_color", "FEE2E2"))).strip()
            if color.startswith("#"):
                color = color[1:]
            comment = threshold.get("comment", THRESHOLDS.get("default_comment", "Needs attention"))
            status = threshold.get("status", comment)
            return {
                "status": status,
                "comment": comment,
                "fill": PatternFill(start_color=color, end_color=color, fill_type="solid"),
            }

    color = str(THRESHOLDS.get("default_color", "FEE2E2")).strip()
    if color.startswith("#"):
        color = color[1:]
    return {
        "status": "Declined",
        "comment": THRESHOLDS.get("default_comment", "Needs attention"),
        "fill": PatternFill(start_color=color, end_color=color, fill_type="solid"),
    }


def get_comment_and_color(change: float):
    decoration = get_trend_decoration(change)
    return decoration["comment"], decoration["fill"]


def _prepare_snapshot(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        raise ValueError("Input file is empty")

    df = df.copy()
    df.columns = [str(col).strip() for col in df.columns]
    columns = list(df.columns)

    name_col = _pick_column(columns, ["Full Name", "FullName", "Name", "Portfolio", "Officer"])
    cr_col = _pick_column(columns, ["Collection Rate (%)", "CollectionRate", "CR%", "CR", "Collection Rate"])
    total_col = _pick_column(columns, ["Total Amount", "TotalAmount", "Amount", "Total", "Amount Due"])
    paid_col = _pick_column(columns, ["Paid Amount", "PaidAmount", "Paid", "Amount Paid"])

    if not name_col:
        raise ValueError("Missing required column: Full Name")
    if not cr_col and (not total_col or not paid_col):
        raise ValueError("Missing required columns. Need either Collection Rate (%) or Total/Paid Amount columns.")

    names = df[name_col].fillna("").astype(str).map(_clean_name).str.replace(r"\s+", " ", regex=True).str.upper()
    mask = names != ""
    if not mask.any():
        raise ValueError("No valid rows found after cleaning Full Name")

    base = df.loc[mask].copy()
    out = pd.DataFrame(index=base.index)
    out["Full Name"] = names.loc[mask]

    if cr_col:
        current_cr = _to_number(base[cr_col]).clip(lower=0, upper=100).round(2)
    else:
        total_amount = _to_number(base[total_col]).replace(0, pd.NA)
        paid_amount = _to_number(base[paid_col])
        current_cr = ((paid_amount / total_amount) * 100).fillna(0.0).clip(lower=0, upper=100).round(2)

    out["Collection Rate (%)"] = current_cr.reindex(base.index)
    out = out[out["Collection Rate (%)"].notna()].copy()

    if out.empty:
        raise ValueError("No usable portfolios found after cleaning the input")

    out = out.sort_values(by=["Collection Rate (%)"], ascending=[True]).reset_index(drop=True)
    return out


def _build_chart(df: pd.DataFrame, chart_path: str) -> str:
    try:
        body = df.copy()
        if body.empty:
            return ""

        body = body.assign(abs_change=body["CR Change (%)"].abs())
        if len(body) > 20:
            body = body.nlargest(20, "abs_change").sort_values("CR Change (%)", ascending=True)
        else:
            body = body.sort_values("CR Change (%)", ascending=True)

        labels = body["Full Name"].tolist()
        values = body["CR Change (%)"].tolist()
        colors = ["#2E7D32" if value > 0 else "#C62828" if value < 0 else "#64748B" for value in values]

        height = max(5, min(14, 0.45 * len(labels) + 2))
        width = max(9, min(16, 0.55 * len(labels) + 6))
        fig, ax = plt.subplots(figsize=(width, height))

        positions = list(range(len(labels)))
        ax.barh(positions, values, color=colors, edgecolor="white")
        ax.axvline(0, color="#263238", linestyle="--", linewidth=1.0)
        ax.set_yticks(positions)
        ax.set_yticklabels(labels)
        ax.set_xlabel("CR Change (%)")
        ax.set_title("DCR Trend Analyser - CR Movement", fontweight="bold")
        ax.grid(axis="x", linestyle=":", alpha=0.35)

        for idx, value in enumerate(values):
            if len(labels) <= 15:
                offset = 0.3 if value >= 0 else -0.3
                ax.text(value + offset, idx, f"{value:.2f}", va="center", fontsize=8)

        plt.tight_layout()
        fig.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        return chart_path
    except Exception:
        plt.close("all")
        return ""


def _auto_width(ws):
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 3, 12), 42)


def _apply_border(cell, border):
    cell.border = border


def _write_workbook(df: pd.DataFrame, output_file: str, chart_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Trend Analysis"

    columns = [
        "Full Name",
        "Morning CR (%)",
        "Current CR (%)",
        "CR Change (%)",
        "Trend Status",
        "Trend Comment",
    ]

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    left = Alignment(horizontal="left", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    for col_idx, header in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center if col_idx > 1 else left
        _apply_border(cell, border)

    numeric_columns = {
        "Morning CR (%)",
        "Current CR (%)",
        "CR Change (%)",
    }
    center_columns = {"Trend Status"}

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        decoration = get_trend_decoration(float(row["CR Change (%)"]))
        row_fill = decoration["fill"]
        for col_idx, header in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(header, ""))
            cell.fill = row_fill
            _apply_border(cell, border)

            if header in numeric_columns:
                cell.alignment = right
                cell.number_format = "0.00"
                if header == "Current CR (%)" and float(row.get(header, 0.0)) < 80:
                    cell.font = Font(color="C62828")
                else:
                    cell.font = Font(color="000000")
            elif header in center_columns:
                cell.alignment = center
                cell.font = Font(color="000000")
            else:
                cell.alignment = left
                cell.font = Font(color="000000")

    ws.freeze_panes = "A2"
    _auto_width(ws)

    summary = wb.create_sheet("Summary")
    title_fill = PatternFill(start_color="0D47A1", end_color="0D47A1", fill_type="solid")
    title_font = Font(bold=True, color="FFFFFF", size=13)
    summary.merge_cells("A1:B1")
    title = summary["A1"]
    title.value = "DCR TREND SUMMARY"
    title.fill = title_fill
    title.font = title_font
    title.alignment = center

    summary_headers = ["Metric", "Value"]
    for col_idx, header in enumerate(summary_headers, start=1):
        cell = summary.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        _apply_border(cell, border)

    summary_rows = [
        ("Total Portfolios", int(len(df))),
        ("Morning Avg CR (%)", round(float(df["Morning CR (%)"].mean()), 2)),
        ("Current Avg CR (%)", round(float(df["Current CR (%)"].mean()), 2)),
        ("Average Change (%)", round(float(df["CR Change (%)"].mean()), 2)),
        ("Most Improved Rows", int((df["Trend Status"] == "Most Improved").sum())),
        ("Excellent Improvement Rows", int((df["Trend Status"] == "Excellent Improvement").sum())),
        ("Strong Improvement Rows", int((df["Trend Status"] == "Strong Improvement").sum())),
        ("Slight Improvement Rows", int((df["Trend Status"] == "Slight Improvement").sum())),
        ("No Change Rows", int((df["Trend Status"] == "No Change").sum())),
        ("Declined Rows", int((df["CR Change (%)"] < 0).sum())),
        ("Best Improvement (%)", round(float(df["CR Change (%)"].max()), 2)),
        ("Worst Decline (%)", round(float(df["CR Change (%)"].min()), 2)),
    ]

    for row_idx, (metric, value) in enumerate(summary_rows, start=3):
        metric_cell = summary.cell(row=row_idx, column=1, value=metric)
        value_cell = summary.cell(row=row_idx, column=2, value=value)
        metric_cell.font = Font(bold=True)
        metric_cell.alignment = left
        value_cell.alignment = right
        _apply_border(metric_cell, border)
        _apply_border(value_cell, border)

        if isinstance(value, (int, float)):
            if "CR" in metric or "Change" in metric or "%" in metric or "Improvement" in metric or "Decline" in metric:
                value_cell.number_format = "0.00"
            else:
                value_cell.number_format = "#,##0"

    top_start = len(summary_rows) + 5
    summary.cell(row=top_start, column=1, value="TOP 5 IMPROVEMENTS").font = Font(bold=True)
    top_improvements = df.nlargest(5, "CR Change (%)")[["Full Name", "CR Change (%)"]]
    for idx, row in enumerate(top_improvements.itertuples(index=False), start=top_start + 1):
        summary.cell(row=idx, column=1, value=row[0])
        summary.cell(row=idx, column=2, value=float(row[1]))
        summary.cell(row=idx, column=2).number_format = "0.00"

    top_decline_start = top_start + len(top_improvements) + 3
    summary.cell(row=top_decline_start, column=1, value="TOP 5 DECLINES").font = Font(bold=True)
    top_declines = df.nsmallest(5, "CR Change (%)")[["Full Name", "CR Change (%)"]]
    for idx, row in enumerate(top_declines.itertuples(index=False), start=top_decline_start + 1):
        summary.cell(row=idx, column=1, value=row[0])
        summary.cell(row=idx, column=2, value=float(row[1]))
        summary.cell(row=idx, column=2).number_format = "0.00"

    _auto_width(summary)

    if chart_path and os.path.exists(chart_path):
        from openpyxl.drawing.image import Image as XLImage

        chart_ws = wb.create_sheet("Chart")
        chart_ws.add_image(XLImage(chart_path), "A1")

    wb.save(output_file)


def run(morning_path: str, current_path: str, output_dir: str) -> dict:
    os.makedirs(output_dir, exist_ok=True)

    morning_df = _prepare_snapshot(_read_input(morning_path))
    current_df = _prepare_snapshot(_read_input(current_path))

    morning = morning_df.rename(columns={"Collection Rate (%)": "Morning CR (%)"})
    current = current_df.rename(columns={"Collection Rate (%)": "Current CR (%)"})

    merged = morning.merge(current, on="Full Name", how="outer")

    for column in ["Morning CR (%)", "Current CR (%)"]:
        if column in merged.columns:
            merged[column] = pd.to_numeric(merged[column], errors="coerce").fillna(0.0)

    merged["CR Change (%)"] = (merged["Current CR (%)"] - merged["Morning CR (%)"]).round(2)
    trend_meta = merged["CR Change (%)"].apply(lambda value: get_trend_decoration(float(value)))
    merged["Trend Status"] = trend_meta.apply(lambda item: item["status"])
    merged["Trend Comment"] = trend_meta.apply(lambda item: item["comment"])

    merged = merged[
        [
            "Full Name",
            "Morning CR (%)",
            "Current CR (%)",
            "CR Change (%)",
            "Trend Status",
            "Trend Comment",
        ]
    ].copy()
    merged = merged.sort_values(by=["CR Change (%)", "Current CR (%)"], ascending=[False, False]).reset_index(drop=True)

    chart_path = os.path.join(output_dir, "dcr_trend_analyser_chart.png")
    chart_file = _build_chart(merged, chart_path)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"DCR_Trend_Analyser_{timestamp}.xlsx")
    _write_workbook(merged, output_file, chart_file)

    return {
        "success": True,
        "output_file": output_file,
        "chart_file": chart_file,
        "stats": {
            "total_portfolios": int(len(merged)),
            "morning_avg_cr": round(float(merged["Morning CR (%)"].mean()), 2),
            "current_avg_cr": round(float(merged["Current CR (%)"].mean()), 2),
            "avg_change_pct": round(float(merged["CR Change (%)"].mean()), 2),
            "improved_count": int((merged["CR Change (%)"] > 0).sum()),
            "declined_count": int((merged["CR Change (%)"] < 0).sum()),
            "unchanged_count": int((merged["CR Change (%)"] == 0).sum()),
            "best_improvement_pct": round(float(merged["CR Change (%)"].max()), 2),
            "worst_decline_pct": round(float(merged["CR Change (%)"].min()), 2),
        },
    }
